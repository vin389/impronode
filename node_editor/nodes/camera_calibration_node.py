# node_editor/nodes/camera_calibration_node.py
"""
Camera Calibration: one node for single-image calibration from
object / image points (with '?' unknown coordinates and straight-line
points), chessboard calibration, and both combined -- and, in later
phases, manual parameter adjustment and 3D grids. The maths is in
node_editor/calib_engine.py (pure numpy / OpenCV / scipy, tested against
synthetic ground truth). Chessboard detection and image ranking are in
node_editor/chessboard_detect.py, the settings search in chessboard_search.py.
(This node was called 'Camera Calibration (Unified)'; the older 'Camera
Calibration' and 'Chessboard Calibration' nodes it replaces are in
recycle_bin/.)

One solve uses up to two kinds of views: the POINTS image (the table) and
the CHECKED chessboard images. The result stores one camera (K, D), the
pose of the points image and the pose of every chessboard image; the
'world frame' choice decides which pose goes to the rvec / tvec outputs.
"""

from __future__ import annotations

import copy
import csv
import dataclasses
import time
import json
import math
import os
import threading
import tkinter as tk
from pathlib import Path
from tkinter import messagebox, ttk

import cv2
import numpy as np
from PIL import Image, ImageTk

from node_editor import calib_engine as ce
from node_editor import chessboard_search as cbs
from node_editor import project_data
from node_editor.calib_engine import merge_initial, run_plan
from node_editor.base_node import BaseNode
from node_editor.execution import ExecutionMode
from node_editor.chessboard_detect import (
    _detect_worker, _read_rgb, rank_calibration_images, take_ranked_images)
from node_editor.pin_types import PinDef, PinSchema, PinType
from node_editor.project_context import get_project_directory
from node_editor.tooltip import add_tooltip, set_tooltips_enabled, tooltip_toggle_var
from node_editor.ui_style import apply_notebook_tab_style

TABLE_HEADER = ("name", "X", "Y", "Z", "u", "v", "line", "weight", "use")
ERROR_GAINS = ("1", "2", "5", "10", "20", "50", "100", "200")
DEFAULT_FIXED = {"k3"} | set(ce.DIST_NAMES[5:])
POSE_PARAMS = ("rvec x", "rvec y", "rvec z", "tvec x", "tvec y", "tvec z")
GRID_COLORS = {"cyan": "#00e5ff", "yellow": "#ffe600", "orange": "#ff9f1a", "green": "#3ddc84",
               "magenta": "#ff5cf0", "white": "#ffffff", "red": "#ff3b30", "blue": "#4d8dff"}

# Point colours in the image view and the table.
KIND_COLORS = {
    "point": "#3ddc84",       # fully known object point
    "unknown": "#ffa94d",     # one '?' coordinate
    "line": "#4dc9f6",        # image-only point of a straight 3D line
    "invalid": "#ff4d4d",
    "unpicked": "#bbbbbb",
    "unused": "#8a8a8a",
}
PROJ_COLOR = "#ff5cf0"
SELECT_COLOR = "#ffe600"
CURVE_COLOR = "#4dc9f6"

_SYNTAX_HELP = ("X / Y / Z cell: a number; '?' = unknown (solved); '?=40' = unknown starting at 40. "
                "At most ONE '?' per row. All three empty + a line name = an image-only point of "
                "that straight 3D line.")

HELP_TEXT = """\
Camera Calibration
==================
Finds the camera matrix (fx, fy, cx, cy), the lens distortion and the camera
pose (rvec, tvec) from points in ONE image -- even when some 3D coordinates
are unknown, and with straight lines that only need image points.

QUICK START
  1. Connect an image (or 'Open image...' in the Image view).
  2. Points tab: add rows with the 3D object coordinates X, Y, Z (or paste
     them from Excel with Ctrl+V), then select a row and click its point in
     the image (the next row without an image point is selected for you).
  3. Solve tab: choose which parameters are FIXED, then press Solve.
  4. Read the errors in the table / image; outputs go to the pins.

OBJECT COORDINATES (X / Y / Z cells)
  12.5      a known coordinate
  ?         an UNKNOWN coordinate: it is solved together with the camera
  ?=40      unknown, starting the search at 40 (helps when the start is far)
  At most ONE '?' per row: one image point fixes one unknown (its viewing
  ray meets the 3D line of possible positions in one point). Two or three
  '?' cannot be solved -- those cells are shown in red and the row is
  ignored. At least 4 rows must be fully known (no '?') to find a starting
  camera pose.

STRAIGHT LINES ('line' column)
  Points with the same line name (e.g. L1) are images of points on ONE
  straight 3D line -- an edge of a building, a rail, a pole -- whose 3D
  coordinates you do not know. The camera without distortion images a 3D
  line as a straight line, so these points mainly fix the DISTORTION.
  Error of a line point = its distance to the best straight line through
  the line's points after removing the distortion, in pixels (normalised
  image plane distance x mean focal length). At least 3 points per line;
  more, spread along the whole visible line, is better.
  Click mode 'Line' adds a row to the chosen line with every click.
  A row may have both object coordinates and a line name: it then counts
  in both.

WEIGHTS
  'weight' column: per row (0 = ignore, 2 = twice as important ...).
  Solve tab: one weight per error type (known points, '?' points, lines).
  Robust (Huber): errors larger than the given pixels count less -- use it
  when a few points may be clicked wrongly; they then stand out in red.

SOLVER (Solve tab, line under Solve)
  'OpenCV only': when nothing needs more, cv2.calibrateCamera (and
  cv2.solvePnP for the points image's pose in 'two steps') solves the whole
  problem -- the same least-squares problem, fast. Otherwise this node's
  least-squares minimization (scipy, started from OpenCV) runs; the line
  says why: '?' unknowns, straight lines, the huber loss, different
  weights, or fixed parameters OpenCV can only fix together (fx with fy,
  cx with cy, p1 with p2, s1-s4, taux with tauy). 'Always refine' runs the
  minimization anyway (OpenCV stops after 30 iterations; with many free
  distortion coefficients that can leave ~0.0001 px RMS). 'Cancel' stops a
  running Solve; the previous result stays.

FIXED PARAMETERS
  With one image and few points, not every parameter can be found: fix the
  ones you know or that matter less (e.g. k3; p1 p2 for a good lens; cx cy
  at the image centre). 'Initial' is the start value -- and the value used
  when the parameter is fixed. Empty = automatic (focal length = the larger
  image side, principal point = image centre, distortion 0).
  The result is warned about when the points cover only the image centre:
  distortion beyond the points' radius is then an extrapolation.

IMAGE VIEW
  wheel = zoom at the cursor, left-drag on empty space = pan, 'Fit' =
  whole image. left-click = set the selected row's image point (or add a
  row), left-drag a point = move it. Snap: refine a click to the nearest corner
  (sub-pixel, like chessboard corners).
  Drawn: measured points (green known, orange '?', blue line-only, red
  invalid), projected points (magenta +) with the error vector magnified
  by 'Error x', fitted lines (blue curve, bent by the distortion) with
  each point's distance to it, and for '?' points the 3D line of possible
  positions (dashed orange).

TABLE
  Click = select (Ctrl / Shift: several), double-click = edit, click the
  'use' cell = on / off, Delete = remove rows, Ctrl+V = paste rows,
  arrows = move the selection. Import / Export: CSV or Excel with the
  columns name, X, Y, Z, u, v, line, weight, use.
  Pasting without a header: columns are name (optional), then 2 = u v,
  3 = X Y Z, 5 = X Y Z u v, 6 = X Y Z u v line.

CHESSBOARD (Chessboard tab)
  Images come from the 'files' input pin (e.g. an Image Sequence node) or
  from 'Add images...'. Give the INNER corner counts (columns = corners per
  row, rows = corners per column; one less than the squares) and the
  square size (your length unit), then 'Detect'. Every image with a board
  is RANKED: the best (sharp, large, centred) first, then always the one
  most DIFFERENT from those before it. 'Auto-select' checks the first
  'target' images; check / uncheck images yourself in the 'use' column.
  Click a row to see the image with its corners (and, after Solve, the
  projected corners with their errors, and the board's X / Y / Z axes).
  Board coordinates: corner [row 0, column 0] = (0, 0, 0), X along a row,
  Y down the columns, Z = X x Y (into the board), as in OpenCV.
  Square X / square Y: the size of one square along a board row (X) and
  along a board column (Y); an empty Y = the same as X. Usually equal; a printer may scale
  one direction by a fraction of a percent -- measure many squares and
  enter both. When they differ, every image must show which board
  direction is X: columns x rows must then be the FULL board's inner
  corners, and an image whose corners would fit the board either way round
  (a board with as many columns as rows, or a small part of a board found
  by a settings search) is marked 'X/Y?', is not auto-checked and cannot be
  used. Such an image keeps its own X along its lattice columns otherwise.

SEARCH SETTINGS (Chessboard tab, 'Search settings...')
  Tries EVERY combination of the given columns, rows, detect widths and
  the three options (on / off / both) on every image -- long, with a
  progress bar, the remaining time and Cancel. Per image, all corners any
  setting found are put on ONE board lattice: two detections are joined
  only when at least 3 corners (not all on one line) coincide and exactly
  one shift / rotation of the lattice explains all of them, and no other
  corner contradicts it; anything else is rejected, never guessed. A
  detection whose grid is not smooth (a false corner, e.g. on the edge of
  an object in front of the board) is rejected too. A corner found by
  several settings counts once; its position is refined once more from
  where they agree. So an image may give a corner set that is not a full
  rectangle -- correct for calibration, since every board image has its
  own pose. Its object points: corner (row r, column c) = (c x square X,
  r x square Y, 0); r and c start at 0 at the top-left of the corners
  found (that origin may be a corner that was not found). With square X
  different from Y, see 'Square X / square Y' above.
  'Images' list: corners, how many settings found a board, the best single
  setting, the gain. 'Settings' list: every setting alone, best first;
  'Copy to the Detect fields' if you prefer one fixed pattern.
  'Use these corners' replaces the board list (then ranked as usual).

COMBINING POINTS AND CHESSBOARD (Solve tab, 'Use')
  points only        the single-image calibration above
  chessboard only    a classic calibration from the checked images
  both, jointly      one least-squares problem with all errors (weights
                     per type) -- the board images fix the intrinsics and
                     distortion, the points image adds its pose (and helps
                     with its own region of the image)
  both, two steps    1) chessboard images only -> camera; 2) that camera
                     fixed, only the pose of the points image (and its
                     '?' unknowns) from the table. Use it when the points
                     image must not change the camera.
  The chessboard images must have the size of the points image.

WORLD FRAME
  rvec / tvec outputs = the pose of ONE view: the points image (its object
  coordinates are the world) or one chessboard image ('Set as world
  frame': e.g. a board lying on the floor at the measuring place defines
  the world). rvecs / tvecs give every checked board's pose.

ADJUST (manual parameters)
  'Add slider', choose a parameter (fx ... distortion ... rvec x/y/z,
  tvec x/y/z of the world-frame view), set its lower / upper end, and drag:
  the table errors, the image drawings, the summary and the weighted cost
  update while you drag; the output pins update when you release. The
  sliders change the CURRENT RESULT (shown as MANUALLY ADJUSTED);
  'Revert to solved values' undoes all manual changes. Type an exact value
  in the value box (Enter). Useful to see how sensitive the errors are to a
  parameter, or to set a value the data cannot determine.

GRIDS
  3D grids drawn into the image with the current camera (distortion
  included), e.g. the floor plane Z = 0 every 100 mm, or a wall. Each grid:
  X, Y, Z values as a list '0 50 120' or start:step:stop '0:100:1000'
  (one value = a plane); frame 'points' = drawn on the points image in its
  object coordinates, 'board' = drawn on every board image in board
  coordinates; colour. Lines along X, Y and Z are drawn through every
  combination of the other two coordinates.

RIGHT-SIDE VIEWS
  Image        the points image or a board image with everything drawn on it
  Errors       error of every item as bars (click = select it) and all
               residual vectors as a scatter plot
  3D           points + camera, all board cameras around the board, or all
               boards in front of the camera. left-drag = rotate, right-drag
               = pan, wheel = zoom, perspective slider, axes icon (lower left)
  Undistort    the shown image without distortion (alpha: 0 = no black
               borders, 1 = all pixels kept); line points must lie on
               straight lines here -- the largest distance is printed
  Sensitivity  leave one board image / table point out at a time and solve
               again (parallel processes): finds influential / bad items
  Runs         every Solve with its settings and result, to compare;
               Restore one, or copy its K, D to the initial values

INPUT PINS
  image                 the image (RGB / grey)
  files                chessboard image files (list of paths)
  object_points (N, 3)  NaN = unknown coordinate  } 'From input pins'
  image_points (N, 2)                             } fills the table
  cmat_in, dvec_in      a camera matrix / distortion, e.g. from the
                        chessboard node: 'Input cmat/dvec -> initial'.
  'Follow input pins' reloads the table from the pins and solves
  whenever they change (e.g. a pose per video frame with fixed intrinsics).

OUTPUT PINS
  cmat (3, 3), dvec (1, n), rvec (3, 1), tvec (3, 1)   (OpenCV convention:
  X_camera = R(rvec) X_object + tvec), rms (all errors, px), rms_types
  [point, ?, line, chessboard] (NaN = none), img_size [w, h], obj_pts
  (N, 3) with '?' solved, img_pts (N, 2), proj_pts (N, 2), pt_err (N,)
  for the rows used as points, in table order.
  rvecs, tvecs (n, 3)   pose of every checked chessboard image (table order)
  cb_err (n,)           RMS error per chessboard image (px)
  cb_obj_pts (n, M, 3), cb_img_pts (n, M, 2)   their corners (M = rows x cols;
                        after a settings search the images may have
                        different corner counts: M = the largest, NaN rows
                        pad the others)

Coordinates follow OpenCV: image (0, 0) is the centre of the top-left pixel.
"""


# ── chessboard candidates ──────────────────────────────────────────────────
# A candidate (one board image) keeps its corners as (N, 2) 'corners' with
# the lattice index (r, c) of every corner in 'rc' (N, 2): object point =
# (c * square, r * square, 0). A plain Detect gives the full nrow x ncol
# grid (row-major); 'Search settings...' can give any set of lattice
# corners, different per image (see node_editor/chessboard_search.py).

def lattice_axes(rc: np.ndarray, board: tuple) -> str | None:
    """Which way the lattice lies on the physical board, from its corner
    counts and the full board's inner corners `board` = (columns, rows):
    'x' = lattice columns c run along the board's X (its rows of
    `columns` corners), 'y' = along its Y, None = cannot tell (fits both
    ways, e.g. a board with as many columns as rows or a small part of a
    board) or fits neither (larger than the board). Only matters when the
    square X and Y sizes differ."""
    ec, er = (int(v) + 1 for v in np.ptp(np.asarray(rc).reshape(-1, 2)[:, ::-1], axis=0))
    ncol, nrow = int(board[0]), int(board[1])
    fit_x = ec <= ncol and er <= nrow
    fit_y = ec <= nrow and er <= ncol
    if fit_x and not fit_y:
        return "x"
    if fit_y and not fit_x:
        return "y"
    return None


def lattice_obj(rc: np.ndarray, sx: float, sy: float | None = None, axes: str | None = "x") -> np.ndarray:
    """Object points (c * step_c, r * step_r, 0) of lattice indices (r, c):
    step_c = square X and step_r = square Y when the lattice columns run
    along the board's X ('x'), swapped when they run along its Y ('y')."""
    sy = sx if sy is None else sy
    step_c, step_r = (sy, sx) if axes == "y" else (sx, sy)
    rc = np.asarray(rc, dtype=np.float64).reshape(-1, 2)
    return np.column_stack([rc[:, 1] * step_c, rc[:, 0] * step_r, np.zeros(len(rc))])


def stack_padded(arrays: list) -> np.ndarray:
    """(n, M, d): stacked as they are when all have M rows, else padded
    with NaN rows to the longest."""
    m = max(len(a) for a in arrays)
    if all(len(a) == m for a in arrays):
        return np.stack(arrays).astype(np.float64)
    out = np.full((len(arrays), m, arrays[0].shape[1]), np.nan)
    for k, a in enumerate(arrays):
        out[k, :len(a)] = a
    return out


def view_descriptor(rc: np.ndarray, corners: np.ndarray, half: float) -> np.ndarray:
    """(9, 2) image positions of a 3 x 3 patch of board points (+-`half`
    squares) around the centre of the found corners, through the
    lattice -> image homography. Comparable between images whose corner
    sets differ (used to rank the views)."""
    src = np.asarray(rc, np.float64)[:, ::-1].reshape(-1, 1, 2)            # (c, r)
    H, _ = cv2.findHomography(src, np.asarray(corners, np.float64).reshape(-1, 1, 2), 0)
    m = src.reshape(-1, 2).mean(axis=0)
    patch = np.array([(m[0] + dx, m[1] + dy) for dy in (-half, 0, half) for dx in (-half, 0, half)])
    if H is None:
        return np.repeat(np.asarray(corners, np.float64).reshape(-1, 2).mean(axis=0)[None], 9, axis=0)
    return cv2.perspectiveTransform(patch.reshape(-1, 1, 2), H).reshape(-1, 2)


def rank_board_cands(cands: list, min_sharp: float) -> list:
    """rank_calibration_images() for candidates; when the corner sets differ
    between images, the per-corner distance is taken on view_descriptor()
    points instead of the corners themselves."""
    first = cands[0]["rc"] if cands else None
    if all(np.array_equal(c["rc"], first) for c in cands):
        return rank_calibration_images(cands, min_sharp)
    half = max(1.0, 0.25 * min(float(np.ptp(c["rc"], axis=0).min()) for c in cands))
    proxies = [{"found": True, "sharpness": c["sharpness"], "coverage": c["coverage"],
                "center_dist": c["center_dist"], "size": c["size"], "_cand": c,
                "corners": view_descriptor(c["rc"], c["corners"], half)} for c in cands]
    ranked = []
    for q in rank_calibration_images(proxies, min_sharp):
        c = q["_cand"]
        c["quality"], c["pick_dist"] = q["quality"], q["pick_dist"]
        ranked.append(c)
    return ranked


def support_color(frac: float) -> str:
    """red (found by one setting) -> yellow -> green (found by all)."""
    f = min(1.0, max(0.0, frac))
    r, g = (255, int(80 + 175 * 2 * f)) if f < 0.5 else (int(255 * (2 - 2 * f)), 220)
    return f"#{r:02x}{g:02x}40"


def draw_lattice(c, ts, rc, pts, support, n_settings: int, used: bool, r: float = 3.5) -> None:
    """Lattice edges between neighbouring corners (r, c)-(r, c+1) / (r+1, c),
    corners coloured by how many settings found them, the lattice origin
    and its +c / +r directions."""
    rc = np.asarray(rc).reshape(-1, 2)
    idx = {tuple(k): i for i, k in enumerate(rc.tolist())}
    edge = "#d8d8e8" if used else "#8a8a8a"
    for (rr, cc), i in idx.items():
        for nb in ((rr, cc + 1), (rr + 1, cc)):
            j = idx.get(nb)
            if j is not None:
                c.create_line(*ts(*pts[i]), *ts(*pts[j]), fill=edge, width=1, tags="ov")
    for i, (x, y) in enumerate(pts):
        X, Y = ts(x, y)
        col = support_color((support[i] - 1) / max(1, n_settings - 1)) if used else "#9a9a9a"
        c.create_oval(X - r, Y - r, X + r, Y + r, outline=col, fill=col, tags="ov")
    o = idx.get((0, 0))
    if o is None:                                      # origin not found: mark the first corner of row 0
        o = min(idx.items(), key=lambda kv: (kv[0][0], kv[0][1]))[1]
    X, Y = ts(*pts[o])
    c.create_oval(X - 8, Y - 8, X + 8, Y + 8, outline="#ffffff", width=2, tags="ov")
    for nb, text in (((rc[o][0], rc[o][1] + 1), "+c (X)"), ((rc[o][0] + 1, rc[o][1]), "+r (Y)")):
        j = idx.get(tuple(int(v) for v in nb))
        if j is not None:
            X1, Y1 = ts(*pts[j])
            c.create_line(X, Y, X + 2.2 * (X1 - X), Y + 2.2 * (Y1 - Y), fill="#ffffff", width=2, arrow="last",
                          tags="ov")
            c.create_text(X + 2.6 * (X1 - X), Y + 2.6 * (Y1 - Y), text=text, fill="#ffffff",
                          font=("Arial", 8, "bold"), tags="ov")


# ── table rows ─────────────────────────────────────────────────────────────

def new_row(rid: int, name: str, X="", Y="", Z="", u=math.nan, v=math.nan, line="",
            weight=1.0, use=True) -> dict:
    return {"id": int(rid), "name": str(name), "X": str(X).strip(), "Y": str(Y).strip(),
            "Z": str(Z).strip(), "u": float(u), "v": float(v), "line": str(line).strip(),
            "weight": float(weight), "use": bool(use)}


def _num_text(v: float) -> str:
    if not math.isfinite(v):
        return ""
    return repr(round(float(v), 6))


def classify_row(r: dict) -> dict:
    """What a row contributes: kind (point / unknown / line / invalid /
    unpicked / unused), red cells, message, object point (NaN = unknown),
    unknown axis and its initial value."""
    out = {"kind": "invalid", "bad": set(), "msg": "", "obj": None, "axis": None, "init": None,
           "line": r["line"]}
    if not r["use"]:
        out["kind"] = "unused"
        return out
    cells = [r["X"], r["Y"], r["Z"]]
    filled = [c != "" for c in cells]
    obj, unk, init = [math.nan] * 3, [], None
    if all(filled):
        for a, text in enumerate(cells):
            try:
                val, is_unk, ini = ce.parse_object_cell(text)
            except ValueError:
                out["bad"].add("XYZ"[a])
                out["msg"] = f"{'XYZ'[a]} '{text}' is not a number, '?' or '?=value'"
                continue
            if is_unk:
                unk.append(a)
                init = ini
            else:
                obj[a] = val
        if out["bad"]:
            return out
        if len(unk) >= 2:
            out["bad"] = {"XYZ"[a] for a in unk}
            out["msg"] = ("two or more unknown coordinates: one image point can fix only ONE "
                          "(its viewing ray meets a 3D line, not a plane)")
            return out
        out["obj"], out["axis"], out["init"] = obj, (unk[0] if unk else None), init
    elif any(filled):
        out["bad"] = {"XYZ"[a] for a in range(3) if not filled[a]}
        out["msg"] = "fill all of X, Y, Z (use ? for an unknown), or leave all three empty for a line point"
        return out
    elif not r["line"]:
        out["msg"] = "no object coordinates and no line name: not used"
        return out
    if not (r["weight"] >= 0.0):
        out["msg"] = "weight must be 0 or more"
        return out
    if not (math.isfinite(r["u"]) and math.isfinite(r["v"])):
        out["kind"] = "unpicked"
        out["msg"] = "no image point yet: select the row and click the point in the image"
        return out
    out["kind"] = ("unknown" if out["axis"] is not None else "point") if out["obj"] is not None else "line"
    return out


def build_view(rows: list, infos: list) -> tuple:
    """The engine View for the usable rows: (view, row index per object
    point, [(line name, [row indices])], messages)."""
    obj, img, w, unk_init, pt_rows = [], [], [], {}, []
    groups: dict = {}
    issues = []
    for i, (r, info) in enumerate(zip(rows, infos)):
        if info["kind"] in ("invalid", "unpicked"):
            issues.append(f"row {i + 1} ({r['name']}): {info['msg']}")
            continue
        if info["kind"] == "unused" or r["weight"] <= 0.0:
            continue
        if info["obj"] is not None:
            if info["axis"] is not None and info["init"] is not None:
                unk_init[len(obj)] = info["init"]
            obj.append(info["obj"])
            img.append((r["u"], r["v"]))
            w.append(r["weight"])
            pt_rows.append(i)
        if r["line"]:
            groups.setdefault(r["line"], []).append(i)
    lines, line_rows, line_w = [], [], []
    for name, idx in groups.items():
        if len(idx) < 3:
            issues.append(f"line {name}: needs at least 3 image points (has {len(idx)}); not used")
            continue
        lines.append(np.array([(rows[i]["u"], rows[i]["v"]) for i in idx], dtype=np.float64))
        line_rows.append((name, idx))
        line_w.append(float(np.mean([rows[i]["weight"] for i in idx])))
    view = ce.View("image", obj=np.array(obj, dtype=np.float64).reshape(-1, 3),
                   img=np.array(img, dtype=np.float64).reshape(-1, 2), weights=np.array(w, dtype=np.float64),
                   unknown_init=unk_init, lines=lines, line_weights=line_w)
    return view, pt_rows, line_rows, issues


def write_table(path: str, rows: list) -> None:
    recs = [[r["name"], r["X"], r["Y"], r["Z"], _num_text(r["u"]), _num_text(r["v"]), r["line"],
             _num_text(r["weight"]), "1" if r["use"] else "0"] for r in rows]
    tmp = Path(path).with_name(Path(path).name + ".tmp")
    if Path(path).suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "points"
        ws.append(list(TABLE_HEADER))
        for c in ws[1]:
            c.font = Font(bold=True)
        for rec in recs:
            row = []
            for k, v in enumerate(rec):
                if k in (4, 5, 7) or (k in (1, 2, 3) and v and "?" not in v):
                    try:
                        row.append(float(v) if v != "" else None)
                        continue
                    except ValueError:
                        pass
                row.append(int(v) if k == 8 else v)
            ws.append(row)
        readme = wb.create_sheet("README")
        for line in (["Camera Calibration point table"], [""],
                     ["name", "point name (free text)"],
                     ["X, Y, Z", "object coordinates: number, ? (unknown) or ?=value (unknown with a start value)"],
                     ["u, v", "image point in pixels (OpenCV: (0, 0) = centre of the top-left pixel); empty = not picked"],
                     ["line", "name of a straight 3D line this image point lies on (X, Y, Z may then be empty)"],
                     ["weight", "importance of the row (0 = ignore)"], ["use", "1 = used, 0 = ignored"]):
            readme.append(line)
        readme.column_dimensions["A"].width = 12
        wb.save(tmp)
    else:
        with open(tmp, "w", newline="", encoding="utf-8-sig") as f:
            wr = csv.writer(f)
            wr.writerow(TABLE_HEADER)
            wr.writerows(recs)
    tmp.replace(path)


def read_table(path: str) -> list:
    """Rows (header first) of a CSV or Excel file as lists of strings."""
    if Path(path).suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["points"] if "points" in wb.sheetnames else wb.worksheets[0]
        raw = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        return raw
    with open(path, newline="", encoding="utf-8-sig") as f:
        return [list(r) for r in csv.reader(f)]


_HEADER_ALIASES = {"name": "name", "x": "X", "y": "Y", "z": "Z", "u": "u", "v": "v", "img_x": "u",
                   "img_y": "v", "image_x": "u", "image_y": "v", "line": "line", "weight": "weight",
                   "w": "weight", "use": "use"}


def _is_cell_like(text: str) -> bool:
    try:
        ce.parse_object_cell(text)
        return True
    except ValueError:
        return text.strip() == ""


def parse_records(raw: list, next_id, next_name) -> tuple[list, list]:
    """Records (lists of strings) -> (rows, errors). With a header line the
    columns are found by name; without one see HELP_TEXT ('Pasting')."""
    raw = [list(r) for r in raw if any(str(v).strip() for v in r)]
    if not raw:
        return [], ["nothing to read"]
    head = [_HEADER_ALIASES.get(str(h).strip().lower()) for h in raw[0]]
    errors, rows = [], []
    header_cols = None
    if sum(h is not None for h in head) >= 2:
        header_cols = {h: k for k, h in enumerate(head) if h is not None}
        body, first_line = raw[1:], 2
    else:
        body, first_line = raw, 1
    for ln, rec in enumerate(body, start=first_line):
        rec = [str(v).strip() for v in rec]
        while rec and rec[-1] == "":
            rec.pop()
        cols = header_cols
        if cols is None:                         # no header: decide per line from its number of cells
            has_name = bool(rec) and not _is_cell_like(rec[0])
            keys = {2: ["u", "v"], 3: ["X", "Y", "Z"], 5: ["X", "Y", "Z", "u", "v"]}.get(
                len(rec) - has_name, ["X", "Y", "Z", "u", "v", "line"])
            cols = {k: j + has_name for j, k in enumerate(keys)}
            if has_name:
                cols["name"] = 0
        rec = rec + [""] * 12

        def get(key, default=""):
            return rec[cols[key]] if key in cols else default
        try:
            u = float(get("u")) if get("u") else math.nan
            v = float(get("v")) if get("v") else math.nan
            w = float(get("weight")) if get("weight") else 1.0
        except ValueError:
            errors.append(f"line {ln}: u, v and weight must be numbers")
            continue
        use = get("use", "1").lower() not in ("0", "false", "no", "n")
        rows.append(new_row(next_id(), get("name") or next_name(), get("X"), get("Y"), get("Z"),
                            u, v, get("line"), w, use))
    return rows, errors


# ── image view ─────────────────────────────────────────────────────────────

class ImageView:
    """Zoomable image with overlays. Screen x of image column u is
    ox + (u + 0.5) * s, so integer (u, v) is a pixel centre (OpenCV).
    Mouse, as in the Image Display node: wheel = zoom at the cursor;
    left-drag on empty space = pan; left-drag on a point = move it; a left
    click without moving = on_click."""

    PREVIEW_MAX = 1600

    def __init__(self, parent, on_click, on_drag_point, on_motion):
        self.on_click, self.on_drag_point, self.on_motion = on_click, on_drag_point, on_motion
        self.canvas = tk.Canvas(parent, bg="#1b1b24", highlightthickness=0, cursor="crosshair")
        self.image = self.preview = None
        self.k = 1.0
        self.view = None
        self.user_view = False
        self.serial = 0
        self._photo, self._photo_key = None, None
        self._press = None
        self._drag = None                              # [row index, moved] while a point is grabbed
        self._pan0 = None                              # (x, y, ox, oy) while panning
        self.draw_overlay = lambda _v: None            # set by the node
        self.hit_point = lambda _x, _y: None           # row index under the cursor, set by the node
        self.empty_text = "No image."
        c = self.canvas
        c.bind("<Configure>", lambda _e: self.render())
        for seq in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            c.bind(seq, self._on_wheel)
        c.bind("<ButtonPress-1>", self._on_press)
        c.bind("<B1-Motion>", self._on_b1_motion)
        c.bind("<ButtonRelease-1>", self._on_release)
        c.bind("<Motion>", self._on_move)

    def set_image(self, img) -> None:
        if img is self.image:
            return
        if img is not None:
            a = np.asarray(img)
            if a.dtype != np.uint8:
                a = cv2.normalize(a.astype(np.float64), None, 0, 255, cv2.NORM_MINMAX).astype(np.uint8)
            if a.ndim == 2:
                a = np.repeat(a[:, :, None], 3, axis=2)
            elif a.ndim == 3 and a.shape[2] == 4:
                a = a[:, :, :3]
            old = None if self.image is None else self.image.shape[:2]
            self.image = np.ascontiguousarray(a)
            h, w = self.image.shape[:2]
            f = min(1.0, self.PREVIEW_MAX / max(h, w))
            self.preview = self.image if f >= 1.0 else cv2.resize(
                self.image, (max(1, round(w * f)), max(1, round(h * f))), interpolation=cv2.INTER_AREA)
            self.k = self.preview.shape[1] / w
            if old != (h, w):
                self.user_view = False
        else:
            self.image = self.preview = None
        self.serial += 1
        self.render()

    def _fit(self, cw, ch, size):
        w, h = size
        s = min((cw - 8) / w, (ch - 8) / h)
        return ((cw - w * s) / 2.0, (ch - h * s) / 2.0, s)

    def fit(self) -> None:
        self.user_view = False
        self.render()

    def to_screen(self, u, v):
        ox, oy, s = self.view
        return ox + (u + 0.5) * s, oy + (v + 0.5) * s

    def to_image(self, x, y):
        ox, oy, s = self.view
        return (x - ox) / s - 0.5, (y - oy) / s - 0.5

    def ensure_visible(self, u: float, v: float) -> None:
        if self.view is None:
            return
        cw, ch = self.canvas.winfo_width(), self.canvas.winfo_height()
        x, y = self.to_screen(u, v)
        if 20 <= x <= cw - 20 and 20 <= y <= ch - 20:
            return
        s = self.view[2]
        self.view = (cw / 2.0 - (u + 0.5) * s, ch / 2.0 - (v + 0.5) * s, s)
        self.user_view = True
        self.render()

    def render(self, size=None) -> None:
        """Draw the image (or, without one, a frame of `size` = (w, h)) and the overlay."""
        c = self.canvas
        try:
            cw, ch = c.winfo_width(), c.winfo_height()
        except tk.TclError:
            return
        if cw < 20 or ch < 20:
            return
        if size is not None:
            self._size = size
        size = (self.image.shape[1], self.image.shape[0]) if self.image is not None else getattr(self, "_size", None)
        c.delete("ov")
        if size is None:
            c.delete("img")
            self._photo_key = None
            self.view = None
            c.create_text(cw / 2, ch / 2, fill="#9aa0b5", justify="center", text=self.empty_text, tags="ov")
            return
        if not self.user_view or self.view is None:
            self.view = self._fit(cw, ch, size)
        ox, oy, s = self.view
        if self.image is None:
            c.delete("img")
            self._photo_key = None
            x0, y0 = self.to_screen(-0.5, -0.5)
            x1, y1 = self.to_screen(size[0] - 0.5, size[1] - 0.5)
            c.create_rectangle(x0, y0, x1, y1, fill="#2a2a36", outline="#555566", tags="ov")
            c.create_text(cw / 2, 14, fill="#9aa0b5", text=self.empty_text, tags="ov")
        else:
            src, pk = (self.preview, self.k) if s <= self.k else (self.image, 1.0)
            sh, sw = src.shape[:2]
            c0 = max(0, int(math.floor(-ox / s * pk)))
            c1 = min(sw, int(math.ceil((cw - ox) / s * pk)))
            r0 = max(0, int(math.floor(-oy / s * pk)))
            r1 = min(sh, int(math.ceil((ch - oy) / s * pk)))
            if c1 > c0 and r1 > r0:
                px, py = ox + c0 / pk * s, oy + r0 / pk * s
                dw = max(1, int(round(ox + c1 / pk * s - px)))
                dh = max(1, int(round(oy + r1 / pk * s - py)))
                key = (self.serial, c0, r0, c1, r1, dw, dh)
                if key != self._photo_key:
                    crop = np.ascontiguousarray(src[r0:r1, c0:c1])
                    interp = cv2.INTER_AREA if dw < crop.shape[1] else (
                        cv2.INTER_NEAREST if s >= 3.0 else cv2.INTER_LINEAR)
                    self._photo = ImageTk.PhotoImage(Image.fromarray(cv2.resize(crop, (dw, dh), interpolation=interp)))
                    self._photo_key = key
                    c.delete("img")
                    c.create_image(px, py, anchor="nw", image=self._photo, tags="img")
                else:
                    c.coords("img", px, py)
            else:
                c.delete("img")
                self._photo_key = None
        self.draw_overlay(self)

    @staticmethod
    def wheel_steps(e) -> int:
        if getattr(e, "num", None) == 4:
            return 1
        if getattr(e, "num", None) == 5:
            return -1
        return int(e.delta / 120) if getattr(e, "delta", 0) else 0

    def _on_wheel(self, e) -> str:
        if self.view is None:
            return "break"
        steps = self.wheel_steps(e)
        if steps:
            ox, oy, s = self.view
            ns = min(max(s * (1.25 ** steps), 0.01), 64.0)
            f = ns / s
            self.view = (e.x - (e.x - ox) * f, e.y - (e.y - oy) * f, ns)
            self.user_view = True
            self.render()
        return "break"

    CLICK_PX = 3                                       # moving less than this is still a click

    def _on_press(self, e) -> None:
        self.canvas.focus_set()
        self._pan0 = None
        if self.view is None:
            return
        self._press = (e.x, e.y)
        hit = self.hit_point(e.x, e.y)
        self._drag = None if hit is None else [hit, False]

    def _on_b1_motion(self, e) -> None:
        p = self._press
        if p is None or self.view is None:
            return
        if self._pan0 is None and (self._drag is None or not self._drag[1]) \
                and abs(e.x - p[0]) + abs(e.y - p[1]) <= self.CLICK_PX:
            return
        if self._drag is not None:                     # move the grabbed point
            self._drag[1] = True
            u, v = self.to_image(e.x, e.y)
            self.on_drag_point(self._drag[0], u, v, False)
            return
        if self._pan0 is None:                         # pan: keep the grabbed image spot under the mouse
            self._pan0 = (p[0], p[1], self.view[0], self.view[1])
            self.canvas.configure(cursor="fleur")
        x0, y0, ox, oy = self._pan0
        self.view = (ox + e.x - x0, oy + e.y - y0, self.view[2])
        self.user_view = True
        self.render()

    def _on_release(self, e) -> None:
        p, self._press = self._press, None
        drag, self._drag = self._drag, None
        panned, self._pan0 = self._pan0 is not None, None
        if panned:
            self.canvas.configure(cursor="crosshair")
        if p is None or self.view is None or panned:
            return
        u, v = self.to_image(e.x, e.y)
        if drag is not None and drag[1]:
            self.on_drag_point(drag[0], u, v, True)
            return
        if abs(e.x - p[0]) + abs(e.y - p[1]) > self.CLICK_PX:
            return
        self.on_click(u, v, drag[0] if drag is not None else None)

    def _on_move(self, e) -> None:
        if self.view is not None:
            self.on_motion(*self.to_image(e.x, e.y))


# ── 3D preview ─────────────────────────────────────────────────────────────

_CUBE_V = np.array([[x, y, z] for x in (-1, 1) for y in (-1, 1) for z in (-1, 1)], dtype=float)
_CUBE_F = ((0, 1, 3, 2), (4, 6, 7, 5), (0, 4, 5, 1), (2, 3, 7, 6), (0, 2, 6, 4), (1, 5, 7, 3))


def view_rotation(yaw: float, pitch: float) -> np.ndarray:
    """World -> view rotation as in the other 3D previews: view x = right,
    view y = depth (away from the eye), view z = up on the screen."""
    cy, sy, cp, sp = math.cos(yaw), math.sin(yaw), math.cos(pitch), math.sin(pitch)
    return np.array([[1, 0, 0], [0, cp, -sp], [0, sp, cp]]) @ np.array([[cy, -sy, 0], [sy, cy, 0], [0, 0, 1]])


class View3D:
    """Rotatable 3D preview, same mouse conventions as the other 3D views:
    left-drag = rotate, right-drag = pan, wheel = zoom; perspective 0..1;
    lower-left X / Y / Z axes icon with a small cube. A scene is a list of
    ("line", (N, 3) points, colour, width), ("dot", (N, 3), colour, radius)
    and ("text", (3,) point, text, colour)."""

    def __init__(self, parent, persp_var):
        self.canvas = tk.Canvas(parent, bg="#ffffff", highlightthickness=0, cursor="fleur")
        self.persp_var = persp_var
        self.yaw, self.pitch, self.zoom = -0.9, 0.55, 1.0
        self.target, self.radius = np.zeros(3), 1.0
        self.scene, self.fitted, self.scale = [], False, 1.0
        self.empty_text = ""
        self._drag = None
        c = self.canvas
        c.bind("<Configure>", lambda _e: self.draw())
        c.bind("<ButtonPress-1>", lambda e: self._press(e, 1))
        c.bind("<ButtonPress-3>", lambda e: self._press(e, 3))
        c.bind("<B1-Motion>", self._move)
        c.bind("<B3-Motion>", self._move)
        for seq in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            c.bind(seq, self._wheel)

    def set_scene(self, scene: list, refit: bool = False) -> None:
        self.scene = scene
        if refit or not self.fitted:
            self.fit(redraw=False)
        self.draw()

    def fit(self, redraw: bool = True) -> None:
        pts = [np.asarray(p[1], float).reshape(-1, 3) for p in self.scene if p[0] in ("line", "dot")]
        pts = np.vstack(pts) if pts else np.zeros((1, 3))
        pts = pts[np.all(np.isfinite(pts), axis=1)]
        if len(pts):
            lo, hi = pts.min(axis=0), pts.max(axis=0)
            self.target = (lo + hi) / 2.0
            self.radius = max(float(np.linalg.norm(hi - lo)) / 2.0, 1e-9)
        self.zoom, self.fitted = 1.0, True
        if redraw:
            self.draw()

    def reset(self) -> None:
        self.yaw, self.pitch = -0.9, 0.55
        self.fit()

    def _persp(self) -> float:
        try:
            return min(1.0, max(0.0, float(self.persp_var.get())))
        except (tk.TclError, ValueError):
            return 0.3

    def project(self, P) -> np.ndarray:
        c = self.canvas
        w, h = c.winfo_width(), c.winfo_height()
        rot = (np.asarray(P, float).reshape(-1, 3) - self.target) @ view_rotation(self.yaw, self.pitch).T
        p = self._persp()
        if p > 1e-6:
            d0 = 3.0 * self.radius / p
            f = d0 / np.maximum(d0 + rot[:, 1], d0 * 1e-3)
        else:
            f = np.ones(len(rot))
        self.scale = min(w, h) * 0.42 * self.zoom / self.radius
        return np.column_stack([w / 2 + self.scale * rot[:, 0] * f, h / 2 - self.scale * rot[:, 2] * f])

    def draw(self) -> None:
        c = self.canvas
        try:
            w, h = c.winfo_width(), c.winfo_height()
        except tk.TclError:
            return
        if w < 30 or h < 30:
            return
        c.delete("all")
        if not self.scene:
            c.create_text(w / 2, h / 2, text=self.empty_text, fill="#888888", width=w - 40, justify="center")
            return
        for prim in self.scene:
            kind = prim[0]
            if kind == "line":
                q = self.project(prim[1])
                if len(q) >= 2 and np.all(np.isfinite(q)):
                    c.create_line(*q.ravel().tolist(), fill=prim[2], width=prim[3])
            elif kind == "dot":
                r = prim[3]
                for x, y in self.project(prim[1]):
                    if math.isfinite(x) and math.isfinite(y):
                        c.create_oval(x - r, y - r, x + r, y + r, fill=prim[2], outline="")
            elif kind == "text":
                x, y = self.project(prim[1])[0]
                if math.isfinite(x) and math.isfinite(y):
                    c.create_text(x + 4, y - 4, text=prim[2], fill=prim[3], anchor="sw", font=("Arial", 7))
        self._axes_icon(h)

    def _axes_icon(self, ch: float) -> None:
        c = self.canvas
        R = view_rotation(self.yaw, self.pitch)
        corner = np.array([32.0, float(ch) - 32.0])
        axis_len = 22.0
        for ax_v, col, lab in ((np.array([1.0, 0, 0]), "#cc0000", "X"), (np.array([0, 1.0, 0]), "#00aa00", "Y"),
                               (np.array([0, 0, 1.0]), "#0000cc", "Z")):
            rv = R @ ax_v
            e = corner + np.array([rv[0], -rv[2]]) * axis_len
            c.create_line(corner[0], corner[1], e[0], e[1], fill=col, width=2)
            c.create_text(e[0] + (e[0] - corner[0]) * 0.25, e[1] + (e[1] - corner[1]) * 0.25,
                          text=lab, fill=col, font=("Arial", 8, "bold"))
        colors = ("#ffcccc", "#ff9999", "#cceecc", "#99dd99", "#ccccff", "#9999ff")
        verts = (_CUBE_V * (axis_len * 0.32)) @ R.T
        for f, quad in enumerate(_CUBE_F):
            q = verts[list(quad)]
            if float(np.cross(q[1] - q[0], q[2] - q[0])[1]) < 0.0:
                c.create_polygon([(corner[0] + verts[i][0], corner[1] - verts[i][2]) for i in quad],
                                 fill=colors[f], outline="#333333", width=1.3)

    def _press(self, e, button: int) -> None:
        self._drag = (button, e.x, e.y)

    def _move(self, e) -> None:
        if self._drag is None:
            return
        button, x0, y0 = self._drag
        dx, dy = e.x - x0, e.y - y0
        self._drag = (button, e.x, e.y)
        if button == 1:
            self.yaw += dx * 0.01
            self.pitch += dy * 0.01
        elif self.scale > 1e-12:
            rt = view_rotation(self.yaw, self.pitch).T
            self.target = self.target - (rt[:, 0] * dx - rt[:, 2] * dy) / self.scale
        self.draw()

    def _wheel(self, e) -> str:
        steps = ImageView.wheel_steps(e)
        if steps:
            self.zoom = max(0.05, self.zoom * (1.12 ** steps))
            self.draw()
        return "break"


def frustum_lines(R, t, K, size, depth: float, color: str, width: int = 1, M=None) -> list:
    """Scene lines of a camera (world -> camera: R, t): rays to the image
    corners at `depth`, the image rectangle and a small 'up' triangle on its
    top edge. M: optional display transform applied to everything."""
    R, t = np.asarray(R, float), np.asarray(t, float).reshape(3)
    C = -R.T @ t
    w, h = size
    fx, fy, cx, cy = K[0][0], K[1][1], K[0][2], K[1][2]

    def ray(u, v):
        return C + R.T @ (np.array([(u - cx) / fx, (v - cy) / fy, 1.0]) * depth)
    corners = [ray(0, 0), ray(w, 0), ray(w, h), ray(0, h)]
    up = [ray(0.35 * w, 0), ray(0.5 * w, -0.18 * h), ray(0.65 * w, 0)]
    T = (lambda P: np.asarray(P) @ M.T) if M is not None else (lambda P: np.asarray(P))
    out = [("line", T([C, q]), color, width) for q in corners]
    out.append(("line", T(corners + [corners[0]]), color, width))
    out.append(("line", T(up), color, width))
    return out


def axes_lines(origin, length: float, M=None, R=None) -> list:
    """X / Y / Z axes (red / green / blue) of a frame at `origin` (columns of R)."""
    R = np.eye(3) if R is None else np.asarray(R, float)
    T = (lambda P: np.asarray(P) @ M.T) if M is not None else (lambda P: np.asarray(P))
    o = np.asarray(origin, float)
    return [("line", T([o, o + R[:, k] * length]), col, 3)
            for k, col in enumerate(("#d62728", "#2ca02c", "#1f77b4"))]


# ── point table ────────────────────────────────────────────────────────────

class Sheet:
    """Spreadsheet-like table drawn on a canvas (only visible rows are drawn).
    Click selects (Ctrl / Shift: several), double-click edits a cell,
    clicking the 'use' cell toggles it."""

    ROW_H = 20
    HEAD_H = 22
    COLUMNS = (("#", "#", 34), ("use", "use", 34), ("name", "name", 90), ("X", "X", 66), ("Y", "Y", 66),
               ("Z", "Z", 66), ("u", "u (px)", 70), ("v", "v (px)", 70), ("line", "line", 46),
               ("weight", "weight", 50), ("err", "error px", 70))
    EDITABLE = ("name", "X", "Y", "Z", "u", "v", "line", "weight")

    def __init__(self, parent, node):
        self.node = node
        self.frame = tk.Frame(parent)
        self.canvas = tk.Canvas(self.frame, bg="#ffffff", highlightthickness=1, highlightbackground="#c8cdd8",
                                yscrollincrement=self.ROW_H, height=260)
        vsb = ttk.Scrollbar(self.frame, orient="vertical", command=self._yview)
        hsb = ttk.Scrollbar(self.frame, orient="horizontal", command=self._xview)
        self.canvas.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        vsb.pack(side="right", fill="y")
        hsb.pack(side="bottom", fill="x")
        self.canvas.pack(fill="both", expand=True)
        self.selected: list[int] = []
        self._anchor = None
        self._editor = None
        c = self.canvas
        c.bind("<Configure>", lambda _e: self.redraw())
        c.bind("<ButtonPress-1>", self._on_click)
        c.bind("<Double-Button-1>", self._on_double)
        for seq in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            c.bind(seq, self._on_wheel)
        c.bind("<Delete>", lambda _e: node.delete_selected())
        c.bind("<Up>", lambda _e: self.step(-1))
        c.bind("<Down>", lambda _e: self.step(1))
        c.bind("<Control-v>", lambda _e: node.paste_rows())
        c.bind("<Control-V>", lambda _e: node.paste_rows())

    def _yview(self, *args):
        self.canvas.yview(*args)
        self.redraw()

    def _xview(self, *args):
        self.canvas.xview(*args)
        self.redraw()

    def _on_wheel(self, e) -> str:
        self.canvas.yview_scroll(-3 * ImageView.wheel_steps(e), "units")
        self.redraw()
        return "break"

    def _col_at(self, x):
        left = 0
        for i, (_k, _t, w) in enumerate(self.COLUMNS):
            if left <= x < left + w:
                return i
            left += w
        return None

    def _row_at(self, y):
        i = int((y - self.HEAD_H) // self.ROW_H)
        return i if 0 <= i < len(self.node.rows) and y >= self.HEAD_H else None

    def redraw(self) -> None:
        c = self.canvas
        try:
            ch = c.winfo_height()
        except tk.TclError:
            return
        rows = self.node.rows
        total_w = sum(w for _k, _t, w in self.COLUMNS)
        c.configure(scrollregion=(0, 0, total_w, self.HEAD_H + len(rows) * self.ROW_H))
        c.delete("cell")
        top = c.canvasy(0)
        first = max(0, int((top - self.HEAD_H) // self.ROW_H))
        last = min(len(rows), int((top + ch) // self.ROW_H) + 1)
        sel = set(self.selected)
        for i in range(first, last):
            y0 = self.HEAD_H + i * self.ROW_H
            left = 0
            for key, _title, w in self.COLUMNS:
                text, fg, bg = self.node.cell(i, key)
                if i in sel and key in ("#", "name"):
                    bg = "#fff3a8"
                c.create_rectangle(left, y0, left + w, y0 + self.ROW_H, fill=bg, outline="#d4d4d4", tags="cell")
                c.create_text(left + 4, y0 + self.ROW_H / 2, anchor="w", text=text, fill=fg,
                              font=("Arial", 8, "bold" if (i in sel and key == "name") else "normal"), tags="cell")
                left += w
            if i in sel:
                c.create_rectangle(1, y0 + 1, total_w - 1, y0 + self.ROW_H - 1, outline="#e0a800", width=2, tags="cell")
        left = 0
        for _key, title, w in self.COLUMNS:
            c.create_rectangle(left, top, left + w, top + self.HEAD_H, fill="#e3e8f0", outline="#a8b2c4", tags="cell")
            c.create_text(left + 4, top + self.HEAD_H / 2, anchor="w", text=title, font=("Arial", 8, "bold"), tags="cell")
            left += w
        if not rows:
            c.create_text(10, top + self.HEAD_H + 26, anchor="nw", fill="#777777", font=("Arial", 9), width=total_w - 20,
                          text="No points yet. 'Add row' and type X, Y, Z (or paste rows from Excel with "
                               "Ctrl+V), then click each point in the image.", tags="cell")

    def ensure_visible(self, i: int) -> None:
        c = self.canvas
        top, h = c.canvasy(0), c.winfo_height()
        y0 = self.HEAD_H + i * self.ROW_H
        total = self.HEAD_H + max(1, len(self.node.rows)) * self.ROW_H
        if y0 < top + self.HEAD_H or y0 + self.ROW_H > top + h:
            c.yview_moveto(max(0.0, (y0 - self.HEAD_H - h / 3) / total))
        self.redraw()

    def select(self, indices: list) -> None:
        self.selected = sorted(set(i for i in indices if 0 <= i < len(self.node.rows)))
        self._anchor = self.selected[0] if self.selected else None
        if self.selected:
            self.ensure_visible(self.selected[0])
        else:
            self.redraw()

    def step(self, d: int) -> str:
        n = len(self.node.rows)
        if n:
            i = (self.selected[-1] + d) if self.selected else 0
            self.select([min(max(i, 0), n - 1)])
            self.node.on_selection_changed()
        return "break"

    def _on_click(self, e) -> None:
        self.canvas.focus_set()
        self.close_editor(commit=True)
        x, y = self.canvas.canvasx(e.x), self.canvas.canvasy(e.y)
        if y < self.canvas.canvasy(0) + self.HEAD_H:
            return
        i, col = self._row_at(y), self._col_at(x)
        if i is not None and col is not None and self.COLUMNS[col][0] == "use":
            self.node.toggle_use(i)
            return
        if i is None:
            self.selected = []
        elif e.state & 0x4:
            self.selected = [j for j in self.selected if j != i] if i in self.selected else sorted(self.selected + [i])
            self._anchor = i
        elif e.state & 0x1 and self._anchor is not None:
            a, b = sorted((self._anchor, i))
            self.selected = list(range(a, b + 1))
        else:
            self.selected, self._anchor = [i], i
        self.redraw()
        self.node.on_selection_changed()

    def _on_double(self, e) -> None:
        x, y = self.canvas.canvasx(e.x), self.canvas.canvasy(e.y)
        i, col = self._row_at(y), self._col_at(x)
        if i is None or col is None:
            return
        key, _t, w = self.COLUMNS[col]
        if key not in self.EDITABLE:
            return
        left = sum(cw for _k, _tt, cw in self.COLUMNS[:col])
        y0 = self.HEAD_H + i * self.ROW_H
        r = self.node.rows[i]
        widget = tk.Entry(self.canvas, font=("Arial", 8))
        value = r[key]
        widget.insert(0, _num_text(value) if isinstance(value, float) else str(value))
        widget.select_range(0, tk.END)
        widget.bind("<Return>", lambda _e: self.close_editor(commit=True, move=1))
        widget.bind("<Tab>", lambda _e: self.close_editor(commit=True, move=0, next_col=True) or "break")
        # FocusOut may arrive after this editor was replaced by the next one: close only itself.
        widget.bind("<FocusOut>", lambda _e, wd=widget: (self._editor is not None and self._editor[0] is wd
                                                         and self.close_editor(commit=True)))
        widget.bind("<Escape>", lambda _e: self.close_editor(commit=False))
        win = self.canvas.create_window(left, y0, anchor="nw", window=widget, width=w, height=self.ROW_H,
                                        tags="editor")
        self._editor = (widget, win, i, key)
        widget.focus_set()

    def close_editor(self, commit: bool, move: int = 0, next_col: bool = False) -> None:
        ed, self._editor = self._editor, None
        if ed is None:
            return
        widget, win, i, key = ed
        value = widget.get()
        try:
            self.canvas.delete(win)
            widget.destroy()
        except tk.TclError:
            pass
        if commit:
            self.node.edit_cell(i, key, value)
        self.redraw()
        # Enter: edit the same column of the next row; Tab: the next column.
        if commit and (move or next_col):
            keys = [k for k, _t, _w in self.COLUMNS]
            j = i + move
            col = keys.index(key)
            if next_col:
                nxt = [k for k in keys[col + 1:] if k in self.EDITABLE]
                if not nxt:
                    return
                key = nxt[0]
            if 0 <= j < len(self.node.rows):
                self.selected = [j]
                self.ensure_visible(j)
                left = sum(cw for _k, _tt, cw in self.COLUMNS[:keys.index(key)])
                y = self.HEAD_H + j * self.ROW_H + self.ROW_H / 2 - self.canvas.canvasy(0)
                ev = type("E", (), {"x": left + 3 - self.canvas.canvasx(0), "y": y})()
                self._on_double(ev)


# ── node ───────────────────────────────────────────────────────────────────

class CameraCalibrationNode(BaseNode):
    """Camera calibration from object / image points with unknown ('?')
    coordinates and straight-line points, robust weighted least squares."""

    EXECUTION_MODE = ExecutionMode.SYNC
    NODE_TYPE = "camera_calibration"
    # Projects saved when this node was 'Camera Calibration (Unified)' still load.
    NODE_TYPE_ALIASES = ("camera_calibration_unified",)
    DISPLAY_NAME = "Camera Calibration"
    CATEGORY = "process"
    SEARCH_KEYWORDS = ("calibration", "camera", "intrinsics", "distortion", "pose", "solvePnP",
                       "plumb line", "straight line", "unknown", "camera matrix", "unified")
    NODE_WIDTH = 230
    NODE_HEIGHT = 120
    HELP_TEXT = HELP_TEXT

    def __init__(self, node_id: str, canvas: tk.Canvas):
        super().__init__(node_id, canvas)
        self.rows: list = []
        self._next_id = 1
        self._infos: list = []
        # result: dict K, D, points_pose [rvec, tvec] | None, cb_poses {path: [rvec, tvec]},
        # unknowns {row id: (axis, value)}, free, stale, report
        self._res = None
        self._eval = None                # engine Result: the stored result evaluated on the current data
        self._eval_map = None            # dict pt_rows, line_rows, pi (points view index | None), cb_paths
        self._cb_err: dict = {}          # chessboard path -> RMS px (current evaluation)
        self._row_err: dict = {}         # row id -> (text, colour)
        self._eval_unknowns: dict = {}   # row index -> solved '?' value (current evaluation)
        self._issues: list = []
        # inputs
        self._pin_image = None
        self._file_image = None
        self._image_path = ""
        self._pin_obj = self._pin_img = None
        self._pin_K = self._pin_D = None
        self._gray = None
        # chessboard
        self._cb_pin_files = None        # list of paths from the 'files' pin
        self._cb_manual_files: list = [] # picked with 'Add images...' (absolute paths)
        self._cb = None                  # detection: {"cfg": {...}, "cands": [dict, rank order]}
        self._cb_use: set = set()        # paths of the checked board images
        self._world = "points"           # "points" or a board image path
        self._cb_state = None            # detection thread state
        self._cb_text = ""
        self._show = None                # image view: None = points image, else a board path
        self._board_img = (None, None)   # (path, rgb) cache
        self._tree_paths: list = []      # chessboard table row -> path
        self._adj_rows: list = []        # Adjust tab: [{param, lo, hi}]
        self._adj_widgets: list = []
        self._grids: list = []           # Grids tab: [{on, name, xs, ys, zs, frame, color}]
        self._grid_cache: dict = {}
        self._grids_on_var = tk.BooleanVar(value=True)
        self._grid_width_var = tk.StringVar(value="1")
        self._v3_mode_var = tk.StringVar(value="points image (object frame)")
        self._v3_persp_var = tk.DoubleVar(value=0.3)
        self._v3_cam_var = tk.DoubleVar(value=0.3)
        self._und_alpha_var = tk.DoubleVar(value=0.0)
        self._und_pts_var = tk.BooleanVar(value=True)
        self._und_cache: dict = {}
        self._und_text = ""
        self._loo_state = None
        self._loo_rows: list = []
        self._loo_base = None
        self._loo_text = ""
        self._loo_sort = ("note", True)            # 'note' column = sort by influence
        self._loo_boards_var = tk.BooleanVar(value=True)
        self._loo_points_var = tk.BooleanVar(value=True)
        self._runs: list = []                     # Runs tab
        self._det_vars = {"ncol": tk.StringVar(value="9"), "nrow": tk.StringVar(value="6"),
                          "square": tk.StringVar(value="25"), "square_y": tk.StringVar(value=""),
                          "max_width": tk.StringVar(value="1200"),
                          "threads": tk.StringVar(value=str(max(2, min(8, (os.cpu_count() or 4))))),
                          "adaptive": tk.BooleanVar(value=True), "normalize": tk.BooleanVar(value=True),
                          "fast": tk.BooleanVar(value=True)}
        self._sel_vars = {"target": tk.StringVar(value="20"), "min_sharp": tk.StringVar(value="0"),
                          "diversity": tk.StringVar(value="0.03")}
        self._use_points_var = tk.BooleanVar(value=True)
        self._use_cb_var = tk.BooleanVar(value=True)
        self._combine_var = tk.StringVar(value="joint")
        self._det_vars["square"].trace_add("write", lambda *_a: self._on_square_changed())
        self._det_vars["square_y"].trace_add("write", lambda *_a: self._on_square_changed())
        # 'Search settings...' (brute-force detection search): the ranges to try and the last run.
        self._search_cfg = {"cols": "8 9 10", "rows": "5 6 7", "widths": "800 1200 0",
                            "adaptive": "both", "normalize": "both", "fast": "on"}
        self._search: dict | None = None
        self._search_win = None
        self._force_lsq_var = tk.BooleanVar(value=False)       # Solve tab: 'Always refine with least squares'
        self._route_after = None
        # settings
        self._model_var = tk.StringVar(value="simple")
        self._fix_vars = {n: tk.BooleanVar(value=n in DEFAULT_FIXED) for n in ce.INTR_NAMES + ce.DIST_NAMES}
        self._init_vars = {n: tk.StringVar(value="") for n in ce.INTR_NAMES + ce.DIST_NAMES}
        self._aspect_var = tk.BooleanVar(value=False)
        self._w_vars = {t: tk.StringVar(value="1") for t in ("point", "unknown", "line", "chessboard")}
        self._loss_var = tk.StringVar(value="linear")
        self._huber_var = tk.StringVar(value="2")
        self._size_vars = (tk.StringVar(value=""), tk.StringVar(value=""))
        self._follow_var = tk.BooleanVar(value=False)
        # view options
        self._mode_var = tk.StringVar(value="point")
        self._line_var = tk.StringVar(value="L1")
        self._snap_var = tk.BooleanVar(value=False)
        self._show_names_var = tk.BooleanVar(value=True)
        self._show_proj_var = tk.BooleanVar(value=True)
        self._show_lines_var = tk.BooleanVar(value=True)
        self._show_axes_var = tk.BooleanVar(value=True)
        self._gain_var = tk.StringVar(value="10")
        self._face_var = tk.StringVar(value="no result yet\n(double-click to open)")
        for var in ([self._force_lsq_var, self._loss_var, self._huber_var, self._aspect_var, self._use_points_var,
                     self._use_cb_var, self._combine_var, self._model_var]
                    + list(self._w_vars.values()) + list(self._init_vars.values()) + list(self._fix_vars.values())):
            var.trace_add("write", lambda *_a: self._schedule_route_text())
        self._reset_widget_refs()
        self._solve_state = None

    def _reset_widget_refs(self) -> None:
        self._route_var = self._route_lbl = self._cancel_solve_btn = None
        self._sheet = None
        self._iv = None
        self._issues_var = None
        self._summary_var = None
        self._summary_lbl = None
        self._status_var = None
        self._solve_btn = None
        self._param_frame = None
        self._res_vars = {}
        self._result_text = None
        self._mouse_var = None
        self._img_src_var = None
        self._help_popup = None
        self._line_combo = None
        self._cb_tree = None
        self._cb_text_var = None
        self._cb_progress_var = None
        self._cb_detect_btn = None
        self._cb_files_var = None
        self._world_var = None
        self._show_var = None
        self._adj_frame = None
        self._adj_info_var = None
        self._grid_frame = None
        self._grid_msg_var = None
        self._adj_widgets = []
        self._rnb = None
        self._err_canvas = None
        self._err_hover_var = None
        self._v3 = None
        self._und_iv = None
        self._und_info_var = None
        self._und_state = None
        self._loo_tree = None
        self._loo_text_var = None
        self._loo_prog_var = None
        self._loo_btn = None
        self._runs_tree = None

    # ── pins ───────────────────────────────────────────────────────────────
    def get_pin_schema(self) -> PinSchema:
        return PinSchema(
            inputs=[PinDef("image", PinType.IMAGE, "image", optional=True),
                    PinDef("object_points", PinType.ARRAY, "objPts", optional=True),
                    PinDef("image_points", PinType.ARRAY, "imgPts", optional=True),
                    PinDef("cmat_in", PinType.ARRAY, "cmat in", optional=True),
                    PinDef("dvec_in", PinType.ARRAY, "dvec in", optional=True),
                    PinDef("files", PinType.STRING, "files", optional=True)],
            outputs=[PinDef("cmat", PinType.ARRAY, "cmat"),
                     PinDef("dvec", PinType.ARRAY, "dvec"),
                     PinDef("rvec", PinType.ARRAY, "rvec"),
                     PinDef("tvec", PinType.ARRAY, "tvec"),
                     PinDef("rms", PinType.SCALAR, "RMS"),
                     PinDef("rms_types", PinType.ARRAY, "RMS types"),
                     PinDef("img_size", PinType.ARRAY, "imgSize"),
                     PinDef("obj_pts", PinType.ARRAY, "objPts"),
                     PinDef("img_pts", PinType.ARRAY, "imgPts"),
                     PinDef("proj_pts", PinType.ARRAY, "prjPts"),
                     PinDef("pt_err", PinType.ARRAY, "ptErr"),
                     PinDef("rvecs", PinType.ARRAY, "rvecs"),
                     PinDef("tvecs", PinType.ARRAY, "tvecs"),
                     PinDef("cb_err", PinType.ARRAY, "cbErr"),
                     PinDef("cb_obj_pts", PinType.ARRAY, "cbObjPts"),
                     PinDef("cb_img_pts", PinType.ARRAY, "cbImgPts")])

    def _world_pose(self):
        """(rvec, tvec, label) of the world-frame view, or None."""
        res = self._res
        if res is None:
            return None
        if self._world != "points" and self._world in res["cb_poses"]:
            rv, tv = res["cb_poses"][self._world]
            return rv, tv, f"board image {os.path.basename(self._world)}"
        if res.get("points_pose") is not None:
            return res["points_pose"][0], res["points_pose"][1], "points image"
        for c in (self._cb["cands"] if self._cb else []):
            if c["path"] in res["cb_poses"]:
                rv, tv = res["cb_poses"][c["path"]]
                return rv, tv, f"board image {os.path.basename(c['path'])} (first one; none chosen)"
        return None

    def _outputs(self) -> dict:
        res, ev, em = self._res, self._eval, self._eval_map
        if res is None or ev is None or em is None:
            return {}
        size = self._image_size()
        out = {"cmat": np.asarray(res["K"], dtype=np.float64),
               "dvec": np.asarray(res["D"], dtype=np.float64).reshape(1, -1),
               "rms": float(ev.rms.get("all", math.nan)),
               "rms_types": np.array([ev.rms.get(t, math.nan) for t in ce.RESIDUAL_TYPES], dtype=np.float64),
               "img_size": np.array(size if size else (0, 0), dtype=np.int64)}
        wp = self._world_pose()
        if wp is not None:
            out["rvec"] = np.asarray(wp[0], dtype=np.float64).reshape(3, 1)
            out["tvec"] = np.asarray(wp[1], dtype=np.float64).reshape(3, 1)
        pi, pt_rows = em["pi"], em["pt_rows"]
        if pi is not None:
            out["obj_pts"] = np.array([self._solved_obj(i) for i in pt_rows], dtype=np.float64).reshape(-1, 3)
            out["img_pts"] = np.array([(self.rows[i]["u"], self.rows[i]["v"]) for i in pt_rows],
                                      dtype=np.float64).reshape(-1, 2)
            out["proj_pts"] = np.asarray(ev.projected[pi], dtype=np.float64).reshape(-1, 2)
            out["pt_err"] = np.linalg.norm(np.asarray(ev.per_point[pi], dtype=np.float64).reshape(-1, 2), axis=1)
        paths = em["cb_paths"]
        if paths:
            cands = {c["path"]: c for c in self._cb["cands"]}
            out["rvecs"] = np.array([np.ravel(res["cb_poses"][p][0]) for p in paths], dtype=np.float64)
            out["tvecs"] = np.array([np.ravel(res["cb_poses"][p][1]) for p in paths], dtype=np.float64)
            out["cb_err"] = np.array([self._cb_err.get(p, math.nan) for p in paths], dtype=np.float64)
            out["cb_obj_pts"] = stack_padded([self._cand_obj(cands[p]) for p in paths])
            out["cb_img_pts"] = stack_padded([cands[p]["corners"].reshape(-1, 2).astype(np.float64) for p in paths])
        return out

    def compute(self, inputs: dict) -> dict:
        img = inputs.get("image")
        if img is not None and img is not self._pin_image:
            self._pin_image = img
            self._gray = None
            self._on_image_changed()
        elif img is None and self._pin_image is not None:
            self._pin_image = None
            self._gray = None
            self._on_image_changed()
        self._pin_K = inputs.get("cmat_in")
        self._pin_D = inputs.get("dvec_in")
        files = self._coerce_files(inputs.get("files"))
        if files != self._cb_pin_files:
            self._cb_pin_files = files
            self._update_cb_files_label()
        obj, ipt = inputs.get("object_points"), inputs.get("image_points")
        changed = not (self._same(obj, self._pin_obj) and self._same(ipt, self._pin_img))
        self._pin_obj = None if obj is None else np.array(obj, dtype=np.float64, copy=True)
        self._pin_img = None if ipt is None else np.array(ipt, dtype=np.float64, copy=True)
        if changed and self._follow_var.get() and self._pin_obj is not None and self._pin_img is not None:
            if self._load_from_pins(quiet=True, push=False):
                self._solve_sync()
        return self._outputs()

    @staticmethod
    def _same(a, b) -> bool:
        if a is None or b is None:
            return a is None and b is None
        a = np.asarray(a, dtype=np.float64)
        return a.shape == b.shape and np.array_equal(a, b, equal_nan=True)

    def _push(self) -> None:
        self._update_face()
        if self._request_downstream is not None:
            self._request_downstream(self.node_id)

    # ── node face ──────────────────────────────────────────────────────────
    def build_body(self) -> None:
        x, y, w, h = self.x, self.y, self.width, self.height
        self._body_rect = self.canvas.create_rectangle(
            x, y, x + w, y + h, fill=self.BODY_COLOR, outline="#555555", tags=(self.node_id, "node_body"))
        self._title_item = self.canvas.create_text(
            x + w / 2, y + 12, text=self.get_canvas_title(), font=("Arial", 9, "bold"),
            fill=self.TITLE_COLOR, tags=(self.node_id,))
        lbl = tk.Label(self.canvas, textvariable=self._face_var, font=("Arial", 8), bg=self.BODY_COLOR,
                       justify="center")
        self._face_win = self.canvas.create_window(x + w / 2, y + h / 2 + 10, window=lbl, tags=(self.node_id,))
        self._canvas_items += [self._body_rect, self._title_item]
        self._update_face()

    def _update_face(self) -> None:
        ev = self._eval
        if self._res is None or ev is None:
            n = sum(1 for i in self._infos if i["kind"] in ("point", "unknown", "line"))
            self._face_var.set(f"{len(self.rows)} rows, {n} usable\nno result yet")
            return
        K = self._res["K"]
        stale = "  (out of date)" if self._res.get("stale") else ""
        self._face_var.set(f"RMS {ev.rms.get('all', math.nan):.3f} px{stale}\n"
                           f"fx {K[0][0]:.1f}  fy {K[1][1]:.1f}")

    # ── inspector ──────────────────────────────────────────────────────────
    def build_inspector(self, parent: tk.Frame) -> None:
        self._reset_widget_refs()
        win = self._inspector_win
        if win is not None:
            sw, sh = win.winfo_screenwidth(), win.winfo_screenheight()
            win.geometry(f"{min(1500, sw - 80)}x{min(880, sh - 120)}")
            win.minsize(min(1000, sw - 80), min(560, sh - 120))
            for seq in ("<Control-h>", "<Control-H>"):
                win.bind(seq, lambda _e: self._open_help())
        apply_notebook_tab_style()

        # Summary bar (bottom), then the two panes.
        bar = tk.Frame(parent, bg="#eef2f7", highlightthickness=1, highlightbackground="#c8cdd8")
        bar.pack(side="bottom", fill="x", pady=(4, 0))
        self._summary_var = tk.StringVar()
        self._summary_lbl = tk.Label(bar, textvariable=self._summary_var, font=("Consolas", 9), bg="#eef2f7",
                                     anchor="w", justify="left")
        self._summary_lbl.pack(side="left", fill="x", expand=True, padx=6, pady=3)
        bar.bind("<Configure>", lambda e: self._summary_lbl is not None and self._summary_lbl.configure(
            wraplength=max(200, e.width - 130)))
        hb = tk.Button(bar, text="Help (Ctrl+H)", font=("Arial", 8), command=self._open_help)
        hb.pack(side="right", padx=4)
        tip_var = tooltip_toggle_var(bar)
        cbx = tk.Checkbutton(bar, text="tooltips", variable=tip_var, font=("Arial", 8), bg="#eef2f7",
                             command=lambda: set_tooltips_enabled(tip_var.get()))
        cbx.pack(side="right")
        add_tooltip(cbx, "Show these help boxes when the mouse rests on a control. Untick when they cover "
                         "something you need to see (applies to all nodes; remembered).")
        add_tooltip(hb, "How to use this node: '?' unknowns, straight lines, weights, fixed parameters.")
        add_tooltip(self._summary_lbl, "Result summary: RMS error of all residuals and per type (pixels), "
                                       "focal lengths, principal point, main distortion coefficients.")

        panes = tk.PanedWindow(parent, orient=tk.HORIZONTAL, sashwidth=6, sashrelief=tk.RAISED)
        panes.pack(fill="both", expand=True)
        left, right = tk.Frame(panes), tk.Frame(panes)
        panes.add(left, width=700, minsize=480, stretch="never")
        panes.add(right, minsize=400, stretch="always")

        lnb = ttk.Notebook(left)
        lnb.pack(fill="both", expand=True)
        tab_pts, tab_cb, tab_solve, tab_adj, tab_grid = (tk.Frame(lnb, padx=4, pady=4) for _ in range(5))
        for t, title in ((tab_pts, "Points"), (tab_cb, "Chessboard"), (tab_solve, "Solve"), (tab_adj, "Adjust"),
                         (tab_grid, "Grids")):
            lnb.add(t, text=title)
        self._build_points_tab(tab_pts)
        self._build_chessboard_tab(tab_cb)
        self._build_solve_tab(tab_solve)
        self._build_adjust_tab(tab_adj)
        self._build_grids_tab(tab_grid)
        lnb.bind("<<NotebookTabChanged>>", lambda _e: self._sync_sliders())

        rnb = ttk.Notebook(right)
        rnb.pack(fill="both", expand=True)
        self._rnb = rnb
        for title, builder in (("Image", self._build_image_tab), ("Errors", self._build_errors_tab),
                               ("3D", self._build_3d_tab), ("Undistort", self._build_undistort_tab),
                               ("Sensitivity", self._build_loo_tab), ("Runs", self._build_runs_tab)):
            t = tk.Frame(rnb)
            rnb.add(t, text=title)
            builder(t)
        rnb.bind("<<NotebookTabChanged>>", lambda _e: self._refresh_right())
        if self._loo_state is not None and not self._loo_state["finished"]:
            self._poll_loo()

        if self._solve_state is not None and not self._solve_state["finished"]:
            self._set_solving(True)
            self._poll_solve()
        if self._cb_state is not None and not self._cb_state["finished"]:
            self._set_detecting(True)
            self._poll_detect()
        self._refresh_all()

    def on_destroy(self) -> None:
        if self._search is not None:
            self._search["state"]["cancel"].set()
        if self._search_win is not None:
            try:
                self._search_win.destroy()
            except tk.TclError:
                pass
            self._search_win = None
        st = self._loo_state
        if st is not None and not st["finished"]:
            st["pool"].shutdown(wait=False, cancel_futures=True)
            st["finished"] = True
        super().on_destroy()

    def close_inspector(self) -> None:
        if self._help_popup is not None:
            try:
                self._help_popup.destroy()
            except tk.TclError:
                pass
        super().close_inspector()
        self._reset_widget_refs()

    # Points tab
    def _build_points_tab(self, tab) -> None:
        tb = tk.Frame(tab)
        tb.pack(fill="x")
        buttons = (("Add row", self._on_add_row, "Add an empty row below the selection; type its X, Y, Z, "
                                                  "then click the point in the image."),
                   ("Delete", self.delete_selected, "Delete the selected rows (key: Delete)."),
                   ("Up", lambda: self._move_rows(-1), "Move the selected rows up."),
                   ("Down", lambda: self._move_rows(1), "Move the selected rows down."),
                   ("Import...", self._on_import, "Read rows from a CSV or Excel file (columns name, X, Y, Z, "
                                                  "u, v, line, weight, use). You can replace or append."),
                   ("Export...", self._on_export, "Write the table to CSV or Excel (.xlsx, with a README "
                                                  "sheet), e.g. to edit it in a spreadsheet."),
                   ("Paste", self.paste_rows, "Append rows copied from Excel / text (Ctrl+V in the table). "
                                              "With a header line, columns are found by name; without one: "
                                              "name (optional), then 2 = u v, 3 = X Y Z, 5 = X Y Z u v, "
                                              "6 = X Y Z u v line."),
                   ("From input pins", lambda: self._load_from_pins(quiet=False),
                    "Replace the table by the arrays at the 'objPts' (N, 3; NaN = unknown) and "
                    "'imgPts' (N, 2) input pins."))
        for text, cmd, tip in buttons:
            b = tk.Button(tb, text=text, font=("Arial", 8), command=cmd)
            b.pack(side="left", padx=(0, 3))
            add_tooltip(b, tip)
        fol = tk.Checkbutton(tab, text="Follow input pins (reload + solve when they change)",
                             variable=self._follow_var, font=("Arial", 8))
        fol.pack(anchor="w")
        add_tooltip(fol, "When the objPts / imgPts input arrays change (e.g. tracked points of a new video "
                         "frame), replace the table by them and solve at once with the current settings. "
                         "Typical: intrinsics fixed, pose per frame.")

        mode = tk.Frame(tab)
        mode.pack(fill="x", pady=(2, 2))
        tk.Label(mode, text="Click in image:", font=("Arial", 8, "bold")).pack(side="left")
        rb1 = tk.Radiobutton(mode, text="Point", value="point", variable=self._mode_var, font=("Arial", 8))
        rb1.pack(side="left")
        add_tooltip(rb1, "A click sets the image point of the selected row, then selects the next row "
                         "without an image point. With no row selected, a click adds a new row.")
        rb2 = tk.Radiobutton(mode, text="Line", value="line", variable=self._mode_var, font=("Arial", 8))
        rb2.pack(side="left")
        add_tooltip(rb2, "Every click adds an image-only point of the straight 3D line named on the right "
                         "(points of one real straight edge / rail / pole). At least 3 per line.")
        self._line_combo = ttk.Combobox(mode, textvariable=self._line_var, width=6, font=("Arial", 8))
        self._line_combo.pack(side="left", padx=(0, 3))
        add_tooltip(self._line_combo, "Name of the line that clicks in Line mode add points to. Pick an "
                                      "existing one or type a new name.")
        nl = tk.Button(mode, text="New line", font=("Arial", 8), command=self._on_new_line)
        nl.pack(side="left")
        add_tooltip(nl, "Start the next line (L1, L2, ...) and switch to Line mode.")
        sn = tk.Checkbutton(mode, text="Snap to corner", variable=self._snap_var, font=("Arial", 8))
        sn.pack(side="left", padx=(8, 0))
        add_tooltip(sn, "Refine each click to the nearest image corner with sub-pixel accuracy "
                        "(cv2.cornerSubPix, 11 x 11 px window) -- for chessboard-like corners, marker "
                        "corners and crossings. Leave it off for points on plain edges or blobs.")

        self._sheet = Sheet(tab, self)
        self._sheet.frame.pack(fill="both", expand=True, pady=(2, 2))
        add_tooltip(self._sheet.canvas, _SYNTAX_HELP + "\n\nClick = select, double-click = edit (Enter = next "
                                        "row, Tab = next column), click 'use' = on / off, Delete, Ctrl+V, arrows.")
        legend = tk.Frame(tab)
        legend.pack(fill="x")
        for kind, text in (("point", "known"), ("unknown", "one '?'"), ("line", "line point"),
                           ("invalid", "invalid"), ("unpicked", "no image point"), ("unused", "not used")):
            tk.Label(legend, text="●", fg=KIND_COLORS[kind], font=("Arial", 10)).pack(side="left")
            tk.Label(legend, text=text, font=("Arial", 8)).pack(side="left", padx=(0, 6))
        tk.Label(tab, text=_SYNTAX_HELP, font=("Arial", 8), fg="#4a5568", justify="left", anchor="w",
                 wraplength=660).pack(fill="x")
        self._issues_var = tk.StringVar()
        tk.Label(tab, textvariable=self._issues_var, font=("Arial", 8), fg="#b00020", justify="left",
                 anchor="w", wraplength=660).pack(fill="x")

    # Solve tab
    def _build_solve_tab(self, tab) -> None:
        top = tk.Frame(tab)
        top.pack(fill="x")
        tk.Label(top, text="Image size", font=("Arial", 8)).pack(side="left")
        for k, var in enumerate(self._size_vars):
            e = tk.Entry(top, textvariable=var, width=6, justify="right")
            e.pack(side="left", padx=(4, 0))
            add_tooltip(e, "Image width / height in pixels. Taken from the image automatically; type it "
                           "only when there is no image.")
        tk.Label(top, text="   Distortion model", font=("Arial", 8)).pack(side="left")
        cb = ttk.Combobox(top, textvariable=self._model_var, state="readonly", width=11,
                          values=tuple(ce.DIST_MODELS))
        cb.pack(side="left", padx=(4, 0))
        cb.bind("<<ComboboxSelected>>", lambda _e: (self._build_param_grid(), self._mark_stale()))
        add_tooltip(cb, "simple: k1 k2 p1 p2 k3 (OpenCV default)\nrational: + k4 k5 k6\n"
                        "thin_prism: + s1 s2 s3 s4\ntilted: + taux tauy (tilted sensor).\n"
                        "More coefficients need more (and wider spread) points; fix the ones you do not need.")

        use = tk.Frame(tab)
        use.pack(fill="x", pady=(4, 0))
        tk.Label(use, text="Use:", font=("Arial", 8, "bold")).pack(side="left")
        for text, var, tip in (
                ("point table", self._use_points_var,
                 "Use the rows of the Points tab (the points image): its pose, '?' unknowns and lines."),
                ("chessboard images", self._use_cb_var,
                 "Use the CHECKED images of the Chessboard tab (their corners and poses).")):
            cbx = tk.Checkbutton(use, text=text, variable=var, font=("Arial", 8), command=self._mark_stale)
            cbx.pack(side="left")
            add_tooltip(cbx, tip)
        tk.Label(use, text="   both:", font=("Arial", 8)).pack(side="left")
        for text, val, tip in (
                ("jointly", "joint", "One least-squares problem with all errors together (weighted per type): "
                                     "the points image also influences the camera parameters."),
                ("two steps", "two_step", "1) the chessboard images alone give the camera parameters; 2) with "
                                          "them fixed, only the points image's pose and its '?' unknowns are "
                                          "solved. The points image cannot change the camera.")):
            rb = tk.Radiobutton(use, text=text, value=val, variable=self._combine_var, font=("Arial", 8),
                                command=self._mark_stale)
            rb.pack(side="left")
            add_tooltip(rb, tip)

        self._param_frame = tk.Frame(tab)
        self._param_frame.pack(fill="x", pady=(4, 2))
        self._build_param_grid()

        row = tk.Frame(tab)
        row.pack(fill="x", pady=(2, 2))
        asp = tk.Checkbutton(row, text="Keep fy / fx (aspect ratio) fixed", variable=self._aspect_var,
                             font=("Arial", 8), command=self._mark_stale)
        asp.pack(side="left")
        add_tooltip(asp, "fy follows fx with the ratio of their initial values (1 when both are automatic): "
                         "square pixels. Recommended with one image and few points.")
        for text, cmd, tip in (
                ("Automatic initial", self._on_auto_initial,
                 "Clear all initial values: focal length = larger image side, principal point = image centre, "
                 "distortion 0."),
                ("Result → initial", self._on_result_to_initial,
                 "Copy the current result into the initial values (e.g. before fixing some of them)."),
                ("Input cmat/dvec → initial", self._on_pins_to_initial,
                 "Copy the camera matrix / distortion from the 'cmat in' / 'dvec in' input pins (e.g. from a "
                 "chessboard calibration) into the initial values. Fix them to keep them.")):
            b = tk.Button(row, text=text, font=("Arial", 8), command=cmd)
            b.pack(side="left", padx=(6, 0))
            add_tooltip(b, tip)

        wbox = tk.LabelFrame(tab, text="Error weights", font=("Arial", 8, "bold"), padx=4, pady=2)
        wbox.pack(fill="x", pady=(4, 2))
        for k, (t, label, tip) in enumerate((
                ("point", "known points", "Weight of the reprojection errors of fully known points."),
                ("unknown", "'?' points", "Weight of the reprojection errors of points with one '?' coordinate."),
                ("line", "line points", "Weight of the straight-line errors (distance of each line point to "
                                        "the best straight line, after removing the distortion)."),
                ("chessboard", "board corners", "Weight of the reprojection errors of the chessboard corners "
                                                "(only matters when combined with the point table, jointly)."))):
            lbl = tk.Label(wbox, text=label, font=("Arial", 8))
            lbl.grid(row=0, column=2 * k, sticky="e")
            e = tk.Entry(wbox, textvariable=self._w_vars[t], width=6, justify="right")
            e.grid(row=0, column=2 * k + 1, padx=(3, 10))
            add_tooltip(lbl, tip + " Each squared error is multiplied by this (and by the row's weight).")
            add_tooltip(e, tip + " Each squared error is multiplied by this (and by the row's weight).")
        lbl = tk.Label(wbox, text="loss", font=("Arial", 8))
        lbl.grid(row=1, column=0, sticky="e")
        loss = ttk.Combobox(wbox, textvariable=self._loss_var, values=("linear", "huber"), state="readonly", width=7)
        loss.grid(row=1, column=1, padx=(3, 10), pady=(2, 0))
        tip = ("linear: ordinary least squares (best when all points are good).\nhuber: robust -- errors "
               "larger than the given pixels count only linearly, so a few wrongly clicked points pull the "
               "result much less; they stand out with large errors.")
        add_tooltip(lbl, tip)
        add_tooltip(loss, tip)
        lbl = tk.Label(wbox, text="huber px", font=("Arial", 8))
        lbl.grid(row=1, column=2, sticky="e")
        e = tk.Entry(wbox, textvariable=self._huber_var, width=6, justify="right")
        e.grid(row=1, column=3, padx=(3, 10), pady=(2, 0))
        add_tooltip(lbl, "Huber threshold in pixels: errors below it are treated normally. About 2-3 x the "
                         "expected clicking accuracy.")
        add_tooltip(e, "Huber threshold in pixels: errors below it are treated normally. About 2-3 x the "
                       "expected clicking accuracy.")

        srow = tk.Frame(tab)
        srow.pack(fill="x", pady=(6, 2))
        self._solve_btn = tk.Button(srow, text="Solve", font=("Arial", 10, "bold"), width=12, bg="#d9f1d8",
                                    command=self._on_solve)
        self._solve_btn.pack(side="left")
        add_tooltip(self._solve_btn, "Run the calibration with the usable rows and the settings above. "
                                     "The result goes to the output pins. The line below says HOW it will be "
                                     "solved.")
        self._cancel_solve_btn = tk.Button(srow, text="Cancel", font=("Arial", 8), state="disabled",
                                           command=self._on_cancel_solve)
        self._cancel_solve_btn.pack(side="left", padx=(4, 0))
        add_tooltip(self._cancel_solve_btn, "Stop a running Solve; the previous result stays. (A running "
                                            "cv2.calibrateCamera call cannot be interrupted: it finishes in "
                                            "the background and its result is thrown away.)")
        self._status_var = tk.StringVar()
        tk.Label(srow, textvariable=self._status_var, font=("Arial", 8), fg="#334455", anchor="w",
                 justify="left", wraplength=520).pack(side="left", padx=8, fill="x", expand=True)
        rrow = tk.Frame(tab)
        rrow.pack(fill="x")
        self._route_var = tk.StringVar()
        self._route_lbl = tk.Label(rrow, textvariable=self._route_var, font=("Arial", 8), anchor="w",
                                   justify="left", wraplength=560)
        self._route_lbl.pack(side="left", fill="x", expand=True)
        add_tooltip(self._route_lbl, "How Solve works with the current data and settings. OpenCV only: "
                                     "cv2.calibrateCamera (and cv2.solvePnP for the points image in 'two "
                                     "steps') solve the whole problem -- fast. Least-squares minimization: "
                                     "this node's own solver (scipy), started from OpenCV, is needed for what "
                                     "OpenCV cannot do: '?' unknowns, straight lines, the huber loss, "
                                     "different weights, or fixed parameters OpenCV cannot fix one by one.")
        cbx = tk.Checkbutton(rrow, text="Always refine", variable=self._force_lsq_var, font=("Arial", 8))
        cbx.pack(side="right")
        add_tooltip(cbx, "Also run the least-squares minimization when OpenCV alone could solve it. OpenCV "
                         "stops after its usual 30 iterations: with many free distortion coefficients (e.g. "
                         "rational k1-k6 all free) the result can then be a hair above the true minimum "
                         "(about 0.0001 px RMS). Slower.")
        self._result_text = tk.Text(tab, height=14, font=("Consolas", 8), wrap="word", relief=tk.FLAT, bg="#f7f8fa")
        self._result_text.pack(fill="both", expand=True, pady=(4, 0))
        add_tooltip(self._result_text, "Full result: camera matrix, distortion, pose (OpenCV: X_camera = "
                                       "R(rvec) X_object + tvec), camera position in object coordinates, "
                                       "errors per type, solver report.")

    def _build_param_grid(self) -> None:
        f = self._param_frame
        if f is None:
            return
        for w in f.winfo_children():
            w.destroy()
        names = list(ce.INTR_NAMES) + list(ce.DIST_NAMES[:ce.DIST_MODELS.get(self._model_var.get(), 5)])
        per_col = 9
        self._res_vars = {}
        for blk in range((len(names) + per_col - 1) // per_col):
            for j, title in enumerate(("param", "fix", "initial", "result")):
                tk.Label(f, text=title, font=("Arial", 8, "bold"), fg="#445").grid(
                    row=0, column=blk * 5 + j, sticky="w", padx=(10 if j == 0 and blk else 0, 2))
            for r, name in enumerate(names[blk * per_col:(blk + 1) * per_col], start=1):
                c0 = blk * 5
                tk.Label(f, text=name, font=("Arial", 8)).grid(row=r, column=c0, sticky="w",
                                                               padx=(10 if blk else 0, 2))
                cb = tk.Checkbutton(f, variable=self._fix_vars[name], command=self._mark_stale)
                cb.grid(row=r, column=c0 + 1)
                add_tooltip(cb, f"Fix {name}: keep it at its initial value (automatic when empty) instead "
                                f"of solving it.")
                e = tk.Entry(f, textvariable=self._init_vars[name], width=10, justify="right")
                e.grid(row=r, column=c0 + 2, padx=(0, 4))
                add_tooltip(e, f"Initial value of {name} (start of the search; the value used when fixed). "
                               f"Empty = automatic.")
                var = tk.StringVar(value="")
                self._res_vars[name] = var
                tk.Label(f, textvariable=var, font=("Consolas", 8), width=12, anchor="e").grid(
                    row=r, column=c0 + 3, sticky="e")
        self._refresh_result_values()

    # Image tab
    def _build_image_tab(self, tab) -> None:
        bar = tk.Frame(tab)
        bar.pack(fill="x")
        ob = tk.Button(bar, text="Open image...", font=("Arial", 8), command=self._on_open_image)
        ob.pack(side="left")
        add_tooltip(ob, "Show an image file (used when nothing is connected to the 'image' input pin). "
                        "Saved with the project (relative path when possible).")
        fb = tk.Button(bar, text="Fit", font=("Arial", 8), command=lambda: self._iv and self._iv.fit())
        fb.pack(side="left", padx=(3, 8))
        add_tooltip(fb, "Show the whole image.")
        for text, var, tip in (("names", self._show_names_var, "Show the row names next to the points."),
                               ("projections", self._show_proj_var,
                                "Show where the result projects each object point (magenta +) and the error "
                                "vector from the clicked point, magnified by 'Error x'."),
                               ("lines", self._show_lines_var,
                                "Show each straight line fitted to its points (blue curve; curved by the lens "
                                "distortion) and every point's distance to it, magnified by 'Error x'."),
                               ("? axes", self._show_axes_var,
                                "For points with a '?' coordinate: the 3D line of all positions the point "
                                "may have (dashed orange) -- the solved position is on it.")):
            cbx = tk.Checkbutton(bar, text=text, variable=var, font=("Arial", 8), command=self._redraw_image)
            cbx.pack(side="left")
            add_tooltip(cbx, tip)
        lbl = tk.Label(bar, text="  Error x", font=("Arial", 8))
        lbl.pack(side="left")
        g = ttk.Combobox(bar, textvariable=self._gain_var, values=ERROR_GAINS, width=4)
        g.pack(side="left")
        g.bind("<<ComboboxSelected>>", lambda _e: self._redraw_image())
        g.bind("<Return>", lambda _e: self._redraw_image())
        tip = ("Magnification of the drawn error vectors / line distances (errors of 0.5 px are invisible "
               "at 1 x). Only the drawing changes.")
        add_tooltip(lbl, tip)
        add_tooltip(g, tip)
        srow = tk.Frame(tab)
        srow.pack(fill="x")
        for text, cmd, tip in (("Points image", lambda: self._show_board(None),
                                "Show the points image (the one of the Points tab; clicks edit its points)."),
                               ("\u25c0 board", lambda: self._step_board(-1),
                                "Show the previous CHECKED chessboard image."),
                               ("board \u25b6", lambda: self._step_board(1),
                                "Show the next CHECKED chessboard image.")):
            b = tk.Button(srow, text=text, font=("Arial", 8), command=cmd)
            b.pack(side="left", padx=(0, 3))
            add_tooltip(b, tip)
        self._img_src_var = tk.StringVar()
        tk.Label(srow, textvariable=self._img_src_var, font=("Arial", 8), fg="#555555", anchor="w").pack(
            side="left", fill="x", padx=(6, 0))
        self._iv = ImageView(tab, self._on_image_click, self._on_image_drag, self._on_image_motion)
        self._iv.canvas.pack(fill="both", expand=True)
        self._iv.draw_overlay = self._draw_overlay
        self._iv.hit_point = self._hit_point
        self._iv.empty_text = "No image: connect the 'image' pin or use 'Open image...'."
        add_tooltip(self._iv.canvas, "wheel = zoom, left-drag on empty space = pan, left-click = set the "
                                     "selected row's point (or add one), left-drag a point = move it.")
        bottom = tk.Frame(tab)
        bottom.pack(fill="x")
        self._mouse_var = tk.StringVar()
        tk.Label(bottom, textvariable=self._mouse_var, font=("Consolas", 8), fg="#555555").pack(side="left")
        tk.Label(bottom, text="wheel = zoom, left-drag = pan", font=("Arial", 8), fg="#888888").pack(side="right")
        self._iv.set_image(self._shown_image())

    # ── refresh ────────────────────────────────────────────────────────────
    def _refresh_all(self) -> None:
        self._infos = [classify_row(r) for r in self.rows]
        self._update_issues()
        self._update_row_errors()
        if self._sheet is not None:
            self._sheet.redraw()
        if self._line_combo is not None:
            self._line_combo.configure(values=self._line_names())
        self._refresh_result_values()
        self._refresh_summary()
        self._refresh_cb_tree()
        self._update_image_source_label()
        self._redraw_image()
        self._update_face()
        self._refresh_right()
        self._schedule_route_text()

    def _update_issues(self) -> None:
        if self._infos and len(self._infos) == len(self.rows):
            _v, _p, _l, issues = build_view(self.rows, self._infos)
        else:
            issues = []
        self._issues = issues
        if self._issues_var is not None:
            more = f"\n... and {len(issues) - 6} more" if len(issues) > 6 else ""
            self._issues_var.set("\n".join(issues[:6]) + more)

    def _line_names(self) -> list:
        return sorted({r["line"] for r in self.rows if r["line"]})

    def cell(self, i: int, key: str):
        """(text, fg, bg) of one table cell."""
        r, info = self.rows[i], self._infos[i] if i < len(self._infos) else classify_row(self.rows[i])
        fg, bg = "#111111", "#ffffff"
        if key == "#":
            return str(i + 1), "#ffffff", KIND_COLORS[info["kind"]]
        if key == "use":
            return ("✔" if r["use"] else ""), "#1a7f37", "#ffffff"
        if key in ("X", "Y", "Z"):
            text = r[key]
            if key in info["bad"]:
                return text or "(empty)", "#d00000", "#ffe3e3"
            if "?" in text:
                val = self._solved_value(r, "XYZ".index(key))
                return (f"{text} → {val:.4g}" if val is not None else text), "#c25e00", "#fff4e6"
            return text, fg, bg
        if key in ("u", "v"):
            return (f"{r[key]:.2f}" if math.isfinite(r[key]) else ""), fg, bg
        if key == "weight":
            return (f"{r['weight']:g}"), ("#999999" if r["weight"] == 1.0 else fg), bg
        if key == "line":
            return r["line"], "#0b6e99", bg
        if key == "err":
            text, color = self._row_err.get(r["id"], ("", fg))
            return text, color, bg
        return str(r.get(key, "")), fg, bg

    def _solved_value(self, r: dict, axis: int):
        if self._res is None:
            return None
        v = self._res["unknowns"].get(r["id"])
        return None if v is None or v[0] != axis else v[1]

    def _solved_obj(self, i: int) -> list:
        info = self._infos[i]
        o = list(info["obj"])
        if info["axis"] is not None:
            val = self._eval_unknowns.get(i) if getattr(self, "_eval_unknowns", None) else None
            o[info["axis"]] = math.nan if val is None else val
        return o

    def _update_row_errors(self) -> None:
        """Evaluate the stored result on the CURRENT data (instant feedback
        after edits): the points view with its stored pose, every checked
        board image that has a stored pose. Fills the error columns."""
        self._row_err, self._cb_err = {}, {}
        self._eval, self._eval_map, self._eval_unknowns = None, None, {}
        res = self._res
        if res is None:
            return
        views, poses, pt_rows, line_rows, pi = [], [], [], [], None
        if res.get("points_pose") is not None and self._use_points_var.get():
            view, pt_rows, line_rows, _issues = build_view(self.rows, self._infos)
            if len(view.obj) or view.lines:
                pi = 0
                views.append(view)
                poses.append(tuple(np.asarray(a, float).reshape(3) for a in res["points_pose"]))
        cb_paths = []
        if self._use_cb_var.get():
            for path, v in self._cb_views():
                pose = res["cb_poses"].get(path)
                if pose is not None:
                    views.append(v)
                    poses.append(tuple(np.asarray(a, float).reshape(3) for a in pose))
                    cb_paths.append(path)
        if not views:
            return
        K, D = np.asarray(res["K"], float), np.asarray(res["D"], float)
        try:
            prob = self._make_problem(views)
            unk = ce.initial_unknowns(prob, K, D, poses)
            for k, i in enumerate(pt_rows):
                stored = res["unknowns"].get(self.rows[i]["id"])
                if (0, k) in unk and stored is not None and stored[0] == self._infos[i]["axis"]:
                    unk[(0, k)] = stored[1]
            ev = ce.evaluate(prob, K, D, poses, unk)
        except Exception:
            return
        self._eval = ev
        self._eval_map = {"pt_rows": pt_rows, "line_rows": line_rows, "pi": pi, "cb_paths": cb_paths}
        for j, path in enumerate(cb_paths):
            r = np.asarray(ev.per_point[j + (pi is not None)], float)
            self._cb_err[path] = float(np.sqrt(np.mean(np.sum(r * r, axis=1)))) if len(r) else math.nan
        if pi is None:
            return
        self._eval_unknowns = {i: ev.unknowns[(0, k)] for k, i in enumerate(pt_rows) if (0, k) in ev.unknowns}
        rms = ev.rms.get("all", 1.0)
        thr = max(1.0, 3.0 * rms) if math.isfinite(rms) else 1.0

        def color(e):
            return "#1a7f37" if e <= max(0.5, 1.5 * rms) else ("#b07800" if e <= thr else "#d00000")
        for k, i in enumerate(pt_rows):
            e = float(np.linalg.norm(ev.per_point[0][k]))
            self._row_err[self.rows[i]["id"]] = (f"{e:.3f}", color(e))
        for j, (_name, idx) in enumerate(line_rows):
            fit = ev.line_fits[0][j]
            if fit is None:
                continue
            for k, i in enumerate(idx):
                e = abs(float(fit["residual"][k]))
                rid = self.rows[i]["id"]
                if rid in self._row_err:
                    self._row_err[rid] = (self._row_err[rid][0] + f" | L {e:.3f}", self._row_err[rid][1])
                else:
                    self._row_err[rid] = (f"L {e:.3f}", color(e))

    def _refresh_result_values(self) -> None:
        if not self._res_vars:
            return
        res = self._res
        for name, var in self._res_vars.items():
            if res is None:
                var.set("")
                continue
            K, D = res["K"], res["D"]
            if name in ce.INTR_NAMES:
                val = {"fx": K[0][0], "fy": K[1][1], "cx": K[0][2], "cy": K[1][2]}[name]
                var.set(f"{val:.3f}")
            else:
                j = ce.DIST_NAMES.index(name)
                var.set(f"{D[j]:.6g}" if j < len(D) else "0")

    def _refresh_summary(self) -> None:
        ev, res = self._eval, self._res
        if self._summary_var is not None:
            if ev is None or res is None:
                self._summary_var.set("No result yet. Fill the Points table, then press Solve (Solve tab).")
                if self._summary_lbl is not None:
                    self._summary_lbl.configure(fg="#333333")
            else:
                K, D = res["K"], res["D"]
                parts = [f"RMS all {ev.rms.get('all', math.nan):.3f} px"]
                for t, label in (("point", "known"), ("unknown", "?"), ("line", "line"), ("chessboard", "board")):
                    if t in ev.rms:
                        parts.append(f"{label} {ev.rms[t]:.3f}")
                parts.append(f"| fx {K[0][0]:.1f} fy {K[1][1]:.1f} cx {K[0][2]:.1f} cy {K[1][2]:.1f}")
                parts.append(f"k1 {D[0]:+.4f} k2 {D[1]:+.4f}" if len(D) >= 2 else "")
                text = "  ".join(p for p in parts if p)
                warn = self._coverage_warning()
                if res.get("adjusted"):
                    text += "\nMANUALLY ADJUSTED with the sliders (Adjust tab: 'Revert' gives back the solved values)."
                if res.get("stale"):
                    text += "\nThe table or settings changed since this result -- errors shown for the old " \
                            "camera; press Solve to update."
                if warn:
                    text += "\n" + warn
                self._summary_var.set(text)
                if self._summary_lbl is not None:
                    self._summary_lbl.configure(fg="#9a5b00" if (res.get("stale") or warn or res.get("adjusted"))
                                                else "#14532d")
        if self._result_text is not None:
            self._result_text.configure(state="normal")
            self._result_text.delete("1.0", tk.END)
            self._result_text.insert("1.0", self._result_report())
            self._result_text.configure(state="disabled")

    def _result_report(self) -> str:
        res, ev, em = self._res, self._eval, self._eval_map
        if res is None:
            return ""
        K = np.asarray(res["K"], float)
        D = np.asarray(res["D"], float).ravel()
        lines = ["camera matrix (cmat):"]
        lines += ["  [" + "  ".join(f"{v:12.4f}" for v in row) + " ]" for row in K]
        lines.append("distortion (dvec): " + "  ".join(
            f"{n}={D[j]:.6g}" for j, n in enumerate(ce.DIST_NAMES[:len(D)])))
        wp = self._world_pose()
        if wp is not None:
            rv, tv = np.asarray(wp[0], float).reshape(3), np.asarray(wp[1], float).reshape(3)
            R, _ = cv2.Rodrigues(rv)
            C = -R.T @ tv
            lines.append(f"world frame = {wp[2]}  (outputs rvec / tvec; X_camera = R(rvec) X_world + tvec)")
            lines.append(f"  rvec: [{rv[0]: .6f} {rv[1]: .6f} {rv[2]: .6f}]   "
                         f"({math.degrees(float(np.linalg.norm(rv))):.3f} deg)")
            lines.append(f"  tvec: [{tv[0]: .4f} {tv[1]: .4f} {tv[2]: .4f}]")
            lines.append(f"  camera position in world coordinates (-R^T tvec): [{C[0]:.4f} {C[1]:.4f} {C[2]:.4f}]")
        if ev is not None:
            lines.append("RMS error (px): " + "  ".join(f"{k} {v:.4f}" for k, v in ev.rms.items()))
        if em is not None:
            used = []
            if em["pi"] is not None:
                used.append(f"{len(em['pt_rows'])} table points, {len(em['line_rows'])} lines")
            if em["cb_paths"]:
                used.append(f"{len(em['cb_paths'])} chessboard images")
            lines.append("used: " + ", ".join(used))
        if res.get("report"):
            lines.append(res["report"])
        return "\n".join(lines)

    def _coverage_warning(self) -> str:
        """Warn when the measurements reach only the image centre while
        higher-order distortion is solved (it is then extrapolated near the
        corners)."""
        em = self._eval_map
        if self._res is None or em is None:
            return ""
        size = self._image_size()
        if not size:
            return ""
        idx = set(em["pt_rows"]) | {i for _n, rr in em["line_rows"] for i in rr}
        pts = [np.array([(self.rows[i]["u"], self.rows[i]["v"]) for i in idx], dtype=float).reshape(-1, 2)]
        if em["cb_paths"]:
            cands = {c["path"]: c for c in self._cb["cands"]}
            pts += [cands[p]["corners"].reshape(-1, 2) for p in em["cb_paths"]]
        pts = np.vstack(pts)
        if len(pts) == 0:
            return ""
        K = self._res["K"]
        c = np.array([K[0][2], K[1][2]])
        corners = np.array([[0, 0], [size[0], 0], [0, size[1]], [size[0], size[1]]], dtype=float)
        reach = float(np.max(np.linalg.norm(pts - c, axis=1)) / np.max(np.linalg.norm(corners - c, axis=1)))
        free = [n for n in ("k2", "k3", "k4", "k5", "k6") if n in self._res.get("free", [])]
        if reach < 0.75 and free:
            return (f"The measurements reach only {100 * reach:.0f}% of the way to the image corners while "
                    f"{', '.join(free)} {'is' if len(free) == 1 else 'are'} solved: the distortion near the "
                    f"corners is extrapolated. Add points / lines / board views near the edges, or fix "
                    f"{free[-1]}.")
        return ""

    # ── table editing ──────────────────────────────────────────────────────
    def _new_id(self) -> int:
        self._next_id += 1
        return self._next_id - 1

    def _next_name(self, prefix: str = "P") -> str:
        used = {r["name"] for r in self.rows}
        k = 1
        while f"{prefix}{k:03d}" in used:
            k += 1
        return f"{prefix}{k:03d}"

    def _changed(self, mark_stale: bool = True, push: bool = True) -> None:
        if mark_stale:
            self._mark_stale(refresh=False)
        self._refresh_all()
        if push:
            self._push()
        else:
            self._update_face()

    def _mark_stale(self, refresh: bool = True) -> None:
        if self._res is not None and not self._res.get("stale"):
            self._res["stale"] = True
        if refresh:
            self._refresh_summary()
            self._update_face()

    def on_selection_changed(self) -> None:
        sel = self._sheet.selected if self._sheet else []
        if self._iv is not None and len(sel) == 1:
            r = self.rows[sel[0]]
            if math.isfinite(r["u"]) and math.isfinite(r["v"]):
                self._iv.ensure_visible(r["u"], r["v"])
        self._redraw_image()

    def toggle_use(self, i: int) -> None:
        self.rows[i]["use"] = not self.rows[i]["use"]
        self._changed()

    def edit_cell(self, i: int, key: str, value: str) -> None:
        if not (0 <= i < len(self.rows)):
            return
        r = self.rows[i]
        value = value.strip()
        if key in ("u", "v", "weight"):
            if value == "":
                new = 1.0 if key == "weight" else math.nan
            else:
                try:
                    new = float(value)
                except ValueError:
                    messagebox.showerror("Edit", f"'{value}' is not a number.", parent=self._dlg_parent())
                    return
            if new == r[key] or (math.isnan(new) and math.isnan(r[key])):
                return
            r[key] = new
        else:
            if key in ("X", "Y", "Z") and value:
                try:
                    ce.parse_object_cell(value)
                except ValueError:
                    messagebox.showwarning("Edit", f"'{value}' is not a number, '?' or '?=value'. It is kept, "
                                                   f"shown in red, and the row is not used until it is fixed.",
                                           parent=self._dlg_parent())
            if value == r[key]:
                return
            r[key] = value
        self._changed()

    def _on_add_row(self) -> None:
        sel = self._sheet.selected if self._sheet else []
        at = (sel[-1] + 1) if sel else len(self.rows)
        self.rows.insert(at, new_row(self._new_id(), self._next_name()))
        self._changed()
        if self._sheet is not None:
            self._sheet.select([at])

    def delete_selected(self) -> str:
        sel = self._sheet.selected if self._sheet else []
        if not sel:
            return "break"
        keep = set(range(len(self.rows))) - set(sel)
        self.rows = [r for i, r in enumerate(self.rows) if i in keep]
        if self._sheet is not None:
            self._sheet.select([min(sel[0], len(self.rows) - 1)] if self.rows else [])
        self._changed()
        return "break"

    def _move_rows(self, d: int) -> None:
        sel = self._sheet.selected if self._sheet else []
        if not sel:
            return
        if (d < 0 and sel[0] == 0) or (d > 0 and sel[-1] == len(self.rows) - 1):
            return
        order = list(range(len(self.rows)))
        for i in (sel if d < 0 else reversed(sel)):
            order[i], order[i + d] = order[i + d], order[i]
        self.rows = [self.rows[i] for i in order]
        self._sheet.select([i + d for i in sel])
        self._changed(mark_stale=False)

    def paste_rows(self) -> str:
        try:
            text = self.canvas.winfo_toplevel().clipboard_get()
        except tk.TclError:
            messagebox.showinfo("Paste", "The clipboard is empty.", parent=self._dlg_parent())
            return "break"
        raw = [ln.split("\t") if "\t" in ln else next(csv.reader([ln])) for ln in text.splitlines() if ln.strip()]
        self._add_records(raw, "Paste", replace=False)
        return "break"

    def _add_records(self, raw: list, title: str, replace: bool) -> None:
        names = []

        def next_name():
            used = {r["name"] for r in self.rows} | set(names)
            k = 1
            while f"P{k:03d}" in used:
                k += 1
            names.append(f"P{k:03d}")
            return names[-1]
        rows, errors = parse_records(raw, self._new_id, next_name)
        if errors:
            messagebox.showerror(title, f"{len(errors)} problem(s) -- nothing was changed:\n\n"
                                 + "\n".join(errors[:12]), parent=self._dlg_parent())
            return
        if not rows:
            return
        self.rows = rows if replace else self.rows + rows
        self._changed()
        if self._sheet is not None:
            self._sheet.select([len(self.rows) - len(rows)])

    def _load_from_pins(self, quiet: bool, push: bool = True) -> bool:
        obj, img = self._pin_obj, self._pin_img
        msg = None
        if obj is None or img is None:
            msg = "Connect both the 'objPts' (N, 3) and the 'imgPts' (N, 2) input pins and run the graph."
        else:
            obj, img = obj.reshape(-1, 3) if obj.size % 3 == 0 else None, img.reshape(-1, 2) if img.size % 2 == 0 else None
            if obj is None or img is None or len(obj) != len(img):
                msg = "objPts must be (N, 3) and imgPts (N, 2) with the same N."
        if msg:
            if not quiet:
                messagebox.showerror("From input pins", msg, parent=self._dlg_parent())
            return False
        old = {r["name"]: r for r in self.rows}
        rows = []
        for k, (o, p) in enumerate(zip(obj, img)):
            name = f"P{k + 1:03d}"
            cells = ["?" if not math.isfinite(c) else repr(float(c)) for c in o]
            prev = old.get(name)
            rows.append(new_row(self._new_id(), name, *cells, p[0], p[1],
                                weight=prev["weight"] if prev else 1.0, use=prev["use"] if prev else True))
        self.rows = rows
        self._changed(push=push)
        return True

    def _on_new_line(self) -> None:
        names = self._line_names()
        k = 1
        while f"L{k}" in names:
            k += 1
        self._line_var.set(f"L{k}")
        self._mode_var.set("line")
        if self._line_combo is not None:
            self._line_combo.configure(values=names + [f"L{k}"])

    # ── image ──────────────────────────────────────────────────────────────
    def _current_image(self):
        """The POINTS image (input pin first, else the opened file)."""
        return self._pin_image if self._pin_image is not None else self._file_image

    def _cb_size(self):
        if self._cb and self._cb["cands"]:
            return tuple(int(v) for v in self._cb["cands"][0]["size"])
        return None

    def _image_size(self):
        """Size of the camera's images: the points image, else the typed size,
        else the chessboard images."""
        img = self._current_image()
        if img is not None:
            a = np.asarray(img)
            return int(a.shape[1]), int(a.shape[0])
        try:
            w, h = int(float(self._size_vars[0].get())), int(float(self._size_vars[1].get()))
            if w > 0 and h > 0:
                return w, h
        except ValueError:
            pass
        return self._cb_size()

    def _shown_image(self):
        if self._show is None:
            return self._current_image()
        if self._board_img[0] != self._show:
            self._board_img = (self._show, _read_rgb(self._show))
        return self._board_img[1]

    def _on_image_changed(self) -> None:
        size = self._image_size()
        if self._current_image() is not None and size:
            self._size_vars[0].set(str(size[0]))
            self._size_vars[1].set(str(size[1]))
        if self._iv is not None:
            self._iv.set_image(self._shown_image())
        self._update_image_source_label()

    def _update_image_source_label(self) -> None:
        if self._img_src_var is None:
            return
        if self._show is not None:
            used = self._used_board_paths()
            pos = f"{used.index(self._show) + 1}/{len(used)} checked" if self._show in used else "not checked"
            err = self._cb_err.get(self._show)
            err_txt = f", RMS {err:.3f} px" if err is not None and math.isfinite(err) else ""
            world = ", WORLD FRAME" if self._world == self._show else ""
            self._img_src_var.set(f"board image {os.path.basename(self._show)} ({pos}{err_txt}{world})")
            return
        size = self._image_size()
        dims = f"{size[0]} x {size[1]}" if size else "size unknown"
        if self._pin_image is not None:
            self._img_src_var.set(f"points image: input pin ({dims})")
        elif self._file_image is not None:
            self._img_src_var.set(f"points image: {os.path.basename(self._image_path)} ({dims})")
        else:
            self._img_src_var.set("no points image")

    def _show_board(self, path) -> None:
        self._show = path
        if self._iv is not None:
            self._iv.empty_text = ("Cannot read this board image." if path is not None else
                                   "No image: connect the 'image' pin or use 'Open image...'.")
            self._iv.set_image(self._shown_image())
        self._update_image_source_label()
        self._redraw_image()
        self._refresh_right()

    def _used_board_paths(self) -> list:
        return [c["path"] for c in (self._cb["cands"] if self._cb else []) if c["path"] in self._cb_use]

    def _step_board(self, d: int) -> None:
        used = self._used_board_paths()
        if not used:
            return
        if self._show in used:
            k = (used.index(self._show) + d) % len(used)
        else:
            k = 0 if d > 0 else len(used) - 1
        self._show_board(used[k])
        self._select_tree_path(used[k])

    def _project_dir(self):
        d = get_project_directory()
        return Path(d) if d else None

    def _rel(self, path: str) -> str:
        base = self._project_dir()
        if base is None:
            return path
        try:
            return os.path.relpath(path, base)
        except ValueError:                                     # another drive
            return path

    def _abs(self, path: str) -> str:
        p = Path(path)
        if not p.is_absolute() and self._project_dir() is not None:
            p = self._project_dir() / p
        return os.path.normpath(str(p))

    def _load_image_file(self, path: str) -> bool:
        try:
            buf = np.fromfile(self._abs(path), dtype=np.uint8)
            bgr = cv2.imdecode(buf, cv2.IMREAD_COLOR) if buf.size else None
        except Exception:
            bgr = None
        if bgr is None:
            return False
        self._file_image = cv2.cvtColor(bgr, cv2.COLOR_BGR2RGB)
        self._gray = None
        return True

    def _on_open_image(self) -> None:
        path = self.ask_open_filename(parent=self._dlg_parent(), title="Open an image", filetypes=[
            ("Images", "*.png *.jpg *.jpeg *.bmp *.tif *.tiff *.webp"), ("All files", "*.*")])
        if not path:
            return
        if not self._load_image_file(path):
            messagebox.showerror("Open image", f"Cannot read the image:\n{path}", parent=self._dlg_parent())
            return
        self._image_path = self._rel(path)
        self._show = None
        self._on_image_changed()
        self._refresh_all()

    def _snap(self, u: float, v: float):
        img = self._current_image()
        if img is None:
            return u, v
        if self._gray is None:
            a = np.asarray(img)
            if a.dtype != np.uint8:
                a = cv2.normalize(a.astype(np.float64), None, 0, 255, cv2.NORM_MINMAX).astype(np.uint8)
            self._gray = a if a.ndim == 2 else cv2.cvtColor(a[:, :, :3], cv2.COLOR_RGB2GRAY)
        h, w = self._gray.shape[:2]
        if not (6 <= u < w - 6 and 6 <= v < h - 6):
            return u, v
        pts = np.array([[[u, v]]], dtype=np.float32)
        try:
            cv2.cornerSubPix(self._gray, pts, (5, 5), (-1, -1),
                             (cv2.TERM_CRITERIA_EPS | cv2.TERM_CRITERIA_COUNT, 40, 0.01))
        except cv2.error:
            return u, v
        nu, nv = float(pts[0, 0, 0]), float(pts[0, 0, 1])
        return (nu, nv) if math.hypot(nu - u, nv - v) <= 8 else (u, v)

    def _on_image_click(self, u: float, v: float, hit) -> None:
        if self._show is not None:
            if self._mouse_var is not None:
                self._mouse_var.set("Board images are not edited by clicks -- press 'Points image' to edit points.")
            return
        size = self._image_size()
        if size and not (-0.5 <= u <= size[0] - 0.5 and -0.5 <= v <= size[1] - 0.5):
            return
        if hit is not None and self._mode_var.get() == "point":
            self._sheet and self._sheet.select([hit])             # click on a point = select its row
            self.on_selection_changed()
            return
        if self._snap_var.get():
            u, v = self._snap(u, v)
        sel = self._sheet.selected if self._sheet else []
        if self._mode_var.get() == "line":
            group = self._line_var.get().strip() or "L1"
            self._line_var.set(group)
            k = sum(1 for r in self.rows if r["line"] == group) + 1
            name = f"{group}_{k:02d}"
            self.rows.append(new_row(self._new_id(), name, u=u, v=v, line=group))
            self._changed()
            self._sheet and self._sheet.select([len(self.rows) - 1])
            return
        if len(sel) == 1:
            i = sel[0]
            self.rows[i]["u"], self.rows[i]["v"] = u, v
            nxt = next((j for j in range(i + 1, len(self.rows))
                        if not math.isfinite(self.rows[j]["u"]) and self.rows[j]["use"]), None)
            self._changed()
            if self._sheet is not None:
                self._sheet.select([nxt] if nxt is not None else [])
            return
        self.rows.append(new_row(self._new_id(), self._next_name(), u=u, v=v))
        self._changed()
        if self._sheet is not None:
            self._sheet.select([len(self.rows) - 1])

    def _on_image_drag(self, i: int, u: float, v: float, done: bool) -> None:
        if not (0 <= i < len(self.rows)) or self._show is not None:
            return
        if done and self._snap_var.get():
            u, v = self._snap(u, v)
        self.rows[i]["u"], self.rows[i]["v"] = u, v
        if done:
            self._changed()
            if self._sheet is not None:
                self._sheet.select([i])
        else:
            self._redraw_image()
            if self._mouse_var is not None:
                self._mouse_var.set(f"{self.rows[i]['name']}: u {u:.2f}  v {v:.2f}")

    def _on_image_motion(self, u: float, v: float) -> None:
        if self._mouse_var is not None:
            self._mouse_var.set(f"u {u:9.2f}   v {v:9.2f}")

    def _hit_point(self, x: float, y: float):
        iv = self._iv
        if iv is None or iv.view is None or self._show is not None:
            return None
        best, best_d = None, 8.0
        for i, r in enumerate(self.rows):
            if math.isfinite(r["u"]) and math.isfinite(r["v"]):
                X, Y = iv.to_screen(r["u"], r["v"])
                d = math.hypot(X - x, Y - y)
                if d < best_d:
                    best, best_d = i, d
        return best

    def _redraw_image(self) -> None:
        if self._iv is not None:
            self._iv.render(size=self._cb_size() if self._show is not None else self._image_size())

    def _gain(self) -> float:
        try:
            return max(0.0, float(self._gain_var.get()))
        except ValueError:
            return 10.0

    def _draw_axes(self, c, ts, rv, tv, K, D, length: float, label: str = "") -> None:
        """The object frame's X / Y / Z axes (red / green / blue) at its origin."""
        rv, tv = np.asarray(rv, float).reshape(3), np.asarray(tv, float).reshape(3)
        pts = np.array([[0, 0, 0], [length, 0, 0], [0, length, 0], [0, 0, length]], dtype=float)
        R, _ = cv2.Rodrigues(rv)
        if np.any((pts @ R.T + tv)[:, 2] <= 1e-9):
            return
        px = ce.project(pts, rv, tv, K, D)
        o = ts(*px[0])
        for k, (col, name) in enumerate((("#ff3b30", "X"), ("#34c759", "Y"), ("#0a84ff", "Z")), start=1):
            e = ts(*px[k])
            c.create_line(*o, *e, fill=col, width=3, arrow=tk.LAST, tags="ov")
            c.create_text(e[0] + 8, e[1], text=name, fill=col, font=("Arial", 9, "bold"), tags="ov")
        if label:
            c.create_text(o[0] - 8, o[1] - 8, text=label, anchor="se", fill="#ffffff",
                          font=("Arial", 9, "bold"), tags="ov")

    def _draw_board_overlay(self, iv: ImageView) -> None:
        c, ts = iv.canvas, iv.to_screen
        cand = next((x for x in (self._cb["cands"] if self._cb else []) if x["path"] == self._show), None)
        if cand is None:
            return
        pts = cand["corners"].reshape(-1, 2)
        n = len(pts)
        used = self._show in self._cb_use
        pose = self._res["cb_poses"].get(self._show) if self._res is not None else None
        self._draw_grids(c, ts, "board", pose)
        if pose is not None and self._show_proj_var.get():
            K, D = np.asarray(self._res["K"], float), np.asarray(self._res["D"], float)
            proj = ce.project(self._cand_obj(cand), pose[0], pose[1], K, D)
            gain = self._gain()
            for p, q in zip(pts, proj):
                c.create_line(*ts(*p), *ts(*(p + (q - p) * gain)), fill=PROJ_COLOR, width=2, tags="ov")
        if cand.get("support") is not None:
            draw_lattice(c, ts, cand["rc"], pts, cand["support"], cand.get("n_settings") or 1, used)
        else:
            for k, (x, y) in enumerate(pts):
                X, Y = ts(x, y)
                t = k / max(1, n - 1)
                col = f"#{int(255 * t):02x}{int(220 * (1 - t)):02x}30" if used else "#9a9a9a"
                r = 6 if k == 0 else 3
                c.create_oval(X - r, Y - r, X + r, Y + r, outline=col, width=2 if k == 0 else 1.5, tags="ov")
        if pose is not None:
            K, D = np.asarray(self._res["K"], float), np.asarray(self._res["D"], float)
            sq = self._square_size()
            self._draw_axes(c, ts, pose[0], pose[1], K, D, 3 * sq,
                            "WORLD FRAME" if self._world == self._show else "")
        msg = ("checked" if used else "NOT checked (not used)") + (
            "" if pose is not None else " -- no pose yet (Solve)")
        c.create_text(8, 8, anchor="nw", fill=SELECT_COLOR, font=("Arial", 9, "bold"), tags="ov",
                      text=f"Board image: {os.path.basename(self._show)} -- {msg}")

    def _draw_overlay(self, iv: ImageView) -> None:
        if self._show is not None:
            self._draw_board_overlay(iv)
            return
        c = iv.canvas
        ts = iv.to_screen
        gain = self._gain()
        ev, em, res = self._eval, self._eval_map, self._res
        sel = set(self._sheet.selected) if self._sheet else set()
        infos = self._infos if len(self._infos) == len(self.rows) else [classify_row(r) for r in self.rows]
        self._draw_grids(c, ts, "points", res.get("points_pose") if res is not None else None)

        if ev is not None and em is not None and em["pi"] is not None:
            pt_rows, line_rows, pi = em["pt_rows"], em["line_rows"], em["pi"]
            K, D = np.asarray(res["K"], float), np.asarray(res["D"], float)
            rv, tv = (np.asarray(a, float).reshape(3) for a in res["points_pose"])
            # straight lines: fitted curve + point-to-line distances
            if self._show_lines_var.get():
                for j, (name, idx) in enumerate(line_rows):
                    fit = ev.line_fits[pi][j]
                    if fit is None or len(fit["curve"]) < 2:
                        continue
                    c.create_line(*[q for p in fit["curve"] for q in ts(*p)], fill=CURVE_COLOR, width=1,
                                  dash=(6, 3), tags="ov")
                    x0, y0 = ts(*fit["curve"][-1])
                    c.create_text(x0 + 4, y0, text=name, anchor="w", fill=CURVE_COLOR,
                                  font=("Arial", 8, "bold"), tags="ov")
                    for k, i in enumerate(idx):
                        p = np.array([self.rows[i]["u"], self.rows[i]["v"]])
                        q = p + (fit["feet"][k] - p) * gain
                        c.create_line(*ts(*p), *ts(*q), fill="#ffd400", width=2, tags="ov")
            # '?' axes: the 3D line of possible positions of each '?' point
            if self._show_axes_var.get():
                objs = np.array([self._solved_obj(i) for i in pt_rows], dtype=float).reshape(-1, 3)
                spans = np.nanmax(objs, axis=0) - np.nanmin(objs, axis=0) if len(objs) > 1 else np.ones(3)
                ext = float(np.nanmax(spans)) if np.any(np.isfinite(spans)) else 1.0
                R, _ = cv2.Rodrigues(rv)
                for k, i in enumerate(pt_rows):
                    ax = infos[i]["axis"]
                    if ax is None or not np.all(np.isfinite(objs[k])):
                        continue
                    s = np.linspace(-0.5, 0.5, 41) * max(ext, 1e-6)
                    seg = np.repeat(objs[k][None], len(s), axis=0)
                    seg[:, ax] += s
                    ok = (seg @ R.T + tv)[:, 2] > 1e-9
                    if ok.sum() < 2:
                        continue
                    px = ce.project(seg[ok], rv, tv, K, D)
                    c.create_line(*[q for p in px for q in ts(*p)], fill=KIND_COLORS["unknown"], width=1,
                                  dash=(3, 4), tags="ov")
            # projections + magnified error vectors
            if self._show_proj_var.get():
                for k, i in enumerate(pt_rows):
                    pr = ev.projected[pi][k]
                    if not np.all(np.isfinite(pr)):
                        continue
                    p = np.array([self.rows[i]["u"], self.rows[i]["v"]])
                    q = p + (pr - p) * gain
                    c.create_line(*ts(*p), *ts(*q), fill=PROJ_COLOR, width=2, tags="ov")
                    X, Y = ts(*pr)
                    c.create_line(X - 5, Y, X + 6, Y, fill=PROJ_COLOR, tags="ov")
                    c.create_line(X, Y - 5, X, Y + 6, fill=PROJ_COLOR, tags="ov")

        names = self._show_names_var.get()
        for i, (r, info) in enumerate(zip(self.rows, infos)):
            if not (math.isfinite(r["u"]) and math.isfinite(r["v"])):
                continue
            X, Y = ts(r["u"], r["v"])
            col = KIND_COLORS[info["kind"]]
            if i in sel:
                continue
            c.create_oval(X - 4, Y - 4, X + 4, Y + 4, outline=col, width=2, tags="ov")
            if names:
                c.create_text(X + 6, Y - 6, text=r["name"], anchor="sw", fill=col, font=("Arial", 7), tags="ov")
        for i in sel:
            if i >= len(self.rows):
                continue
            r = self.rows[i]
            if not (math.isfinite(r["u"]) and math.isfinite(r["v"])):
                continue
            X, Y = ts(r["u"], r["v"])
            c.create_oval(X - 9, Y - 9, X + 9, Y + 9, outline=SELECT_COLOR, width=2, tags="ov")
            for a, b, cc, d in ((-14, 0, -4, 0), (4, 0, 14, 0), (0, -14, 0, -4), (0, 4, 0, 14)):
                c.create_line(X + a, Y + b, X + cc, Y + d, fill=SELECT_COLOR, tags="ov")
            c.create_text(X + 11, Y - 11, text=r["name"], anchor="sw", fill=SELECT_COLOR,
                          font=("Arial", 8, "bold"), tags="ov")
        if self._mode_var.get() == "line":
            c.create_text(8, 8, anchor="nw", fill=CURVE_COLOR, font=("Arial", 9, "bold"), tags="ov",
                          text=f"Line mode: clicks add points to line {self._line_var.get()}")
        elif sel and len(sel) == 1:
            c.create_text(8, 8, anchor="nw", fill=SELECT_COLOR, font=("Arial", 9, "bold"), tags="ov",
                          text=f"Click to set the image point of {self.rows[min(sel)]['name']}")

    # ── solving ────────────────────────────────────────────────────────────
    def _float_var(self, var, default):
        text = var.get().strip()
        if not text:
            return default
        return float(text)

    def _user_initial(self) -> dict:
        """Typed initial values (None = empty / automatic). Raises ValueError."""
        out = {}
        for n, v in self._init_vars.items():
            t = v.get().strip()
            out[n] = float(t) if t else None
        return out

    def _make_problem(self, views: list) -> ce.Problem:
        size = self._image_size() or (1000, 1000)
        model = self._model_var.get() if self._model_var.get() in ce.DIST_MODELS else "simple"
        fixed = {n for n, v in self._fix_vars.items() if v.get()}
        weights = {t: self._float_var(v, 1.0) for t, v in self._w_vars.items()}
        prob = ce.Problem(size, list(views), dist_model=model, fixed=fixed,
                          fix_aspect=bool(self._aspect_var.get()), weights=weights,
                          loss=self._loss_var.get(), huber_px=self._float_var(self._huber_var, 2.0))
        prob.K, prob.D = merge_initial(ce.guess_K(size), np.zeros(0), self._user_initial(), prob)
        return prob

    def _prepare_solve(self) -> dict:
        """The solve plan, or raises ValueError with a message for the user."""
        self._infos = [classify_row(r) for r in self.rows]
        size = self._image_size()
        if size is None:
            raise ValueError("The image size is unknown: connect / open the image, or type the size in the "
                             "Solve tab.")
        pts = None
        if self._use_points_var.get():
            view, pt_rows, line_rows, _issues = build_view(self.rows, self._infos)
            if len(view.obj) or view.lines:
                known = sum(1 for i in pt_rows if self._infos[i]["axis"] is None)
                if known < 4:
                    raise ValueError(f"At least 4 usable rows with fully known X, Y, Z (no '?') and an image "
                                     f"point are needed to find a starting pose of the points image; there are "
                                     f"{known}. (Or untick 'Use: point table'.)")
                pts = (view, pt_rows, line_rows)
        if self._use_cb_var.get() and self._cb:
            unknown = [c for c in self._cb["cands"] if c["path"] in self._cb_use and self._cand_axes(c) is None]
            if unknown:
                sx, sy = self._square_xy()
                raise ValueError(
                    f"{len(unknown)} checked board image(s) (marked 'X/Y?' in the Chessboard tab) cannot tell the "
                    f"board's X direction from its Y: square X ({sx:g}) and Y ({sy:g}) differ, and their corners "
                    f"would fit the {self._cb['cfg'].get('ncol')} x {self._cb['cfg'].get('nrow')} board either way "
                    f"round (or not at all). Uncheck them, make the board size (columns x rows) the FULL board, "
                    f"or use one square size.")
        cbs = self._cb_views() if self._use_cb_var.get() else []
        if cbs and self._cb_size() != tuple(size):
            raise ValueError(f"The chessboard images are {self._cb_size()[0]} x {self._cb_size()[1]} but the "
                             f"points image is {size[0]} x {size[1]}: they must come from the same camera "
                             f"setting (same size).")
        if pts is None and not cbs:
            raise ValueError("Nothing to solve: fill the Points table and / or detect and check chessboard "
                             "images (and tick them under 'Use').")
        try:
            prob = self._make_problem(([pts[0]] if pts else []) + [v for _p, v in cbs])
        except ValueError:
            raise ValueError("An initial value, weight or the huber threshold is not a number.")
        if prob.huber_px <= 0 or any(w < 0 for w in prob.weights.values()):
            raise ValueError("Weights must be 0 or more and the huber threshold more than 0.")
        return {"prob": prob, "pts": None if pts is None else True, "pt_rows": pts[1] if pts else [],
                "cb_paths": [p for p, _v in cbs], "mode": self._combine_var.get(),
                "user_init": self._user_initial()}

    def _schedule_route_text(self) -> None:
        if self._route_after is not None or self._route_var is None:
            return
        try:
            self._route_after = self.canvas.after(150, self._update_route_text)
        except tk.TclError:
            self._route_after = None

    def _update_route_text(self) -> None:
        """'Solver: ...' line under Solve: OpenCV only, or why the
        least-squares minimization is needed."""
        self._route_after = None
        if self._route_var is None:
            return
        color = "#7a4a00"
        try:
            plan = self._prepare_solve()
            plan["force_lsq"] = bool(self._force_lsq_var.get())
            route = ce.solver_route(plan)
            if plan["force_lsq"]:
                text = ("Least-squares minimization (started from OpenCV), because 'Always refine' is ticked"
                        + ("." if not route["reasons"] else "; also needed for: " + "; ".join(route["reasons"])))
            else:
                text = route["text"]
                if not route["reasons"]:
                    color = "#1b5e20"
        except ValueError as e:
            text, color = f"not ready -- {e}", "#777777"
        except Exception as e:                                   # never break the UI for a label
            text, color = f"? ({type(e).__name__}: {e})", "#777777"
        try:
            self._route_var.set("Solver: " + text)
            self._route_lbl.configure(fg=color)
        except tk.TclError:
            pass

    def _on_cancel_solve(self) -> None:
        st = self._solve_state
        if st is None or st["finished"]:
            return
        st["plan"]["cancel"].set()
        self._solve_state = None                       # its result, if it still comes, is ignored
        self._set_solving(False)
        if self._status_var is not None:
            self._status_var.set("Solve cancelled; the previous result is kept.")

    def _on_solve(self) -> None:
        if self._solve_state is not None and not self._solve_state["finished"]:
            return
        if self._sheet is not None:
            self._sheet.close_editor(commit=True)
        try:
            plan = self._prepare_solve()
        except ValueError as e:
            messagebox.showerror("Solve", str(e), parent=self._dlg_parent())
            return
        plan["force_lsq"] = bool(self._force_lsq_var.get())
        plan["cancel"] = threading.Event()
        state = {"finished": False, "result": None, "error": None, "rows": [dict(r) for r in self.rows],
                 "plan": plan, "use": set(self._cb_use)}

        def work():
            try:
                state["result"] = run_plan(plan)
            except ce.Cancelled:
                state["error"] = "cancelled"
            except Exception as e:                                  # reported in the UI
                state["error"] = f"{e}" if isinstance(e, ValueError) else f"{type(e).__name__}: {e}"
            finally:
                state["finished"] = True
        self._solve_state = state
        self._set_solving(True)
        threading.Thread(target=work, daemon=True).start()
        self._poll_solve()

    def _poll_solve(self) -> None:
        st = self._solve_state
        if st is None:
            return
        if not st["finished"]:
            try:
                self.canvas.after(40, self._poll_solve)
            except tk.TclError:
                pass
            return
        self._set_solving(False)
        if st["error"] == "cancelled":
            return
        if st["error"]:
            if self._status_var is not None:
                self._status_var.set("Failed: " + st["error"])
            messagebox.showerror("Solve", st["error"], parent=self._dlg_parent())
            return
        self._apply_result(st["result"], st["plan"], st["rows"], st["use"])

    def _solve_sync(self) -> None:
        try:
            plan = self._prepare_solve()
            plan["force_lsq"] = bool(self._force_lsq_var.get())
            self._apply_result(run_plan(plan), plan, [dict(r) for r in self.rows], set(self._cb_use), push=False)
        except Exception as e:
            if self._status_var is not None:
                self._status_var.set(f"Automatic solve failed: {e}")

    def _apply_result(self, out: dict, plan: dict, rows_at_solve: list, use_at_solve: set,
                      push: bool = True) -> None:
        prob = plan["prob"]
        unknowns = {}
        for k, i in enumerate(plan["pt_rows"]):
            if k in out["unknowns"] and i < len(rows_at_solve):
                info = classify_row(rows_at_solve[i])
                unknowns[rows_at_solve[i]["id"]] = (info["axis"], float(out["unknowns"][k]))
        lay_free = [n for n in ce.INTR_NAMES + ce.DIST_NAMES[:ce.DIST_MODELS[prob.dist_model]]
                    if n not in prob.fixed and not (prob.fix_aspect and n == "fy")]
        changed = [dict(r) for r in self.rows] != rows_at_solve or set(self._cb_use) != use_at_solve
        pp = out["points_pose"]
        route = out.get("route") or {}
        report = (f"method: {route.get('text', '?')}\n"
                  f"solver: {'converged' if out['success'] else 'NOT converged'}"
                  f"{(' -- ' + out['message']) if out['message'] else ''}; "
                  + (f"{out['nfev']} least-squares evaluations, " if out["nfev"] else "")
                  + f"{out['seconds'] * 1000:.0f} ms; "
                  f"{len(lay_free)} camera parameters solved: {', '.join(lay_free) or 'none'}")
        if out["report"]:
            report += "\n" + out["report"].replace("; ", "\n")
        self._res = {"K": np.asarray(out["K"]).tolist(), "D": np.asarray(out["D"]).ravel().tolist(),
                     "points_pose": None if pp is None else [np.ravel(pp[0]).tolist(), np.ravel(pp[1]).tolist()],
                     "cb_poses": {p: [np.ravel(rv).tolist(), np.ravel(tv).tolist()]
                                  for p, (rv, tv) in out["cb_poses"].items()},
                     "unknowns": unknowns, "free": lay_free, "stale": changed, "report": report,
                     "adjusted": False, "solved_rms": float(out["rms"])}
        self._res["solved"] = {"K": [list(r) for r in self._res["K"]], "D": list(self._res["D"]),
                               "points_pose": None if pp is None else [list(a) for a in self._res["points_pose"]],
                               "cb_poses": {p: [list(a) for a in v] for p, v in self._res["cb_poses"].items()}}
        if self._world != "points" and self._world not in self._res["cb_poses"]:
            self._world = "points"
        if self._status_var is not None:
            self._status_var.set(f"Done: RMS {out['rms']:.3f} px in {out['seconds'] * 1000:.0f} ms"
                                 + ("" if out["success"] else " (NOT converged)"))
        self._refresh_all()
        self._add_run(plan, out)
        self._refresh_right()
        self._sync_sliders()
        if push:
            self._push()

    def _set_solving(self, running: bool) -> None:
        if self._solve_btn is not None:
            self._solve_btn.configure(state="disabled" if running else "normal",
                                      text="Solving..." if running else "Solve")
        if self._cancel_solve_btn is not None:
            self._cancel_solve_btn.configure(state="normal" if running else "disabled")
        if running and self._status_var is not None:
            self._status_var.set("Solving ...")

    # ── manual adjustment (Adjust tab) ─────────────────────────────────────
    def _adj_param_names(self) -> list:
        nd = ce.DIST_MODELS.get(self._model_var.get(), 5)
        if self._res is not None:
            nd = max(nd, len(self._res["D"]))
        return list(ce.INTR_NAMES) + list(ce.DIST_NAMES[:nd]) + list(POSE_PARAMS)

    def _adj_pose_key(self):
        """Key of the pose the pose sliders move: the world-frame view."""
        res = self._res
        if res is None:
            return None
        if self._world != "points" and self._world in res["cb_poses"]:
            return self._world
        if res.get("points_pose") is not None:
            return "points"
        return next((c["path"] for c in (self._cb["cands"] if self._cb else []) if c["path"] in res["cb_poses"]),
                    None)

    def _param_get(self, name: str):
        res = self._res
        if res is None:
            return None
        if name in ce.INTR_NAMES:
            r, c = {"fx": (0, 0), "fy": (1, 1), "cx": (0, 2), "cy": (1, 2)}[name]
            return float(res["K"][r][c])
        if name in ce.DIST_NAMES:
            j = ce.DIST_NAMES.index(name)
            return float(res["D"][j]) if j < len(res["D"]) else 0.0
        key = self._adj_pose_key()
        if key is None:
            return None
        pose = res["points_pose"] if key == "points" else res["cb_poses"][key]
        k = POSE_PARAMS.index(name)
        return float(pose[k // 3][k % 3])

    def _param_set(self, name: str, value: float) -> None:
        res = self._res
        if res is None:
            return
        if name in ce.INTR_NAMES:
            r, c = {"fx": (0, 0), "fy": (1, 1), "cx": (0, 2), "cy": (1, 2)}[name]
            res["K"][r][c] = float(value)
        elif name in ce.DIST_NAMES:
            j = ce.DIST_NAMES.index(name)
            while len(res["D"]) <= j:
                res["D"].append(0.0)
            res["D"][j] = float(value)
        else:
            key = self._adj_pose_key()
            if key is None:
                return
            pose = res["points_pose"] if key == "points" else res["cb_poses"][key]
            k = POSE_PARAMS.index(name)
            pose[k // 3][k % 3] = float(value)
        res["adjusted"] = True

    def _default_bounds(self, name: str, v: float):
        size = self._image_size() or (1000, 1000)
        if name in ("fx", "fy"):
            d = 0.2 * abs(v) or 100.0
        elif name == "cx":
            d = 0.1 * size[0]
        elif name == "cy":
            d = 0.1 * size[1]
        elif name in ("k1", "k2", "k3", "k4", "k5", "k6"):
            d = max(0.3, abs(v))
        elif name in ("taux", "tauy"):
            d = 0.05
        elif name in ce.DIST_NAMES:
            d = max(0.01, 2 * abs(v))
        elif name.startswith("rvec"):
            d = 0.2
        else:
            key = self._adj_pose_key()
            pose = None if key is None else (self._res["points_pose"] if key == "points" else self._res["cb_poses"][key])
            d = max(0.2 * float(np.linalg.norm(pose[1])), 1.0) if pose is not None else 100.0
        return v - d, v + d

    def _build_adjust_tab(self, tab) -> None:
        tk.Label(tab, font=("Arial", 8), fg="#4a5568", justify="left", anchor="w", wraplength=660, text=(
            "Move a parameter by hand and see at once how the errors (table, image, summary) change. "
            "The sliders change the CURRENT RESULT, so the output pins follow them; 'Revert' gives back the "
            "last solved values. Pose sliders move the world-frame view (rvec in radians).")).pack(fill="x")
        bar = tk.Frame(tab)
        bar.pack(fill="x", pady=(4, 2))
        for text, cmd, tip in (
                ("Add slider", self._on_add_slider, "Add a slider for one more parameter."),
                ("Revert to solved values", self._on_revert_adjust,
                 "Undo all manual changes: back to the values of the last Solve."),
                ("Adjusted → initial", self._on_result_to_initial,
                 "Copy the (adjusted) camera matrix / distortion into the Solve tab's initial values, e.g. to "
                 "solve again from here with some parameters fixed.")):
            b = tk.Button(bar, text=text, font=("Arial", 8), command=cmd)
            b.pack(side="left", padx=(0, 4))
            add_tooltip(b, tip)
        self._adj_frame = tk.Frame(tab)
        self._adj_frame.pack(fill="x")
        self._adj_info_var = tk.StringVar()
        tk.Label(tab, textvariable=self._adj_info_var, font=("Consolas", 9), fg="#14532d", justify="left",
                 anchor="w").pack(fill="x", pady=(6, 0))
        self._build_adj_rows()

    def _build_adj_rows(self) -> None:
        f = self._adj_frame
        if f is None:
            return
        for w in f.winfo_children():
            w.destroy()
        self._adj_widgets = []
        names = self._adj_param_names()
        for j, text in enumerate(("parameter", "lower", "", "upper", "value")):
            tk.Label(f, text=text, font=("Arial", 8, "bold"), anchor="w").grid(row=0, column=j, sticky="w")
        for k, spec in enumerate(self._adj_rows):
            row = f
            r = k + 1
            pvar = tk.StringVar(value=spec["param"])
            cb = ttk.Combobox(row, textvariable=pvar, values=names, state="readonly", width=9)
            cb.grid(row=r, column=0, padx=2, pady=1, sticky="w")
            add_tooltip(cb, "Parameter moved by this slider. Pose: rvec x/y/z (rotation vector, radians) and "
                            "tvec x/y/z of the world-frame view.")
            lo_var, hi_var = tk.StringVar(value=spec["lo"]), tk.StringVar(value=spec["hi"])
            lo = tk.Entry(row, textvariable=lo_var, width=10, justify="right")
            lo.grid(row=r, column=1, padx=2, pady=1, sticky="w")
            add_tooltip(lo, "Lower end of the slider (Enter to apply).")
            val = tk.DoubleVar(value=0.0)
            sc = tk.Scale(row, orient=tk.HORIZONTAL, length=280, showvalue=False, variable=val, from_=0, to=1,
                          resolution=0.001)
            sc.grid(row=r, column=2, padx=2, pady=1, sticky="w")
            add_tooltip(sc, "Drag: the errors and drawings update immediately; downstream nodes update when you "
                            "release the mouse. Arrow keys after a click move it in small steps.")
            hi = tk.Entry(row, textvariable=hi_var, width=10, justify="right")
            hi.grid(row=r, column=3, padx=2, pady=1, sticky="w")
            add_tooltip(hi, "Upper end of the slider (Enter to apply).")
            ent_var = tk.StringVar()
            ent = tk.Entry(row, textvariable=ent_var, width=13, justify="right")
            ent.grid(row=r, column=4, padx=2, pady=1, sticky="w")
            add_tooltip(ent, "Exact value: type it and press Enter (the slider range widens if needed).")
            rm = tk.Button(row, text="✕", font=("Arial", 7), command=lambda kk=k: self._on_remove_slider(kk))
            rm.grid(row=r, column=5, padx=2, pady=1, sticky="w")
            add_tooltip(rm, "Remove this slider (the value stays).")
            w = {"pvar": pvar, "lo": lo_var, "hi": hi_var, "val": val, "scale": sc, "entry": ent_var}
            self._adj_widgets.append(w)
            cb.bind("<<ComboboxSelected>>", lambda _e, kk=k: self._on_slider_param(kk))
            for e in (lo, hi):
                e.bind("<Return>", lambda _e, kk=k: self._on_slider_bounds(kk))
                e.bind("<FocusOut>", lambda _e, kk=k: self._on_slider_bounds(kk))
            ent.bind("<Return>", lambda _e, kk=k: self._on_slider_entry(kk))
            sc.configure(command=lambda _v, kk=k: self._on_slider_move(kk))
            sc.bind("<ButtonRelease-1>", lambda _e: self._push())
            sc.bind("<KeyRelease>", lambda _e: self._push())
        self._sync_sliders()

    def _configure_scale(self, k: int) -> None:
        spec, w = self._adj_rows[k], self._adj_widgets[k]
        try:
            lo, hi = float(spec["lo"]), float(spec["hi"])
        except ValueError:
            return
        if hi <= lo:
            hi = lo + 1e-9
        w["scale"].configure(from_=lo, to=hi, resolution=(hi - lo) / 2000.0)

    def _sync_sliders(self) -> None:
        """Slider positions / value boxes from the current result."""
        self._adj_syncing = True
        try:
            for k, w in enumerate(self._adj_widgets):
                spec = self._adj_rows[k]
                v = self._param_get(spec["param"])
                if v is None:
                    w["entry"].set("")
                    continue
                try:
                    lo, hi = float(spec["lo"]), float(spec["hi"])
                except ValueError:
                    lo, hi = self._default_bounds(spec["param"], v)
                if not (lo <= v <= hi) or hi <= lo:
                    lo, hi = self._default_bounds(spec["param"], v)
                spec["lo"], spec["hi"] = f"{lo:.6g}", f"{hi:.6g}"
                w["lo"].set(spec["lo"])
                w["hi"].set(spec["hi"])
                self._configure_scale(k)
                w["val"].set(v)
                w["entry"].set(f"{v:.8g}")
        finally:
            self._adj_syncing = False
        self._update_adj_info()

    def _on_add_slider(self) -> None:
        if self._res is None:
            messagebox.showinfo("Adjust", "Solve first: the sliders move the parameters of a result.",
                                parent=self._dlg_parent())
            return
        used = {s["param"] for s in self._adj_rows}
        name = next((n for n in ("fx", "k1", "k2", "cx", "cy", "p1", "p2", "fy") + POSE_PARAMS if n not in used),
                    "fx")
        self._adj_rows.append({"param": name, "lo": "", "hi": ""})
        self._build_adj_rows()

    def _on_remove_slider(self, k: int) -> None:
        if 0 <= k < len(self._adj_rows):
            del self._adj_rows[k]
            self._build_adj_rows()

    def _on_slider_param(self, k: int) -> None:
        self._adj_rows[k] = {"param": self._adj_widgets[k]["pvar"].get(), "lo": "", "hi": ""}
        self._sync_sliders()

    def _on_slider_bounds(self, k: int) -> None:
        w, spec = self._adj_widgets[k], self._adj_rows[k]
        try:
            lo, hi = float(w["lo"].get()), float(w["hi"].get())
            if hi <= lo:
                raise ValueError
        except ValueError:
            w["lo"].set(spec["lo"])
            w["hi"].set(spec["hi"])
            return
        spec["lo"], spec["hi"] = w["lo"].get().strip(), w["hi"].get().strip()
        v = self._param_get(spec["param"])
        self._configure_scale(k)
        if v is not None and not (lo <= v <= hi):
            self._on_slider_value(k, min(max(v, lo), hi), push=True)
        elif v is not None:
            self._adj_syncing = True
            w["val"].set(v)
            self._adj_syncing = False

    def _on_slider_entry(self, k: int) -> None:
        w, spec = self._adj_widgets[k], self._adj_rows[k]
        try:
            v = float(w["entry"].get())
        except ValueError:
            return
        lo, hi = float(spec["lo"]), float(spec["hi"])
        if not (lo <= v <= hi):
            spec["lo"], spec["hi"] = (f"{x:.6g}" for x in self._default_bounds(spec["param"], v))
            w["lo"].set(spec["lo"])
            w["hi"].set(spec["hi"])
            self._configure_scale(k)
        self._on_slider_value(k, v, push=True)

    def _on_slider_move(self, k: int) -> None:
        if getattr(self, "_adj_syncing", False) or self._res is None:
            return
        self._on_slider_value(k, float(self._adj_widgets[k]["val"].get()), push=False)

    def _on_slider_value(self, k: int, v: float, push: bool) -> None:
        self._param_set(self._adj_rows[k]["param"], v)
        w = self._adj_widgets[k]
        self._adj_syncing = True
        try:
            w["val"].set(v)
            w["entry"].set(f"{v:.8g}")
            # another slider may show the same parameter
            for j, ww in enumerate(self._adj_widgets):
                if j != k and self._adj_rows[j]["param"] == self._adj_rows[k]["param"]:
                    ww["val"].set(v)
                    ww["entry"].set(f"{v:.8g}")
        finally:
            self._adj_syncing = False
        self._schedule_live()
        if push:
            self._push()

    def _schedule_live(self) -> None:
        """Re-evaluate at most once per idle period while a slider is dragged."""
        if getattr(self, "_live_pending", False):
            return
        self._live_pending = True

        def run():
            self._live_pending = False
            self._update_row_errors()
            if self._sheet is not None:
                self._sheet.redraw()
            self._refresh_result_values()
            self._refresh_summary()
            self._refresh_cb_tree()
            self._update_image_source_label()
            self._redraw_image()
            self._update_face()
            self._update_adj_info()
            self._refresh_right()
        try:
            self.canvas.after_idle(run)
        except tk.TclError:
            self._live_pending = False

    def _update_adj_info(self) -> None:
        if getattr(self, "_adj_info_var", None) is None:
            return
        ev, res = self._eval, self._res
        if res is None or ev is None:
            self._adj_info_var.set("No result yet: Solve first.")
            return
        key = self._adj_pose_key()
        pose_txt = ("points image" if key == "points" else os.path.basename(key)) if key else "none"
        solved = res.get("solved_rms")
        parts = [f"RMS all {ev.rms.get('all', math.nan):.4f} px"]
        parts += [f"{t} {v:.4f}" for t, v in ev.rms.items() if t != "all"]
        text = "   ".join(parts) + f"\nweighted cost {ev.cost:.6g}"
        if solved is not None:
            text += f"   (at the last Solve: RMS all {solved:.4f} px)"
        text += f"\npose sliders move: {pose_txt}" + ("\nMANUALLY ADJUSTED" if res.get("adjusted") else "")
        self._adj_info_var.set(text)

    def _on_revert_adjust(self) -> None:
        res = self._res
        if res is None or not res.get("solved"):
            return
        snap = res["solved"]
        res["K"] = [list(r) for r in snap["K"]]
        res["D"] = list(snap["D"])
        res["points_pose"] = None if snap["points_pose"] is None else [list(a) for a in snap["points_pose"]]
        res["cb_poses"] = {p: [list(a) for a in v] for p, v in snap["cb_poses"].items()}
        res["adjusted"] = False
        self._refresh_all()
        self._sync_sliders()
        self._push()

    # ── 3D grids (Grids tab) ───────────────────────────────────────────────
    def _build_grids_tab(self, tab) -> None:
        tk.Label(tab, font=("Arial", 8), fg="#4a5568", justify="left", anchor="w", wraplength=660, text=(
            "3D grids projected into the image with the current camera (distortion included). A grid is drawn "
            "in the OBJECT coordinates of the shown image: 'points' grids on the points image (the X / Y / Z "
            "of the table), 'board' grids on the board images (board coordinates). Values: a list "
            "'0 50 120' or start:step:stop like '0:100:1000'. One value = a plane, e.g. Z = '0'.")).pack(fill="x")
        bar = tk.Frame(tab)
        bar.pack(fill="x", pady=(4, 2))
        b = tk.Button(bar, text="Add grid", font=("Arial", 8), command=self._on_add_grid)
        b.pack(side="left")
        add_tooltip(b, "Add one more grid (up to 12).")
        cbx = tk.Checkbutton(bar, text="show grids", variable=self._grids_on_var, font=("Arial", 8),
                             command=self._redraw_image)
        cbx.pack(side="left", padx=(8, 0))
        add_tooltip(cbx, "Draw (or hide) all grids.")
        lbl = tk.Label(bar, text="  line width", font=("Arial", 8))
        lbl.pack(side="left")
        sp = tk.Spinbox(bar, from_=1, to=5, width=3, textvariable=self._grid_width_var, command=self._redraw_image)
        sp.pack(side="left")
        add_tooltip(lbl, "Width of the grid lines on the screen (pixels).")
        add_tooltip(sp, "Width of the grid lines on the screen (pixels).")
        self._grid_frame = tk.Frame(tab)
        self._grid_frame.pack(fill="x")
        self._grid_msg_var = tk.StringVar()
        tk.Label(tab, textvariable=self._grid_msg_var, font=("Arial", 8), fg="#b00020", anchor="w",
                 justify="left", wraplength=660).pack(fill="x", pady=(4, 0))
        self._build_grid_rows()

    def _build_grid_rows(self) -> None:
        f = self._grid_frame
        if f is None:
            return
        for w in f.winfo_children():
            w.destroy()
        for j, text in enumerate(("on", "name", "X values", "Y values", "Z values", "frame", "color")):
            tk.Label(f, text=text, font=("Arial", 8, "bold"), anchor="w").grid(row=0, column=j, sticky="w")
        for k, g in enumerate(self._grids):
            row = f
            r = k + 1
            col = 0
            vars_ = {}
            on = tk.BooleanVar(value=g["on"])
            tk.Checkbutton(row, variable=on, command=lambda kk=k: self._on_grid_edit(kk)).grid(row=r, column=0)
            vars_["on"] = on
            for key, width, tip in (("name", 9, "Name (only for you)."),
                                    ("xs", 17, "X values of the grid lines, e.g. '0:100:1000' (start:step:stop) "
                                               "or '0 250 500'."),
                                    ("ys", 17, "Y values, same syntax."),
                                    ("zs", 17, "Z values, same syntax. One value = a plane at that Z.")):
                var = tk.StringVar(value=g[key])
                e = tk.Entry(row, textvariable=var, width=width)
                col += 1
                e.grid(row=r, column=col, padx=(0, 3), pady=1)
                e.bind("<Return>", lambda _e, kk=k: self._on_grid_edit(kk))
                e.bind("<FocusOut>", lambda _e, kk=k: self._on_grid_edit(kk))
                add_tooltip(e, tip)
                vars_[key] = var
            for key, values, width, tip in (
                    ("frame", ("points", "board"), 7, "points: drawn on the points image in its object coordinates; "
                                                     "board: drawn on every board image in board coordinates."),
                    ("color", tuple(GRID_COLORS), 8, "Line colour.")):
                var = tk.StringVar(value=g[key])
                cb = ttk.Combobox(row, textvariable=var, values=values, state="readonly", width=width)
                col += 1
                cb.grid(row=r, column=col, padx=(0, 3), pady=1)
                cb.bind("<<ComboboxSelected>>", lambda _e, kk=k: self._on_grid_edit(kk))
                add_tooltip(cb, tip)
                vars_[key] = var
            rm = tk.Button(row, text="✕", font=("Arial", 7), command=lambda kk=k: self._on_remove_grid(kk))
            rm.grid(row=r, column=col + 1)
            add_tooltip(rm, "Remove this grid.")
            g["_vars"] = vars_
        self._redraw_image()

    def _on_add_grid(self) -> None:
        if len(self._grids) >= 12:
            return
        n = len(self._grids) + 1
        colors = list(GRID_COLORS)
        self._grids.append({"on": True, "name": f"G{n}", "xs": "0:100:500", "ys": "0:100:500", "zs": "0",
                            "frame": "board" if (self._show is not None or self._res is None or
                                                 self._res.get("points_pose") is None) else "points",
                            "color": colors[(n - 1) % len(colors)]})
        self._build_grid_rows()

    def _on_remove_grid(self, k: int) -> None:
        if 0 <= k < len(self._grids):
            del self._grids[k]
            self._build_grid_rows()

    def _on_grid_edit(self, k: int) -> None:
        if not (0 <= k < len(self._grids)):
            return
        g = self._grids[k]
        for key, var in g.get("_vars", {}).items():
            g[key] = bool(var.get()) if key == "on" else var.get().strip()
        self._redraw_image()

    def _grid_lines(self, g: dict, pose, K, D) -> list:
        xs, ys, zs = (cbs.parse_number_list(g[k]) for k in ("xs", "ys", "zs"))
        for axis, key, vals in (("X", "xs", xs), ("Y", "ys", ys), ("Z", "zs", zs)):
            if vals is None:
                raise ValueError(f"grid {g['name']}: {axis} values '{g[key]}' cannot be read -- use a list "
                                 f"'0 50 120' or start:step:stop '0:100:1000'")
        n_lines = len(ys) * len(zs) * (len(xs) > 1) + len(xs) * len(zs) * (len(ys) > 1) + \
            len(xs) * len(ys) * (len(zs) > 1)
        if n_lines > 3000:
            raise ValueError(f"grid {g['name']}: {n_lines} lines are too many (at most 3000)")
        key = (g["xs"], g["ys"], g["zs"], tuple(np.ravel(pose[0])), tuple(np.ravel(pose[1])),
               tuple(np.ravel(K)), tuple(np.ravel(D)))
        hit = self._grid_cache.get(key)
        if hit is None:
            hit = ce.project_grid(xs, ys, zs, pose[0], pose[1], K, D)
            if len(self._grid_cache) > 64:
                self._grid_cache.clear()
            self._grid_cache[key] = hit
        return hit

    def _draw_grids(self, c, ts, frame: str, pose) -> None:
        if not self._grids_on_var.get() or self._res is None or pose is None:
            if getattr(self, "_grid_msg_var", None) is not None:
                self._grid_msg_var.set("" if self._res is not None or not self._grids else
                                       "Grids are drawn after a Solve (they need the camera).")
            return
        K, D = np.asarray(self._res["K"], float), np.asarray(self._res["D"], float)
        try:
            width = max(1, min(5, int(self._grid_width_var.get())))
        except (ValueError, tk.TclError):
            width = 1
        msgs = []
        for g in self._grids:
            if not g["on"]:
                continue
            if g["frame"] != frame:                    # not drawn here, but report its errors too
                bad = [a for a, k in (("X", "xs"), ("Y", "ys"), ("Z", "zs")) if cbs.parse_number_list(g[k]) is None]
                if bad:
                    msgs.append(f"grid {g['name']}: {bad[0]} values '{g[bad[0].lower() + 's']}' cannot be read "
                                f"-- use a list '0 50 120' or start:step:stop '0:100:1000'")
                continue
            try:
                lines = self._grid_lines(g, pose, K, D)
            except ValueError as e:
                msgs.append(str(e))
                continue
            for pl in lines:
                if len(pl) >= 2 and np.all(np.isfinite(pl)):
                    c.create_line(*[q for p in pl for q in ts(*p)], fill=GRID_COLORS.get(g["color"], "#00e5ff"),
                                  width=width, tags="ov")
        if getattr(self, "_grid_msg_var", None) is not None:
            self._grid_msg_var.set("\n".join(msgs))

    # ── right-side views ───────────────────────────────────────────────────
    def _right_tab(self) -> str:
        nb = self._rnb
        if nb is None:
            return ""
        try:
            return nb.tab(nb.select(), "text")
        except tk.TclError:
            return ""

    def _refresh_right(self) -> None:
        tab = self._right_tab()
        if tab == "Errors":
            self._draw_errors()
        elif tab == "3D":
            self._update_3d()
        elif tab == "Undistort":
            self._update_undistort()
        elif tab == "Runs":
            self._refresh_runs_tree()

    # Errors
    def _build_errors_tab(self, tab) -> None:
        tk.Label(tab, font=("Arial", 8), fg="#4a5568", anchor="w", justify="left", wraplength=760, text=(
            "Top: error of every used item (table point / line point / board image RMS), in table order; the "
            "dashed line is the RMS of all. Click a bar to select that row or board image. Bottom: every "
            "residual vector (projected - measured, px) around (0, 0); circles = 1 x and 3 x the RMS.")).pack(fill="x")
        self._err_canvas = tk.Canvas(tab, bg="#ffffff", highlightthickness=0)
        self._err_canvas.pack(fill="both", expand=True)
        self._err_canvas.bind("<Configure>", lambda _e: self._draw_errors())
        self._err_canvas.bind("<ButtonRelease-1>", self._on_err_click)
        self._err_canvas.bind("<Motion>", self._on_err_motion)
        self._err_hover_var = tk.StringVar()
        tk.Label(tab, textvariable=self._err_hover_var, font=("Consolas", 8), fg="#555555", anchor="w").pack(fill="x")

    def _error_items(self) -> list:
        """[(label, error px, colour, ("row", i) | ("board", path))] of the current evaluation."""
        ev, em = self._eval, self._eval_map
        if ev is None or em is None:
            return []
        items = []
        if em["pi"] is not None:
            done = set()
            for k, i in enumerate(em["pt_rows"]):
                e = float(np.linalg.norm(ev.per_point[em["pi"]][k]))
                items.append((i, self.rows[i]["name"], e, KIND_COLORS[self._infos[i]["kind"]], ("row", i)))
                done.add(i)
            for j, (_n, idx) in enumerate(em["line_rows"]):
                fit = ev.line_fits[em["pi"]][j]
                if fit is None:
                    continue
                for k, i in enumerate(idx):
                    if i not in done:
                        items.append((i, self.rows[i]["name"], abs(float(fit["residual"][k])), KIND_COLORS["line"],
                                      ("row", i)))
            items.sort(key=lambda t: t[0])
        out = [t[1:] for t in items]
        for p in em["cb_paths"]:
            out.append((os.path.basename(p), self._cb_err.get(p, math.nan), "#8e5cd9", ("board", p)))
        return out

    def _draw_errors(self) -> None:
        c = self._err_canvas
        if c is None:
            return
        try:
            w, h = c.winfo_width(), c.winfo_height()
        except tk.TclError:
            return
        if w < 60 or h < 60:
            return
        c.delete("all")
        items = self._error_items()
        self._err_bars = []
        if not items:
            c.create_text(w / 2, h / 2, text="No result yet: Solve first.", fill="#888888")
            return
        ev = self._eval
        rms = ev.rms.get("all", math.nan)
        # bars
        top, bot, left, right = 14, int(h * 0.52), 46, w - 10
        vals = np.array([it[1] for it in items], dtype=float)
        vmax = float(np.nanmax(vals)) * 1.1 if np.any(np.isfinite(vals)) and np.nanmax(vals) > 0 else 1.0
        bw = (right - left) / len(items)
        c.create_line(left, bot, right, bot, fill="#888888")
        c.create_line(left, top, left, bot, fill="#888888")
        for frac in (0.0, 0.5, 1.0):
            y = bot - frac * (bot - top)
            c.create_text(left - 4, y, text=f"{frac * vmax:.3g}", anchor="e", font=("Arial", 7), fill="#555555")
        c.create_text(left, top - 8, text="error (px)", anchor="w", font=("Arial", 7), fill="#555555")
        for k, (label, v, col, ref) in enumerate(items):
            x0 = left + k * bw
            y = bot - (v / vmax) * (bot - top) if math.isfinite(v) else bot
            c.create_rectangle(x0 + 0.1 * bw, y, x0 + 0.9 * bw, bot, fill=col, outline="")
            self._err_bars.append((x0, x0 + bw, label, v, ref))
            if bw >= 14:
                c.create_text(x0 + bw / 2, bot + 3, text=label[:10], anchor="ne", angle=60, font=("Arial", 7),
                              fill="#444444")
        if math.isfinite(rms):
            y = bot - (rms / vmax) * (bot - top)
            c.create_line(left, y, right, y, fill="#d00000", dash=(4, 3))
            c.create_text(right, y - 2, text=f"RMS all {rms:.3f}", anchor="se", font=("Arial", 7), fill="#d00000")
        # residual scatter
        res = []
        em = self._eval_map
        if em["pi"] is not None:
            r = np.asarray(ev.per_point[em["pi"]], float).reshape(-1, 2)
            res.append((r[np.all(np.isfinite(r), axis=1)], KIND_COLORS["point"]))
        for j, _p in enumerate(em["cb_paths"]):
            r = np.asarray(ev.per_point[j + (em["pi"] is not None)], float).reshape(-1, 2)
            res.append((r, "#8e5cd9"))
        allr = np.vstack([r for r, _c in res]) if res else np.zeros((0, 2))
        cy0 = bot + (h - bot) / 2 + 20
        size = min(w - 20, h - bot - 50) / 2
        if len(allr) and size > 20:
            lim = max(float(np.max(np.abs(allr))) * 1.1, 3 * rms if math.isfinite(rms) else 1e-6, 1e-6)
            s = size / lim
            cx0 = w / 2
            c.create_line(cx0 - size, cy0, cx0 + size, cy0, fill="#cccccc")
            c.create_line(cx0, cy0 - size, cx0, cy0 + size, fill="#cccccc")
            for f, dash in ((1, None), (3, (3, 3))):
                if math.isfinite(rms):
                    rr = f * rms * s
                    c.create_oval(cx0 - rr, cy0 - rr, cx0 + rr, cy0 + rr, outline="#d00000", dash=dash)
            for r, col in res:
                for dx, dy in r:
                    X, Y = cx0 + dx * s, cy0 + dy * s
                    c.create_oval(X - 1.5, Y - 1.5, X + 1.5, Y + 1.5, fill=col, outline="")
            c.create_text(cx0 + size, cy0 + 2, text=f"{lim:.3g} px", anchor="ne", font=("Arial", 7), fill="#555555")
            c.create_text(8, h - 6, anchor="sw", font=("Arial", 7), fill="#555555",
                          text="residual dx, dy (px)\ngreen = table points\nviolet = board corners")

    def _err_bar_at(self, x, y):
        for x0, x1, label, v, ref in getattr(self, "_err_bars", []):
            if x0 <= x < x1 and y <= self._err_canvas.winfo_height() * 0.6:
                return label, v, ref
        return None

    def _on_err_motion(self, e) -> None:
        hit = self._err_bar_at(e.x, e.y)
        self._err_hover_var.set("" if hit is None else f"{hit[0]}: {hit[1]:.4f} px  (click to select)")

    def _on_err_click(self, e) -> None:
        hit = self._err_bar_at(e.x, e.y)
        if hit is None:
            return
        kind, key = hit[2]
        if kind == "row":
            self._show_board(None)
            if self._sheet is not None:
                self._sheet.select([key])
            self.on_selection_changed()
        else:
            self._show_board(key)
            self._select_tree_path(key)

    # 3D
    def _build_3d_tab(self, tab) -> None:
        bar = tk.Frame(tab)
        bar.pack(fill="x")
        lbl = tk.Label(bar, text="Show", font=("Arial", 8))
        lbl.pack(side="left")
        cb = ttk.Combobox(bar, textvariable=self._v3_mode_var, state="readonly", width=24,
                          values=("points image (object frame)", "board frame (all boards)", "camera frame"))
        cb.pack(side="left", padx=(4, 8))
        cb.bind("<<ComboboxSelected>>", lambda _e: self._update_3d(refit=True))
        tip = ("points image: its object points (green known, orange solved '?') and its camera.\n"
               "board frame: the board at the origin and the camera of every checked board image around it "
               "(red = the image shown in the Image view). OpenCV's board Z axis points AWAY from the "
               "camera, so the cameras are on the -Z side.\n"
               "camera frame: one camera at the origin; every checked board, and the points image's object "
               "points, where that camera saw them (each from its own shot).")
        add_tooltip(lbl, tip)
        add_tooltip(cb, tip)
        lbl = tk.Label(bar, text="perspective", font=("Arial", 8))
        lbl.pack(side="left")
        sc = tk.Scale(bar, from_=0.0, to=1.0, resolution=0.02, orient=tk.HORIZONTAL, length=110,
                      variable=self._v3_persp_var, showvalue=False, command=lambda _v: self._v3 and self._v3.draw())
        sc.pack(side="left")
        tip = "0 = orthographic (parallel projection); 1 = strong perspective (near / far foreshortening)."
        add_tooltip(lbl, tip)
        add_tooltip(sc, tip)
        lbl = tk.Label(bar, text="  camera size", font=("Arial", 8))
        lbl.pack(side="left")
        sc = tk.Scale(bar, from_=0.05, to=1.0, resolution=0.05, orient=tk.HORIZONTAL, length=90,
                      variable=self._v3_cam_var, showvalue=False, command=lambda _v: self._update_3d())
        sc.pack(side="left")
        add_tooltip(lbl, "Size of the drawn camera pyramids (relative to the scene).")
        add_tooltip(sc, "Size of the drawn camera pyramids (relative to the scene).")
        b = tk.Button(bar, text="Reset view", font=("Arial", 8), command=lambda: self._v3 and self._v3.reset())
        b.pack(side="left", padx=(8, 0))
        add_tooltip(b, "Back to the default direction and zoom, fitted to the scene.")
        self._v3 = View3D(tab, self._v3_persp_var)
        self._v3.canvas.pack(fill="both", expand=True)
        add_tooltip(self._v3.canvas, "left-drag = rotate, right-drag = pan, wheel = zoom. Camera pyramids: the "
                                     "apex is the camera centre, the rectangle is the image (small triangle = "
                                     "image top). Axes: X red, Y green, Z blue.")

    def _scene_3d(self) -> tuple:
        res = self._res
        mode = self._v3_mode_var.get()
        if res is None:
            return [], "No result yet: Solve first."
        K, size = res["K"], self._image_size() or (1000, 1000)
        try:
            cam = float(self._v3_cam_var.get())
        except (tk.TclError, ValueError):
            cam = 0.3
        scene = []
        if mode.startswith("points"):
            pp = res.get("points_pose")
            em = self._eval_map
            if pp is None or em is None or em["pi"] is None:
                return [], "No pose of the points image in this result (the table was not used)."
            objs = np.array([self._solved_obj(i) for i in em["pt_rows"]], dtype=float).reshape(-1, 3)
            R, _ = cv2.Rodrigues(np.asarray(pp[0], float))
            C = -R.T @ np.asarray(pp[1], float)
            ext = float(np.max(np.ptp(np.vstack([objs, C[None]]), axis=0))) or 1.0
            scene += axes_lines(np.zeros(3), 0.15 * ext)
            known = [k for k, i in enumerate(em["pt_rows"]) if self._infos[i]["axis"] is None]
            unk = [k for k, i in enumerate(em["pt_rows"]) if self._infos[i]["axis"] is not None]
            if known:
                scene.append(("dot", objs[known], KIND_COLORS["point"], 3))
            if unk:
                scene.append(("dot", objs[unk], KIND_COLORS["unknown"], 4))
            if len(objs) <= 80 and self._show_names_var.get():
                scene += [("text", objs[k], self.rows[i]["name"], "#555555") for k, i in enumerate(em["pt_rows"])]
            scene += frustum_lines(R, pp[1], K, size, cam * ext, "#d62728", 2)
            return scene, ""
        if not self._cb or not res["cb_poses"]:
            return [], "No board poses in this result (no checked board images were used)."
        by_path = {c["path"]: c for c in self._cb["cands"]}
        order = [c["path"] for c in self._cb["cands"] if c["path"] in res["cb_poses"]]
        sq = self._square_size()
        if mode.startswith("board"):
            # Each board image has its own lattice: draw the shown one (else the first).
            ref = by_path[self._show if self._show in res["cb_poses"] else order[0]]
            grid = self._cand_obj(ref)
            nodes = {tuple(k): P for k, P in zip(ref["rc"].tolist(), grid)}
            for (r, col), P in nodes.items():
                for nb in ((r, col + 1), (r + 1, col)):
                    if nb in nodes:
                        scene.append(("line", np.array([P, nodes[nb]]), "#c8c8c8", 1))
            scene.append(("line", self._lattice_box(ref)[[0, 1, 2, 3, 0]], "#555555", 2))
            centers = [-cv2.Rodrigues(np.asarray(res["cb_poses"][p][0], float))[0].T @ np.asarray(res["cb_poses"][p][1], float)
                       for p in order]
            ext = float(np.max(np.ptp(np.vstack([grid] + [c[None] for c in centers]), axis=0))) or 1.0
            scene += axes_lines(np.zeros(3), 3 * sq)
            for k, p in enumerate(order):
                R, _ = cv2.Rodrigues(np.asarray(res["cb_poses"][p][0], float))
                cur = p == self._show
                col = "#d62728" if cur else ("#b07800" if p == self._world else "#6a3f9e")
                scene += frustum_lines(R, res["cb_poses"][p][1], K, size, cam * ext * 0.5, col, 2 if cur else 1)
                scene.append(("text", centers[k], str(k + 1), col))
            return scene, ""
        # camera frame: X right, Z = optical axis shown horizontal, image up = up
        M = np.array([[1.0, 0, 0], [0, 0, 1.0], [0, -1.0, 0]])
        pts_all = [np.zeros((1, 3))]
        for p in order:
            R, _ = cv2.Rodrigues(np.asarray(res["cb_poses"][p][0], float))
            t = np.asarray(res["cb_poses"][p][1], float)
            box = (self._lattice_box(by_path[p]) @ R.T + t) @ M.T
            X = (self._cand_obj(by_path[p]) @ R.T + t) @ M.T
            cur = p == self._show
            col = "#d62728" if cur else "#6a3f9e"
            scene.append(("line", box[[0, 1, 2, 3, 0]], col, 2 if cur else 1))
            scene.append(("dot", box[0][None], col, 2.5))
            pts_all.append(X)
        pp, em = res.get("points_pose"), self._eval_map
        if pp is not None and em is not None and em["pi"] is not None:
            objs = np.array([self._solved_obj(i) for i in em["pt_rows"]], dtype=float).reshape(-1, 3)
            R, _ = cv2.Rodrigues(np.asarray(pp[0], float))
            X = (objs @ R.T + np.asarray(pp[1], float)) @ M.T
            scene.append(("dot", X, KIND_COLORS["point"], 3))
            pts_all.append(X)
        ext = float(np.max(np.ptp(np.vstack(pts_all), axis=0))) or 1.0
        scene += frustum_lines(np.eye(3), np.zeros(3), K, size, cam * ext, "#222222", 2, M=M)
        scene += axes_lines(np.zeros(3), 0.12 * ext, M=M)
        return scene, ""

    def _update_3d(self, refit: bool = False) -> None:
        if self._v3 is None:
            return
        scene, msg = self._scene_3d()
        self._v3.empty_text = msg
        self._v3.set_scene(scene, refit=refit)

    # Undistort
    def _build_undistort_tab(self, tab) -> None:
        bar = tk.Frame(tab)
        bar.pack(fill="x")
        lbl = tk.Label(bar, text="alpha", font=("Arial", 8))
        lbl.pack(side="left")
        sc = tk.Scale(bar, from_=0.0, to=1.0, resolution=0.05, orient=tk.HORIZONTAL, length=140,
                      variable=self._und_alpha_var, command=lambda _v: self._update_undistort())
        sc.pack(side="left")
        tip = ("cv2.getOptimalNewCameraMatrix alpha: 0 = zoom in so every pixel is valid (no black borders), "
               "1 = keep every source pixel (black, curved borders).")
        add_tooltip(lbl, tip)
        add_tooltip(sc, tip)
        cbx = tk.Checkbutton(bar, text="points and straight lines", variable=self._und_pts_var, font=("Arial", 8),
                             command=lambda: self._und_iv and self._und_iv.render())
        cbx.pack(side="left", padx=(8, 0))
        add_tooltip(cbx, "Draw the measured points without distortion. Line points should now lie on a STRAIGHT "
                         "line (drawn through them); the largest distance is shown below.")
        self._und_info_var = tk.StringVar()
        tk.Label(tab, textvariable=self._und_info_var, font=("Consolas", 8), fg="#555555", anchor="w",
                 justify="left").pack(fill="x")
        self._und_iv = ImageView(tab, lambda *_a: None, lambda *_a: None,
                                 lambda u, v: self._und_info_var.set(self._und_text + f"\nu {u:9.2f}   v {v:9.2f}"))
        self._und_iv.canvas.pack(fill="both", expand=True)
        self._und_iv.draw_overlay = self._draw_und_overlay
        self._und_iv.empty_text = "Solve first; the image shown in the Image view is undistorted here."
        add_tooltip(self._und_iv.canvas, "The image of the Image view without lens distortion. wheel = zoom, "
                                         "left-drag = pan.")

    def _update_undistort(self) -> None:
        iv = self._und_iv
        if iv is None or self._right_tab() != "Undistort":
            return
        res, img = self._res, self._shown_image()
        if res is None or img is None:
            self._und_state = None
            iv.set_image(None)
            self._und_text = ""
            self._und_info_var.set("")
            return
        K, D = np.asarray(res["K"], float), np.asarray(res["D"], float)
        a = np.asarray(img)
        size = (a.shape[1], a.shape[0])
        try:
            alpha = float(self._und_alpha_var.get())
        except (tk.TclError, ValueError):
            alpha = 0.0
        key = (tuple(K.ravel()), tuple(D.ravel()), size, round(alpha, 4))
        if self._und_cache.get("key") != key:
            newK, roi = cv2.getOptimalNewCameraMatrix(K, D, size, alpha, size)
            m1, m2 = cv2.initUndistortRectifyMap(K, D, None, newK, size, cv2.CV_16SC2)
            self._und_cache = {"key": key, "newK": newK, "roi": roi, "m1": m1, "m2": m2, "img_id": None}
        uc = self._und_cache
        if uc.get("img_id") is not id(img):
            uc["img"] = cv2.remap(a, uc["m1"], uc["m2"], cv2.INTER_LINEAR)
            uc["img_id"] = id(img)
            iv.set_image(uc["img"])
        nk = uc["newK"]
        self._und_state = {"K": K, "D": D, "newK": nk, "roi": uc["roi"]}
        self._und_text = (f"new camera matrix: fx {nk[0, 0]:.2f}  fy {nk[1, 1]:.2f}  cx {nk[0, 2]:.2f}  "
                          f"cy {nk[1, 2]:.2f}   valid region x {uc['roi'][0]} y {uc['roi'][1]} "
                          f"w {uc['roi'][2]} h {uc['roi'][3]}")
        self._und_info_var.set(self._und_text)
        iv.render()

    def _draw_und_overlay(self, iv: ImageView) -> None:
        st = getattr(self, "_und_state", None)
        if st is None or iv.view is None:
            return
        c, ts = iv.canvas, iv.to_screen
        x, y, w, h = st["roi"]
        if w > 0 and h > 0:
            c.create_rectangle(*ts(x - 0.5, y - 0.5), *ts(x + w - 0.5, y + h - 0.5), outline="#ffffff",
                               dash=(4, 4), tags="ov")
        if not self._und_pts_var.get():
            return

        def und(pts):
            p = np.asarray(pts, float).reshape(-1, 1, 2)
            return cv2.undistortPointsIter(p, st["K"], st["D"], None, st["newK"],
                                           (cv2.TERM_CRITERIA_COUNT | cv2.TERM_CRITERIA_EPS, 50, 1e-10)).reshape(-1, 2)
        if self._show is not None:
            cand = next((q for q in self._cb["cands"] if q["path"] == self._show), None) if self._cb else None
            if cand is not None:
                for X, Y in (ts(*q) for q in und(cand["corners"].reshape(-1, 2))):
                    c.create_oval(X - 2.5, Y - 2.5, X + 2.5, Y + 2.5, outline="#3ddc84", tags="ov")
            return
        pts = [(i, r) for i, r in enumerate(self.rows) if math.isfinite(r["u"]) and math.isfinite(r["v"]) and r["use"]]
        if not pts:
            return
        uv = und([(r["u"], r["v"]) for _i, r in pts])
        for (i, r), q in zip(pts, uv):
            X, Y = ts(*q)
            col = KIND_COLORS[self._infos[i]["kind"]] if i < len(self._infos) else "#ffffff"
            c.create_oval(X - 3, Y - 3, X + 3, Y + 3, outline=col, width=2, tags="ov")
        worst = []
        for name in self._line_names():
            q = np.array([uvk for (i, r), uvk in zip(pts, uv) if r["line"] == name])
            if len(q) < 2:
                continue
            vx, vy, x0, y0 = cv2.fitLine(q.astype(np.float32), cv2.DIST_L2, 0, 0.01, 0.01).ravel()
            d = np.array([vx, vy])
            s = (q - [x0, y0]) @ d
            span = s.max() - s.min()
            a = np.array([x0, y0]) + d * (s.min() - 0.1 * span)
            b = np.array([x0, y0]) + d * (s.max() + 0.1 * span)
            c.create_line(*ts(*a), *ts(*b), fill=CURVE_COLOR, width=1, tags="ov")
            dist = np.abs((q - [x0, y0]) @ np.array([-vy, vx]))
            worst.append(f"{name} {dist.max():.2f}")
        if worst:
            self._und_info_var.set(self._und_text + "\nlargest distance to the straight line (px): " + ", ".join(worst))

    # Sensitivity (leave one out)
    def _build_loo_tab(self, tab) -> None:
        tk.Label(tab, font=("Arial", 8), fg="#4a5568", anchor="w", justify="left", wraplength=760, text=(
            "Leave one out: solve again without ONE board image / table point at a time (same settings, in "
            "parallel processes). Large changes = that item has a strong influence on the result -- check it "
            "(a wrong click, a blurred board ...). Click a column title to sort; double-click a row to show "
            "the item.")).pack(fill="x")
        bar = tk.Frame(tab)
        bar.pack(fill="x", pady=(2, 2))
        tk.Label(bar, text="Leave out:", font=("Arial", 8)).pack(side="left")
        for text, var, tip in (("board images", self._loo_boards_var, "One run per checked board image."),
                               ("table points", self._loo_points_var,
                                "One run per used table point (known or '?'); needs 5 or more known points.")):
            cbx = tk.Checkbutton(bar, text=text, variable=var, font=("Arial", 8))
            cbx.pack(side="left")
            add_tooltip(cbx, tip)
        self._loo_btn = tk.Button(bar, text="Run", font=("Arial", 8, "bold"), width=8, command=self._on_loo_run)
        self._loo_btn.pack(side="left", padx=(8, 3))
        add_tooltip(self._loo_btn, "Start the leave-one-out runs with the current Solve settings "
                                   "(uses all processor cores but one).")
        b = tk.Button(bar, text="Stop", font=("Arial", 8), command=self._on_loo_stop)
        b.pack(side="left")
        add_tooltip(b, "Stop; finished rows are kept.")
        self._loo_prog_var = tk.DoubleVar(value=0.0)
        ttk.Progressbar(bar, variable=self._loo_prog_var, maximum=100.0, length=140).pack(side="left", padx=6)
        self._loo_text_var = tk.StringVar(value=self._loo_text)
        tk.Label(tab, textvariable=self._loo_text_var, font=("Arial", 8), fg="#334455", anchor="w").pack(fill="x")
        tf = tk.Frame(tab)
        tf.pack(fill="both", expand=True)
        cols = ("item", "rms", "dfx", "dfy", "dcx", "dcy", "dk1", "dk2", "note")
        tree = ttk.Treeview(tf, columns=cols, show="headings", selectmode="browse")
        for col, title, w in (("item", "left out", 150), ("rms", "RMS px", 70), ("dfx", "Δfx", 70),
                              ("dfy", "Δfy", 70), ("dcx", "Δcx", 70), ("dcy", "Δcy", 70),
                              ("dk1", "Δk1", 80), ("dk2", "Δk2", 80), ("note", "note", 160)):
            tree.heading(col, text=title, command=lambda cc=col: self._sort_loo(cc))
            tree.column(col, width=w, anchor="w" if col in ("item", "note") else "e", stretch=col == "note")
        vsb = ttk.Scrollbar(tf, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        tree.pack(fill="both", expand=True)
        tree.tag_configure("strong", background="#ffe3e3")
        tree.bind("<Double-Button-1>", self._on_loo_double)
        add_tooltip(tree, "Δ = value without this item minus the value of the full solve. The 3 most "
                          "influential rows (largest relative change of fx, fy, cx, cy) are shaded red.")
        self._loo_tree = tree
        self._refresh_loo_tree()

    def _loo_jobs(self) -> list:
        base = self._prepare_solve()
        prob = base["prob"]
        has_pts = base["pts"] is not None
        jobs = []
        if self._loo_boards_var.get():
            for j, p in enumerate(base["cb_paths"]):
                views = [v for k, v in enumerate(prob.views) if k != j + has_pts]
                if not views:
                    continue
                jobs.append({"plan": dict(base, prob=dataclasses.replace(prob, views=views),
                                          cb_paths=[q for q in base["cb_paths"] if q != p]),
                             "label": os.path.basename(p), "kind": "board", "key": p})
        if self._loo_points_var.get() and has_pts:
            for i in base["pt_rows"]:
                rows = [dict(r) for r in self.rows]
                rows[i]["use"] = False
                infos = [classify_row(r) for r in rows]
                view, pt_rows, _lines, _iss = build_view(rows, infos)
                if sum(1 for k in pt_rows if infos[k]["axis"] is None) < 4:
                    continue
                views = [view] + list(prob.views[1:])
                jobs.append({"plan": dict(base, prob=dataclasses.replace(prob, views=views), pt_rows=pt_rows),
                             "label": f"row {i + 1} {self.rows[i]['name']}", "kind": "row", "key": self.rows[i]["id"]})
        return jobs

    def _on_loo_run(self) -> None:
        if self._loo_state is not None and not self._loo_state["finished"]:
            return
        if self._res is None or not self._res.get("solved"):
            messagebox.showinfo("Sensitivity", "Solve first: the runs are compared with the full solve.",
                                parent=self._dlg_parent())
            return
        try:
            jobs = self._loo_jobs()
        except ValueError as e:
            messagebox.showerror("Sensitivity", str(e), parent=self._dlg_parent())
            return
        if not jobs:
            messagebox.showinfo("Sensitivity", "Nothing to leave out: tick 'board images' and / or 'table points' "
                                               "(table points need 5 or more known points).", parent=self._dlg_parent())
            return
        n_workers = max(1, min(len(jobs), (os.cpu_count() or 2) - 1))
        pool, futures = ce.start_worker_pool(n_workers, jobs, ce.loo_task)
        snap = self._res["solved"]
        self._loo_state = {"pool": pool, "futures": futures, "finished": False, "t0": time.perf_counter(),
                           "base": {"K": snap["K"], "D": snap["D"]}, "n": len(jobs), "workers": n_workers}
        self._loo_rows = []
        self._loo_base = self._loo_state["base"]
        self._set_loo_text(f"Running {len(jobs)} solves in {n_workers} processes ...")
        self._poll_loo()

    def _on_loo_stop(self) -> None:
        st = self._loo_state
        if st is not None and not st["finished"]:
            st["pool"].shutdown(wait=False, cancel_futures=True)
            st["finished"] = True
            self._set_loo_text(f"Stopped: {len(self._loo_rows)} of {st['n']} done.")
            self._refresh_loo_tree()

    def _set_loo_text(self, text: str) -> None:
        self._loo_text = text
        if self._loo_text_var is not None:
            self._loo_text_var.set(text)

    def _poll_loo(self) -> None:
        st = self._loo_state
        if st is None or st["finished"]:
            return
        done = [f for f in st["futures"] if f.done()]
        for f in done:
            st["futures"].remove(f)
            try:
                self._loo_rows.append(f.result())
            except Exception as e:                       # a broken process: report, keep the rest
                self._loo_rows.append({"label": "?", "kind": "", "key": None, "error": f"{type(e).__name__}: {e}"})
        if self._loo_prog_var is not None:
            try:
                self._loo_prog_var.set(100.0 * len(self._loo_rows) / max(1, st["n"]))
            except tk.TclError:
                pass
        if done:
            self._refresh_loo_tree()
        if st["futures"]:
            self._set_loo_text(f"Running ... {len(self._loo_rows)} / {st['n']} done ({st['workers']} processes)")
            try:
                self.canvas.after(150, self._poll_loo)
            except tk.TclError:
                pass
            return
        st["finished"] = True
        st["pool"].shutdown(wait=False)
        self._set_loo_text(f"Done: {st['n']} solves in {time.perf_counter() - st['t0']:.1f} s "
                           f"({st['workers']} processes).")

    def _loo_values(self, row: dict) -> dict:
        b = self._loo_base
        if row.get("error") or b is None:
            return {}
        K, D, K0, D0 = row["K"], row["D"], b["K"], b["D"]
        size = self._image_size() or (1, 1)
        v = {"rms": row["rms"], "dfx": K[0][0] - K0[0][0], "dfy": K[1][1] - K0[1][1],
             "dcx": K[0][2] - K0[0][2], "dcy": K[1][2] - K0[1][2],
             "dk1": (D[0] if D else 0) - (D0[0] if D0 else 0),
             "dk2": (D[1] if len(D) > 1 else 0) - (D0[1] if len(D0) > 1 else 0)}
        v["influence"] = (abs(v["dfx"]) / K0[0][0] + abs(v["dfy"]) / K0[1][1] + abs(v["dcx"]) / size[0]
                          + abs(v["dcy"]) / size[1])
        return v

    def _sort_loo(self, col: str) -> None:
        self._loo_sort = (col, not self._loo_sort[1]) if self._loo_sort[0] == col else (col, True)
        self._refresh_loo_tree()

    def _refresh_loo_tree(self) -> None:
        tree = self._loo_tree
        if tree is None:
            return
        tree.delete(*tree.get_children())
        rows = [(r, self._loo_values(r)) for r in self._loo_rows]
        col, desc = self._loo_sort
        if col == "item":
            rows.sort(key=lambda t: t[0]["label"], reverse=desc)
        else:
            key = "influence" if col in ("note",) else col
            rows.sort(key=lambda t: (abs(t[1].get(key, -1.0)) if key != "rms" else t[1].get(key, -1.0)),
                      reverse=desc)
        strong = sorted((v.get("influence", 0.0) for _r, v in rows), reverse=True)[:3]
        self._loo_keys = []
        for k, (r, v) in enumerate(rows):
            if not v:
                vals = (r["label"], "", "", "", "", "", "", "", r.get("error") or "")
            else:
                vals = (r["label"], f"{v['rms']:.4f}", f"{v['dfx']:+.3f}", f"{v['dfy']:+.3f}", f"{v['dcx']:+.3f}",
                        f"{v['dcy']:+.3f}", f"{v['dk1']:+.5f}", f"{v['dk2']:+.5f}",
                        "" if r.get("success", True) else "not converged")
            tags = ("strong",) if v and strong and v["influence"] >= strong[-1] and v["influence"] > 0 else ()
            tree.insert("", "end", iid=str(k), values=vals, tags=tags)
            self._loo_keys.append((r.get("kind"), r.get("key")))

    def _on_loo_double(self, e) -> None:
        item = self._loo_tree.identify_row(e.y)
        if not item:
            return
        kind, key = self._loo_keys[int(item)]
        if kind == "board":
            self._show_board(key)
            self._select_tree_path(key)
        elif kind == "row":
            i = next((k for k, r in enumerate(self.rows) if r["id"] == key), None)
            if i is not None:
                self._show_board(None)
                if self._sheet is not None:
                    self._sheet.select([i])
                self.on_selection_changed()

    # Runs
    def _build_runs_tab(self, tab) -> None:
        tk.Label(tab, font=("Arial", 8), fg="#4a5568", anchor="w", justify="left", wraplength=760, text=(
            "Every Solve is listed here, so results with different settings can be compared. 'Restore' makes a "
            "run the current result again (runs are saved with the project, poses included); 'K, D -> initial' "
            "copies its camera into the Solve tab's initial values.")).pack(fill="x")
        bar = tk.Frame(tab)
        bar.pack(fill="x", pady=(2, 2))
        for text, cmd, tip in (("Restore", self._on_run_restore, "Make the selected run the current result "
                                                                  "(outputs, errors, drawings)."),
                               ("K, D → initial", self._on_run_initial, "Copy the selected run's camera "
                                                                             "matrix and distortion into the "
                                                                             "initial values."),
                               ("Delete", self._on_run_delete, "Remove the selected run from the list."),
                               ("Clear all", self._on_runs_clear, "Remove all runs from the list.")):
            b = tk.Button(bar, text=text, font=("Arial", 8), command=cmd)
            b.pack(side="left", padx=(0, 3))
            add_tooltip(b, tip)
        tf = tk.Frame(tab)
        tf.pack(fill="both", expand=True)
        tree = ttk.Treeview(tf, columns=[c for c, _t, _w in self._RUN_COLS], show="headings", selectmode="browse")
        vsb = ttk.Scrollbar(tf, orient="vertical", command=tree.yview)
        hsb = ttk.Scrollbar(tf, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        vsb.pack(side="right", fill="y")
        hsb.pack(side="bottom", fill="x")
        tree.pack(fill="both", expand=True)
        tree.tag_configure("session", foreground="#111111")
        tree.tag_configure("saved", foreground="#777777")
        tree.bind("<Double-Button-1>", lambda _e: self._on_run_restore())
        add_tooltip(tree, "Grey rows come from projects saved by an older version: they keep K and D only. Double-click = Restore. "
                          "solver: OpenCV = cv2.calibrateCamera only, LSQ = least-squares minimization. Columns "
                          "for p1, p2, k3 ... tauy appear when some run has a nonzero value.")
        self._runs_tree = tree
        self._refresh_runs_tree()

    def _add_run(self, plan: dict, out: dict) -> None:
        prob = plan["prob"]
        data = []
        if plan["pts"] is not None:
            data.append(f"{len(plan['pt_rows'])} pts")
        if plan["cb_paths"]:
            data.append(f"{len(plan['cb_paths'])} boards")
        if plan["pts"] is not None and plan["cb_paths"]:
            data.append("two steps" if plan["mode"] == "two_step" else "joint")
        nd = ce.DIST_MODELS[prob.dist_model]
        fixed = [n for n in ce.INTR_NAMES + ce.DIST_NAMES[:nd] if n in prob.fixed]
        route = out.get("route") or {}
        solver = ("OpenCV" if route.get("boards") in ("opencv", None) and route.get("points") in ("opencv", None)
                  else "OpenCV+PnP" if route.get("points") == "pnp" and route.get("boards") == "opencv"
                  else "OpenCV+LSQ" if route.get("boards") == "opencv" else "LSQ")
        self._runs.append({"n": (self._runs[-1]["n"] + 1) if self._runs else 1, "time": time.strftime("%H:%M:%S"),
                           "data": ", ".join(data), "model": prob.dist_model, "solver": solver,
                           "fixed": " ".join(fixed) + (" fy/fx" if prob.fix_aspect else ""),
                           "rms": float(self._eval.rms.get("all", out["rms"]) if self._eval is not None else out["rms"]),
                           "K": np.asarray(out["K"]).tolist(),
                           "D": np.asarray(out["D"]).ravel().tolist(), "full": copy.deepcopy(self._res)})
        del self._runs[:-30]

    _RUN_COLS = (("n", "#", 34), ("time", "time", 64), ("data", "data", 140), ("model", "model", 62),
                 ("solver", "solver", 72), ("fixed", "fixed", 120), ("rms", "RMS px", 64), ("fx", "fx", 74),
                 ("fy", "fy", 74), ("cx", "cx", 70), ("cy", "cy", 70), ("k1", "k1", 72), ("k2", "k2", 72))
    _RUN_EXTRA = ("p1", "p2", "k3", "k4", "k5", "k6", "s1", "s2", "s3", "s4", "taux", "tauy")

    def _refresh_runs_tree(self) -> None:
        tree = self._runs_tree
        if tree is None:
            return
        tree.delete(*tree.get_children())
        # Distortion columns beyond k1 k2, only those with a nonzero value in some run.
        extra = [n for n in self._RUN_EXTRA
                 if any(len(r["D"]) > ce.DIST_NAMES.index(n) and r["D"][ce.DIST_NAMES.index(n)] != 0
                        for r in self._runs)]
        cols = list(self._RUN_COLS) + [(n, n, 80) for n in extra]
        tree.configure(columns=[c for c, _t, _w in cols])
        for col, title, w in cols:
            tree.heading(col, text=title)
            tree.column(col, width=w, minwidth=30, anchor="w" if col in ("time", "data", "model", "solver", "fixed")
                        else "e", stretch=col == "data")
        for k, r in enumerate(self._runs):
            K, D = r["K"], r["D"]
            more = []
            for n in extra:
                j = ce.DIST_NAMES.index(n)
                more.append(f"{D[j]:.6g}" if len(D) > j else "")
            tree.insert("", "end", iid=str(k), tags=("session",) if r.get("full") else ("saved",), values=(
                r["n"], r["time"], r["data"], r["model"], r.get("solver", ""), r["fixed"], f"{r['rms']:.4f}",
                f"{K[0][0]:.2f}", f"{K[1][1]:.2f}", f"{K[0][2]:.2f}", f"{K[1][2]:.2f}", f"{D[0]:.5f}" if D else "",
                f"{D[1]:.5f}" if len(D) > 1 else "", *more))
        if self._runs:
            tree.see(str(len(self._runs) - 1))

    def _selected_run(self):
        tree = self._runs_tree
        sel = tree.selection() if tree is not None else ()
        if not sel:
            messagebox.showinfo("Runs", "Select a run in the list first.", parent=self._dlg_parent())
            return None
        return self._runs[int(sel[0])]

    def _on_run_restore(self) -> None:
        r = self._selected_run()
        if r is None:
            return
        if not r.get("full"):
            messagebox.showinfo("Runs", "This run was loaded with the project and has no poses; use 'K, D -> "
                                        "initial' and Solve with the camera fixed instead.", parent=self._dlg_parent())
            return
        self._res = copy.deepcopy(r["full"])
        self._res["stale"] = True
        self._refresh_all()
        self._sync_sliders()
        self._push()

    def _on_run_initial(self) -> None:
        r = self._selected_run()
        if r is not None:
            self._set_initial(r["K"], r["D"])

    def _on_run_delete(self) -> None:
        tree = self._runs_tree
        sel = tree.selection() if tree is not None else ()
        if sel:
            del self._runs[int(sel[0])]
            self._refresh_runs_tree()

    def _on_runs_clear(self) -> None:
        if self._runs and messagebox.askyesno("Runs", f"Remove all {len(self._runs)} runs from the list?",
                                              parent=self._dlg_parent()):
            self._runs = []
            self._refresh_runs_tree()

    # ── chessboard ─────────────────────────────────────────────────────────
    @staticmethod
    def _coerce_files(value):
        if value is None:
            return None
        if isinstance(value, (str, os.PathLike)):
            items = [str(value)]
        else:
            try:
                items = [str(v) for v in list(value)]
            except TypeError:
                return None
        items = [p for p in items if p.strip()]
        return items or None

    def _cb_files(self) -> list:
        return list(self._cb_pin_files) if self._cb_pin_files else list(self._cb_manual_files)

    def _square_xy(self) -> tuple:
        """(square X, square Y): the board square's size along a board row
        (X) and along a board column (Y); an empty Y box = X."""
        cfg = self._cb["cfg"] if self._cb else {}
        try:
            sx = float(self._det_vars["square"].get())
            if not sx > 0:
                raise ValueError
        except ValueError:
            sx = float(cfg.get("square", 1.0))
        try:
            sy = float(self._det_vars["square_y"].get())     # empty = same as X
            if not sy > 0:
                raise ValueError
        except ValueError:
            sy = sx
        return sx, sy

    def _square_size(self) -> float:
        """Mean square size (for drawing lengths only)."""
        return 0.5 * sum(self._square_xy())

    def _cand_axes(self, cand: dict) -> str | None:
        """See lattice_axes(); always 'x' when square X = Y (the direction
        does not matter then)."""
        sx, sy = self._square_xy()
        if sx == sy:
            return "x"
        cfg = self._cb["cfg"] if self._cb else {}
        return lattice_axes(cand["rc"], (cfg.get("ncol", 0), cfg.get("nrow", 0)))

    def _cand_obj(self, cand: dict) -> np.ndarray:
        """(N, 3) object points of one board image, in the order of its corners."""
        return lattice_obj(cand["rc"], *self._square_xy(), self._cand_axes(cand))

    def _lattice_box(self, cand: dict) -> np.ndarray:
        """(4, 3) object points of the corners of the lattice's bounding box
        (row 0 col 0, row 0 col max, max max, max 0); may be unobserved."""
        rmax, cmax = (int(v) for v in np.max(cand["rc"], axis=0))
        return lattice_obj(np.array([(0, 0), (0, cmax), (rmax, cmax), (rmax, 0)]), *self._square_xy(),
                           self._cand_axes(cand))

    def _cb_views(self) -> list:
        """[(path, View)] of the checked board images, in table (rank) order."""
        if not self._cb:
            return []
        return [(c["path"], ce.View(c["path"], kind="chessboard", obj=self._cand_obj(c),
                                    img=c["corners"].reshape(-1, 2).astype(np.float64)))
                for c in self._cb["cands"] if c["path"] in self._cb_use]

    def _build_chessboard_tab(self, tab) -> None:
        top = tk.Frame(tab)
        top.pack(fill="x")
        b = tk.Button(top, text="Add images...", font=("Arial", 8), command=self._on_add_cb_files)
        b.pack(side="left")
        add_tooltip(b, "Pick chessboard image files (used when nothing is connected to the 'files' input pin).")
        b = tk.Button(top, text="Clear list", font=("Arial", 8), command=self._on_clear_cb_files)
        b.pack(side="left", padx=(3, 0))
        add_tooltip(b, "Forget the picked image files (not the detections).")
        self._cb_files_var = tk.StringVar()
        tk.Label(top, textvariable=self._cb_files_var, font=("Arial", 8), fg="#555555").pack(side="left", padx=6)

        grid = tk.LabelFrame(tab, text="Board and detection", font=("Arial", 8, "bold"), padx=4, pady=2)
        grid.pack(fill="x", pady=(4, 2))
        fields = (("ncol", "columns", "INNER corners along one board row (where four squares meet) -- one less "
                                      "than the squares. OpenCV patternSize width."),
                  ("nrow", "rows", "INNER corners along one board column. OpenCV patternSize height."),
                  ("square", "square X", "Size of one square ALONG A BOARD ROW (the 'columns' direction, "
                                         "object X), in your length unit (e.g. mm). Measure over many squares "
                                         "and divide. May be changed after detection."),
                  ("square_y", "square Y", "Size of one square ALONG A BOARD COLUMN (the 'rows' direction, object "
                                           "Y). Empty = the same as square X. Printers often scale X and Y "
                                           "slightly differently; if yours did, enter both. With X different from Y, columns x rows must be the "
                                           "FULL board's inner corners (used to tell X from Y in each image)."),
                  ("max_width", "detect width", "Corners are searched on a copy shrunk to this width (faster), "
                                                "then refined on the full image. 0 = full size."),
                  ("threads", "threads", "Images analysed at the same time (OpenCV runs them in parallel)."))
        for k, (key, label, tip) in enumerate(fields):
            lbl = tk.Label(grid, text=label, font=("Arial", 8))
            lbl.grid(row=0, column=2 * k, sticky="e", padx=(6 if k else 0, 2))
            e = tk.Entry(grid, textvariable=self._det_vars[key], width=6, justify="right")
            e.grid(row=0, column=2 * k + 1)
            add_tooltip(lbl, tip)
            add_tooltip(e, tip)
        flags = tk.Frame(grid)
        flags.grid(row=1, column=0, columnspan=12, sticky="w")
        for key, label, tip in (("adaptive", "adaptive threshold", "Handle uneven lighting (usually helps)."),
                                ("normalize", "normalize", "Equalize the brightness before detecting."),
                                ("fast", "fast check", "Quickly skip images without a board -- much faster for "
                                                       "long sequences with many empty frames.")):
            cbx = tk.Checkbutton(flags, text=label, variable=self._det_vars[key], font=("Arial", 8))
            cbx.pack(side="left")
            add_tooltip(cbx, tip)
        row = tk.Frame(grid)
        row.grid(row=2, column=0, columnspan=12, sticky="w", pady=(2, 0))
        self._cb_detect_btn = tk.Button(row, text="Detect", font=("Arial", 9, "bold"), bg="#e8dcf5", width=10,
                                        command=self._on_detect)
        self._cb_detect_btn.pack(side="left")
        add_tooltip(self._cb_detect_btn, "Find the board corners in every image (sub-pixel), then rank the images "
                                         "and check the best ones (Auto-select). Runs in the background.")
        b = tk.Button(row, text="Stop", font=("Arial", 8), command=self._on_stop_detect)
        b.pack(side="left", padx=(3, 0))
        add_tooltip(b, "Stop the detection; the images done so far are kept.")
        b = tk.Button(row, text="Search settings...", font=("Arial", 8), command=self._on_open_search)
        b.pack(side="left", padx=(3, 0))
        add_tooltip(b, "Not sure about the pattern size or the options? Try MANY combinations of columns, rows, "
                       "detect width and the three options on every image, and keep, per image, every corner "
                       "any of them found (a corner found by several settings counts once). Shows where the "
                       "corners are and which single setting works best.")
        self._cb_progress_var = tk.DoubleVar(value=0.0)
        ttk.Progressbar(row, variable=self._cb_progress_var, maximum=100.0, length=160).pack(side="left", padx=6)
        self._cb_text_var = tk.StringVar(value=self._cb_text)
        tk.Label(grid, textvariable=self._cb_text_var, font=("Arial", 8), fg="#334455", anchor="w", justify="left",
                 wraplength=640).grid(row=3, column=0, columnspan=12, sticky="w")

        sel = tk.LabelFrame(tab, text="Selection", font=("Arial", 8, "bold"), padx=4, pady=2)
        sel.pack(fill="x", pady=(2, 2))
        for k, (key, label, tip) in enumerate((
                ("target", "target count", "How many images Auto-select checks: 10-30 varied views are usually "
                                           "enough; near-duplicate views do not help."),
                ("min_sharp", "min sharpness", "Images blurrier than this are not checked (0 = any). See the "
                                               "'sharp' column."),
                ("diversity", "min difference", "Auto-select stops when the next ranked image differs from every "
                                                "checked one by less than this fraction of the image diagonal "
                                                "(mean corner distance) -- you may get fewer than the target."))):
            lbl = tk.Label(sel, text=label, font=("Arial", 8))
            lbl.grid(row=0, column=2 * k, sticky="e", padx=(6 if k else 0, 2))
            e = tk.Entry(sel, textvariable=self._sel_vars[key], width=6, justify="right")
            e.grid(row=0, column=2 * k + 1)
            add_tooltip(lbl, tip)
            add_tooltip(e, tip)
        bar = tk.Frame(sel)
        bar.grid(row=1, column=0, columnspan=8, sticky="w", pady=(3, 0))
        for text, cmd, tip in (
                ("Auto-select", self._on_auto_select, "Check the first 'target count' images of the ranking "
                                                       "(see Help) and uncheck the others."),
                ("Check all", lambda: self._set_all_cb(True), "Check every detected image."),
                ("Uncheck all", lambda: self._set_all_cb(False), "Uncheck every image."),
                ("Set as world frame", self._on_set_world, "Use the SELECTED board image's pose as the world frame: "
                                                           "the rvec / tvec outputs become that board's pose (its "
                                                           "corner [0, 0] is the world origin)."),
                ("Points image = world", self._on_world_points, "Use the points image's pose (its object "
                                                                "coordinates) for the rvec / tvec outputs.")):
            b = tk.Button(bar, text=text, font=("Arial", 8), command=cmd)
            b.pack(side="left", padx=(0, 3))
            add_tooltip(b, tip)
        self._world_var = tk.StringVar()
        tk.Label(sel, textvariable=self._world_var, font=("Arial", 8, "bold"), fg="#7a3d00", anchor="w").grid(
            row=2, column=0, columnspan=8, sticky="w")

        tf = tk.Frame(tab)
        tf.pack(fill="both", expand=True, pady=(2, 0))
        cols = ("use", "rank", "src", "file", "sharp", "cover", "err", "world")
        tree = ttk.Treeview(tf, columns=cols, show="headings", height=10, selectmode="browse")
        for col, title, w, anchor in (("use", "use", 40, "center"), ("rank", "rank", 44, "e"), ("src", "src #", 50, "e"),
                                      ("file", "file", 230, "w"), ("sharp", "sharp", 70, "e"),
                                      ("cover", "cover %", 60, "e"), ("err", "RMS px", 70, "e"),
                                      ("world", "world", 50, "center")):
            tree.heading(col, text=title)
            tree.column(col, width=w, anchor=anchor, stretch=(col == "file"))
        vsb = ttk.Scrollbar(tf, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        tree.pack(fill="both", expand=True)
        tree.tag_configure("unused", foreground="#8a8a8a")
        tree.tag_configure("world", foreground="#7a3d00", font=("Arial", 9, "bold"))
        tree.bind("<ButtonRelease-1>", self._on_tree_click)
        tree.bind("<<TreeviewSelect>>", self._on_tree_select)
        tree.bind("<space>", lambda _e: self._toggle_selected_cb())
        self._cb_tree = tree
        add_tooltip(tree, "Ranked board images: rank 1 = best; each next one is the most different from those "
                          "before. Click 'use' = check / uncheck (or Space). Select a row to see the image. "
                          "sharp = Laplacian variance on the board, cover = board area / image area, RMS = "
                          "reprojection error after Solve.")
        self._update_cb_files_label()

    # ── settings search (brute force over detection settings) ───────────────
    _SEARCH_FLAG_CHOICES = {"on": [True], "off": [False], "both": [False, True]}
    _SEARCH_FLAGS = (("adaptive", "adaptive threshold"), ("normalize", "normalize"), ("fast", "fast check"))
    _SEARCH_FLAG_CHOICES_KEYS = ("adaptive", "normalize", "fast")
    _SEARCH_PALETTE = ("#00e5ff", "#ffe600", "#ff9f1a", "#3ddc84", "#ff5cf0", "#4d8dff", "#ff3b30", "#ffffff")

    @staticmethod
    def _fmt_seconds(t: float) -> str:
        t = int(round(t))
        return f"{t // 3600}:{t // 60 % 60:02d}:{t % 60:02d}" if t >= 3600 else f"{t // 60}:{t % 60:02d}"

    def _parse_search_settings(self) -> list:
        """The Setting list from the dialog fields; ValueError with a message."""
        sv = self._search_vars

        def ints(key, label, low):
            arr = cbs.parse_number_list(sv[key].get())
            if arr is None or not len(arr) or any(v != int(v) or v < low for v in arr):
                raise ValueError(f"{label}: whole numbers {low} or more, e.g. '7 8 9' or '7:9'.")
            return list(dict.fromkeys(int(v) for v in arr))

        cols = ints("cols", "columns", 3)
        rows = ints("rows", "rows", 3)
        widths = ints("widths", "detect width", 0)
        flags = [self._SEARCH_FLAG_CHOICES[sv[k].get()] for k, _l in self._SEARCH_FLAGS]
        return list(dict.fromkeys(cbs.build_settings(cols, rows, widths, *flags)))

    def _update_search_count(self) -> None:
        if self._search_win is None:
            return
        n_img = len(self._cb_files())
        try:
            n = len(self._parse_search_settings())
        except ValueError as e:
            self._search_count_var.set(str(e))
            return
        sv = self._search_vars
        parts = [len(set(cbs.parse_number_list(sv[k].get()).tolist())) for k in ("cols", "rows", "widths")]
        parts += [len(self._SEARCH_FLAG_CHOICES[sv[k].get()]) for k, _l in self._SEARCH_FLAGS]
        self._search_count_var.set(
            f"{' x '.join(str(v) for v in parts)} = {n} settings per image  x  {n_img} images  =  "
            f"{n * n_img:,} detections  (repeated ones, e.g. a detect width wider than the image, run once)")

    def _on_open_search(self) -> None:
        win = self._search_win
        if win is not None:
            try:
                win.deiconify()
                win.lift()
                return
            except tk.TclError:
                self._search_win = None
        self._build_search_window()

    def _build_search_window(self) -> None:
        win = tk.Toplevel(self._dlg_parent())
        win.title(f"Search detection settings -- {self.get_inspector_title()}")
        win.geometry("1180x760")
        win.protocol("WM_DELETE_WINDOW", self._on_search_close)
        self._search_win = win
        sv = self._search_vars = {k: tk.StringVar(value=v) for k, v in self._search_cfg.items()}
        tk.Label(win, justify="left", anchor="w", wraplength=1140, font=("Arial", 9), text=(
            "Every image is searched with EVERY combination below. In each image, all corners found by any "
            "setting are put together on one board lattice (a corner found by several settings counts once "
            "and is placed where all agree), so an image can give more corners than any single setting -- not "
            "always a full rectangle. Detections that do not fit the lattice exactly are rejected, never "
            "guessed. Then 'Use these corners' replaces the board list of the Chessboard tab.")).pack(
            fill="x", padx=8, pady=(6, 2))

        f = tk.LabelFrame(win, text="Settings to try", font=("Arial", 8, "bold"), padx=6, pady=3)
        f.pack(fill="x", padx=8)
        for k, (key, label, tip) in enumerate((
                ("cols", "columns", "INNER corners per board row to try, e.g. '7 8 9' or '7:9'. Smaller "
                                    "numbers can find a part of the board when it is cut by the image edge "
                                    "or covered."),
                ("rows", "rows", "INNER corners per board column to try, e.g. '12 13 14'."),
                ("widths", "detect width", "Shrink the image to these widths for the search (0 = full size), "
                                           "e.g. '800 1200 1600 0'. Corners are always refined on the full "
                                           "image."))):
            lbl = tk.Label(f, text=label, font=("Arial", 8))
            lbl.grid(row=0, column=2 * k, sticky="e", padx=(8 if k else 0, 2))
            e = tk.Entry(f, textvariable=sv[key], width=18)
            e.grid(row=0, column=2 * k + 1, sticky="w")
            add_tooltip(lbl, tip)
            add_tooltip(e, tip)
        for k, (key, label) in enumerate(self._SEARCH_FLAGS):
            lbl = tk.Label(f, text=label, font=("Arial", 8))
            lbl.grid(row=1, column=2 * k, sticky="e", padx=(8 if k else 0, 2), pady=(3, 0))
            cb = ttk.Combobox(f, textvariable=sv[key], values=("on", "off", "both"), state="readonly", width=6)
            cb.grid(row=1, column=2 * k + 1, sticky="w", pady=(3, 0))
            tip = "on / off / both (try with and without -- doubles the time)."
            add_tooltip(lbl, tip)
            add_tooltip(cb, tip)
        self._search_count_var = tk.StringVar()
        tk.Label(f, textvariable=self._search_count_var, font=("Arial", 8), fg="#334455", anchor="w").grid(
            row=2, column=0, columnspan=6, sticky="w", pady=(3, 0))
        for v in sv.values():
            v.trace_add("write", lambda *_a: self._on_search_field_changed())

        run = tk.Frame(win)
        run.pack(fill="x", padx=8, pady=(4, 0))
        self._search_start_btn = tk.Button(run, text="Start", font=("Arial", 9, "bold"), bg="#e8dcf5", width=10,
                                           command=self._on_search_start)
        self._search_start_btn.pack(side="left")
        add_tooltip(self._search_start_btn, "Run the search in the background (uses the 'threads' of the "
                                            "Chessboard tab). Images appear in the list as they finish.")
        self._search_cancel_btn = tk.Button(run, text="Cancel", font=("Arial", 8), command=self._on_search_cancel)
        self._search_cancel_btn.pack(side="left", padx=(3, 0))
        add_tooltip(self._search_cancel_btn, "Stop the search. Images already started keep the corners found so "
                                             "far (marked 'partial').")
        self._search_progress_var = tk.DoubleVar(value=0.0)
        ttk.Progressbar(run, variable=self._search_progress_var, maximum=100.0, length=260).pack(side="left",
                                                                                                padx=6)
        self._search_status_var = tk.StringVar(value="Not started.")
        tk.Label(run, textvariable=self._search_status_var, font=("Arial", 8), fg="#334455").pack(side="left")
        self._search_found_var = tk.StringVar(value="")
        tk.Label(win, textvariable=self._search_found_var, font=("Arial", 10, "bold"), fg="#1b5e20",
                 anchor="w").pack(fill="x", padx=8, pady=(2, 2))

        pane = tk.PanedWindow(win, orient="horizontal", sashwidth=5)
        pane.pack(fill="both", expand=True, padx=8)
        nb = ttk.Notebook(pane)
        pane.add(nb, width=560)

        tab = tk.Frame(nb)
        nb.add(tab, text="Images")
        cols = ("file", "corners", "found", "best", "gain", "note")
        tree = ttk.Treeview(tab, columns=cols, show="headings", selectmode="browse")
        for col, title, w, anchor in (("file", "file", 170, "w"), ("corners", "corners", 60, "e"),
                                      ("found", "settings found", 90, "e"), ("best", "best single", 76, "e"),
                                      ("gain", "gain", 46, "e"), ("note", "note", 150, "w")):
            tree.heading(col, text=title)
            tree.column(col, width=w, anchor=anchor, stretch=(col in ("file", "note")))
        vsb = ttk.Scrollbar(tab, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        tree.pack(fill="both", expand=True)
        tree.tag_configure("none", foreground="#8a8a8a")
        tree.tag_configure("gain", foreground="#1b5e20")
        tree.bind("<<TreeviewSelect>>", lambda _e: self._on_search_select())
        add_tooltip(tree, "One row per image (as it finishes). corners = different board corners found (each "
                          "counted once); settings found = how many settings found a board; best single = the "
                          "most corners ONE setting found; gain = corners more than the best single setting. "
                          "Select a row to see where the corners are.")
        self._search_tree = tree

        tab = tk.Frame(nb)
        nb.add(tab, text="Settings")
        cols = ("setting", "images", "corners")
        tree = ttk.Treeview(tab, columns=cols, show="headings", selectmode="browse")
        for col, title, w, anchor in (("setting", "setting", 330, "w"), ("images", "images", 70, "e"),
                                      ("corners", "corners", 80, "e")):
            tree.heading(col, text=title)
            tree.column(col, width=w, anchor=anchor, stretch=(col == "setting"))
        bar = tk.Frame(tab)
        bar.pack(side="bottom", fill="x", pady=(3, 0))
        b = tk.Button(bar, text="Copy to the Detect fields", font=("Arial", 8), command=self._on_search_copy_setting)
        b.pack(side="left")
        add_tooltip(b, "Prefer ONE fixed rectangular pattern for all images? Copy the selected setting to the "
                       "Chessboard tab (columns, rows, detect width, options), then press Detect there.")
        vsb = ttk.Scrollbar(tab, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        tree.pack(fill="both", expand=True)
        add_tooltip(tree, "Every setting alone, best first: images = in how many images it found a board, "
                          "corners = how many corners it found in total (its own full rectangles).")
        self._search_set_tree = tree

        right = tk.Frame(pane)
        pane.add(right)
        top = tk.Frame(right)
        top.pack(fill="x")
        self._search_mode = tk.StringVar(value="image")
        for value, text, tip in (("image", "Selected image", "The selected image with its corners."),
                                 ("all", "All images", "The corners of EVERY image on one frame, one colour per "
                                                       "image: where in the image the calibration has data "
                                                       "(the edges and corners matter for the distortion).")):
            rb = tk.Radiobutton(top, text=text, value=value, variable=self._search_mode, font=("Arial", 8),
                                command=self._on_search_select)
            rb.pack(side="left")
            add_tooltip(rb, tip)
        iv = ImageView(right, lambda *_a: None, lambda *_a: None, self._on_search_motion)
        iv.draw_overlay = self._draw_search_overlay
        iv.empty_text = "Select an image in the list."
        iv.canvas.pack(fill="both", expand=True)
        self._search_iv = iv
        self._search_hover_var = tk.StringVar(value="")
        tk.Label(right, textvariable=self._search_hover_var, font=("Consolas", 8), anchor="w").pack(fill="x")
        tk.Label(right, justify="left", anchor="w", wraplength=560, font=("Arial", 8), fg="#444444", text=(
            "Dot colour = how many settings found that corner: red = one, yellow = half, green = all. Lines "
            "join neighbouring board corners. White ring = the lattice origin, object point (0, 0, 0); arrows "
            "= +c (object X) and +r (object Y). Wheel = zoom, drag = pan, hover a corner for details.")).pack(
            fill="x")

        bottom = tk.Frame(win)
        bottom.pack(fill="x", padx=8, pady=6)
        b = tk.Button(bottom, text="Use these corners", font=("Arial", 9, "bold"), bg="#d7ecd9",
                      command=self._on_search_apply)
        b.pack(side="left")
        add_tooltip(b, "Replace the board images of the Chessboard tab with the search result (then they are "
                       "ranked and auto-selected as after Detect).")
        b = tk.Button(bottom, text="Close", font=("Arial", 8), command=self._on_search_close)
        b.pack(side="right")
        add_tooltip(b, "Close this window. A running search goes on; open 'Search settings...' again to see it.")

        self._update_search_count()
        self._search_fill_trees(full=True)
        self._set_search_running(self._search is not None and not self._search["state"]["finished"])
        if self._search is not None:
            self._poll_search()

    def _on_search_field_changed(self) -> None:
        for k, v in self._search_vars.items():
            self._search_cfg[k] = v.get()
        self._update_search_count()

    def _set_search_running(self, running: bool) -> None:
        if self._search_win is None:
            return
        try:
            self._search_start_btn.configure(state="disabled" if running else "normal",
                                             text="Searching..." if running else "Start")
            self._search_cancel_btn.configure(state="normal" if running else "disabled")
        except tk.TclError:
            pass

    def _on_search_start(self) -> None:
        if self._search is not None and not self._search["state"]["finished"]:
            return
        files = self._cb_files()
        if not files:
            messagebox.showinfo("Search", "Connect chessboard images to the 'files' input pin, or press 'Add "
                                          "images...' in the Chessboard tab.", parent=self._search_win)
            return
        try:
            settings = self._parse_search_settings()
            threads = max(1, int(self._det_vars["threads"].get()))
        except ValueError as e:
            messagebox.showerror("Search", str(e) if str(e) else "threads must be a whole number.",
                                 parent=self._search_win)
            return
        total = len(settings) * len(files)
        if total > 20000 and not messagebox.askokcancel(
                "Search", f"{total:,} detections can take a long time (a failed detection of a large pattern "
                          f"can take a second or more). Start anyway? You can cancel at any time.",
                parent=self._search_win):
            return
        paths = [self._abs(p) for p in files]
        state = cbs.new_search_state(len(paths), len(settings))
        self._search = {"settings": settings, "paths": paths, "state": state, "t0": time.monotonic(),
                        "rows": set(), "sel": None, "t_end": None}
        threading.Thread(target=cbs.search_worker, args=(paths, settings, threads, state), daemon=True).start()
        self._search_fill_trees(full=True)
        self._set_search_running(True)
        self._poll_search()

    def _on_search_cancel(self) -> None:
        if self._search is not None and not self._search["state"]["finished"]:
            self._search["state"]["cancel"].set()

    def _on_search_close(self) -> None:
        win, self._search_win = self._search_win, None
        if win is not None:
            try:
                win.destroy()
            except tk.TclError:
                pass

    def _search_result(self, i):
        se = self._search
        return None if se is None or i is None else se["state"]["results"][i]

    def _poll_search(self) -> None:
        se = self._search
        if se is None:
            return
        st = se["state"]
        finished = st["finished"]
        if finished and se["t_end"] is None:
            se["t_end"] = time.monotonic()
        if self._search_win is not None:
            try:
                self._search_update_status()
                self._search_fill_trees(full=finished)
            except tk.TclError:
                self._search_win = None
        if not finished:
            try:
                self.canvas.after(250, self._poll_search)
            except tk.TclError:
                pass
            return
        self._set_search_running(False)
        if st["error"] and self._search_win is not None:
            messagebox.showerror("Search", f"The search failed: {st['error']}", parent=self._search_win)

    def _search_update_status(self) -> None:
        se = self._search
        st = se["state"]
        done, total = st["done"], max(1, st["total"])
        elapsed = (se["t_end"] or time.monotonic()) - se["t0"]
        self._search_progress_var.set(100.0 * done / total)
        text = (f"{done:,} / {total:,} detections ({100.0 * done / total:.1f}%)   images "
                f"{st['images_done']} / {len(se['paths'])}   elapsed {self._fmt_seconds(elapsed)}")
        if st["finished"]:
            text += "   -- cancelled" if st["cancel"].is_set() else "   -- done"
        elif done:
            text += f"   remaining ~{self._fmt_seconds(elapsed * (total - done) / done)}"
        self._search_status_var.set(text)
        live = dict(st["live"])
        n_img = sum(1 for v in live.values() if v > 0)
        self._search_found_var.set(f"Corners found: {sum(live.values()):,} in {n_img} image(s)   "
                                   f"(the same corner of the same image counts once)")

    def _search_fill_trees(self, full: bool) -> None:
        """Add rows of newly finished images; with `full`, rebuild both lists."""
        se = self._search
        tree, stree = self._search_tree, self._search_set_tree
        if full:
            tree.delete(*tree.get_children())
            stree.delete(*stree.get_children())
            if se is not None:
                se["rows"] = set()
        if se is None:
            return
        n_set = len(se["settings"])
        for i, r in enumerate(list(se["state"]["results"])):
            if r is None or i in se["rows"]:
                continue
            se["rows"].add(i)
            note = []
            if r.get("error"):
                note.append(r["error"])
            elif not r.get("found"):
                note.append("no board")
            if r.get("partial"):
                note.append("partial (cancelled)")
            if r.get("rejected"):
                note.append(f"{r['rejected']} detection(s) rejected")
            if r.get("n_lattices", 1) > 1:
                note.append(f"{r['n_lattices'] - 1} separate part(s) dropped")
            n = len(r["rc"]) if r.get("found") else 0
            gain = n - r.get("best_single", 0) if r.get("found") else 0
            tag = "none" if not n else ("gain" if gain > 0 else "")
            tree.insert("", "end", iid=str(i), tags=(tag,), values=(
                os.path.basename(r["path"]), n or "", f"{len(r.get('settings_merged', [])) if n else 0} / {n_set}",
                r.get("best_single") or "", f"+{gain}" if gain > 0 else "", ", ".join(note)))
        if full:
            hits = se["state"]["setting_hits"]
            order = sorted(range(n_set), key=lambda k: (-hits[k][0], -hits[k][1], k))
            for k in order:
                stree.insert("", "end", iid=str(k), values=(se["settings"][k].label(), hits[k][0], hits[k][1]))

    def _on_search_select(self) -> None:
        se = self._search
        iv = self._search_iv
        if se is None or iv is None:
            return
        sel = self._search_tree.selection()
        se["sel"] = int(sel[0]) if sel else None
        if self._search_mode.get() == "all":
            sizes = [r["size"] for r in se["state"]["results"] if r is not None and r.get("size")]
            iv.set_image(None)
            iv.empty_text = "Corners of all images"
            iv.render(size=tuple(sizes[0]) if sizes else None)
            return
        r = self._search_result(se["sel"])
        iv.empty_text = "Select an image in the list."
        iv.set_image(_read_rgb(r["path"]) if r is not None else None)

    def _draw_search_overlay(self, iv: ImageView) -> None:
        se = self._search
        if se is None or iv.view is None:
            return
        c, ts = iv.canvas, iv.to_screen
        n_set = len(se["settings"])
        if self._search_mode.get() == "all":
            found = [r for r in se["state"]["results"] if r is not None and r.get("found")]
            for k, r in enumerate(found):
                col = self._SEARCH_PALETTE[k % len(self._SEARCH_PALETTE)]
                for x, y in r["corners"]:
                    X, Y = ts(x, y)
                    c.create_oval(X - 1.5, Y - 1.5, X + 1.5, Y + 1.5, outline=col, fill=col, tags="ov")
            c.create_text(8, 8, anchor="nw", fill=SELECT_COLOR, font=("Arial", 9, "bold"), tags="ov",
                          text=f"All images: {sum(len(r['rc']) for r in found):,} corners in {len(found)} images "
                               f"(one colour per image)")
            return
        r = self._search_result(se["sel"])
        if r is None:
            return
        if not r.get("found"):
            c.create_text(8, 8, anchor="nw", fill=SELECT_COLOR, font=("Arial", 9, "bold"), tags="ov",
                          text=f"{os.path.basename(r['path'])}: no board found")
            return
        draw_lattice(c, ts, r["rc"], r["corners"], r["support"], n_set, True)
        c.create_text(8, 8, anchor="nw", fill=SELECT_COLOR, font=("Arial", 9, "bold"), tags="ov",
                      text=f"{os.path.basename(r['path'])}: {len(r['rc'])} corners, found by "
                           f"{len(r['settings_merged'])} of {n_set} settings")

    def _on_search_motion(self, u: float, v: float) -> None:
        se, iv = self._search, self._search_iv
        if se is None or iv is None or iv.view is None or self._search_mode.get() != "image":
            return
        r = self._search_result(se["sel"])
        if r is None or not r.get("found"):
            self._search_hover_var.set("")
            return
        d = np.hypot(r["corners"][:, 0] - u, r["corners"][:, 1] - v)
        k = int(np.argmin(d))
        if d[k] * iv.view[2] > 10:
            self._search_hover_var.set("")
            return
        rr, cc = (int(x) for x in r["rc"][k])
        sx, sy = self._square_xy()
        axes = "x"
        if sx != sy:
            try:
                axes = lattice_axes(r["rc"], (int(self._det_vars["ncol"].get()), int(self._det_vars["nrow"].get())))
            except ValueError:
                axes = None
        X, Y, _Z = lattice_obj(r["rc"][k:k + 1], sx, sy, axes)[0]
        where = f"object ({X:g}, {Y:g}, 0)" if axes is not None else "object: X/Y direction unknown"
        self._search_hover_var.set(
            f"corner row {rr}, col {cc}  ->  {where}   image ({r['corners'][k][0]:.2f}, "
            f"{r['corners'][k][1]:.2f})   found by {r['support'][k]} / {len(se['settings'])} settings   "
            f"spread {r['spread'][k]:.2f} px")

    def _on_search_copy_setting(self) -> None:
        se = self._search
        sel = self._search_set_tree.selection() if self._search_win is not None else ()
        if se is None or not sel:
            messagebox.showinfo("Search", "Select a setting in the list first.", parent=self._search_win)
            return
        st = se["settings"][int(sel[0])]
        for key, value in (("ncol", str(st.ncol)), ("nrow", str(st.nrow)), ("max_width", str(st.max_width))):
            self._det_vars[key].set(value)
        for key in ("adaptive", "normalize", "fast"):
            self._det_vars[key].set(getattr(st, key))
        messagebox.showinfo("Search", f"Copied '{st.label()}' to the Chessboard tab. Press Detect there to use "
                                      f"it for all images.", parent=self._search_win)

    def _on_search_apply(self) -> None:
        se = self._search
        if se is None:
            messagebox.showinfo("Search", "Run a search first.", parent=self._search_win)
            return
        if not se["state"]["finished"]:
            messagebox.showinfo("Search", "The search is still running: wait for it, or Cancel it first.",
                                parent=self._search_win)
            return
        found = [r for r in se["state"]["results"] if r is not None and r.get("found")]
        if not found:
            messagebox.showinfo("Search", "No board was found in any image.", parent=self._search_win)
            return
        ref = tuple(found[0]["size"])
        kept = [r for r in found if tuple(r["size"]) == ref]
        if self._cb and not messagebox.askokcancel(
                "Search", f"Replace the {len(self._cb['cands'])} board image(s) of the Chessboard tab with the "
                          f"{len(kept)} of the search?", parent=self._search_win):
            return
        n_set = len(se["settings"])
        cands = [{"path": r["path"], "source_index": int(r["source_index"]),
                  "corners": np.asarray(r["corners"], dtype=np.float32).reshape(-1, 2),
                  "rc": np.asarray(r["rc"], dtype=np.int32).reshape(-1, 2),
                  "support": np.asarray(r["support"], dtype=np.int32), "n_settings": n_set,
                  "sharpness": float(r["sharpness"]), "coverage": float(r["coverage"]),
                  "center_dist": float(r["center_dist"]), "size": ref} for r in kept]
        try:
            ncol, nrow = int(self._det_vars["ncol"].get()), int(self._det_vars["nrow"].get())
            max_width = int(self._det_vars["max_width"].get())
        except ValueError:
            ncol = nrow = max_width = 0
        sx, sy = self._square_xy()
        cfg = {"ncol": ncol, "nrow": nrow, "square": sx, "square_y": sy, "max_width": max_width,
               "search": dict(self._search_cfg)}
        self._cb = {"cfg": cfg, "cands": cands}
        self._cb_use = set()
        self._world = "points"
        if self._res is not None:
            self._res["cb_poses"] = {}
            self._mark_stale(refresh=False)
        self._show = None
        text = (f"Search: {sum(len(c['rc']) for c in cands):,} corners in {len(cands)} images "
                f"({n_set} settings each).")
        if len(kept) < len(found):
            text += f" {len(found) - len(kept)} image(s) with a different size than the first were dropped."
        self._set_cb_text(text)
        self._on_auto_select()
        self._on_search_close()

    def _update_cb_files_label(self) -> None:
        if self._cb_files_var is None:
            return
        if self._cb_pin_files:
            self._cb_files_var.set(f"{len(self._cb_pin_files)} files from the 'files' input pin")
        elif self._cb_manual_files:
            self._cb_files_var.set(f"{len(self._cb_manual_files)} picked files")
        else:
            self._cb_files_var.set("no files: connect 'files' or 'Add images...'")

    def _on_add_cb_files(self) -> None:
        paths = self.ask_open_filenames(parent=self._dlg_parent(), title="Chessboard images", filetypes=[
            ("Images", "*.png *.jpg *.jpeg *.bmp *.tif *.tiff *.webp"), ("All files", "*.*")])
        if paths:
            known = set(self._cb_manual_files)
            self._cb_manual_files += [os.path.normpath(p) for p in paths if os.path.normpath(p) not in known]
            self._update_cb_files_label()

    def _on_clear_cb_files(self) -> None:
        self._cb_manual_files = []
        self._update_cb_files_label()

    def _on_square_changed(self) -> None:
        if self._cb is not None and self._cb_use:
            self._mark_stale(refresh=False)
            self._refresh_all()

    def _set_detecting(self, running: bool) -> None:
        if self._cb_detect_btn is not None:
            self._cb_detect_btn.configure(state="disabled" if running else "normal",
                                          text="Detecting..." if running else "Detect")

    def _set_cb_text(self, text: str) -> None:
        self._cb_text = text
        if self._cb_text_var is not None:
            self._cb_text_var.set(text)

    def _on_detect(self) -> None:
        if self._cb_state is not None and not self._cb_state["finished"]:
            return
        files = self._cb_files()
        if not files:
            messagebox.showinfo("Detect", "Connect chessboard images to the 'files' input pin, or press "
                                          "'Add images...'.", parent=self._dlg_parent())
            return
        try:
            ncol, nrow = int(self._det_vars["ncol"].get()), int(self._det_vars["nrow"].get())
            max_width, threads = int(self._det_vars["max_width"].get()), int(self._det_vars["threads"].get())
            square = float(self._det_vars["square"].get())
            square_y = float(self._det_vars["square_y"].get() or square)
            if ncol < 2 or nrow < 2 or threads < 1 or square <= 0 or square_y <= 0 or max_width < 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("Detect", "columns / rows must be 2 or more, threads 1 or more, square X and Y "
                                           "more than 0, detect width 0 or more.", parent=self._dlg_parent())
            return
        flags = 0
        for key, bit in (("adaptive", cv2.CALIB_CB_ADAPTIVE_THRESH), ("normalize", cv2.CALIB_CB_NORMALIZE_IMAGE),
                         ("fast", cv2.CALIB_CB_FAST_CHECK)):
            if self._det_vars[key].get():
                flags |= bit
        cfg = {"ncol": ncol, "nrow": nrow, "square": square, "square_y": square_y, "detect_flags": flags,
               "max_width": max_width,
               "subpix_win": 0, "subpix_criteria": (cv2.TERM_CRITERIA_EPS + cv2.TERM_CRITERIA_MAX_ITER, 30, 0.001),
               "threads": threads}
        state = {"cancel": threading.Event(), "done": 0, "found": 0, "total": len(files),
                 "results": None, "error": None, "finished": False, "cfg": cfg}
        job = {"paths": [self._abs(p) for p in files], "indices": list(range(len(files))), "cfg": cfg}
        self._cb_state = state
        threading.Thread(target=_detect_worker, args=(job, state), daemon=True).start()
        self._set_detecting(True)
        self._set_cb_text(f"Detecting 0 / {len(files)} ...")
        self._poll_detect()

    def _on_stop_detect(self) -> None:
        if self._cb_state is not None and not self._cb_state["finished"]:
            self._cb_state["cancel"].set()

    def _poll_detect(self) -> None:
        st = self._cb_state
        if st is None:
            return
        if self._cb_progress_var is not None:
            try:
                self._cb_progress_var.set(100.0 * st["done"] / max(1, st["total"]))
            except tk.TclError:
                pass
        if not st["finished"]:
            self._set_cb_text(f"Detecting {st['done']} / {st['total']} ... boards found: {st['found']}")
            try:
                self.canvas.after(100, self._poll_detect)
            except tk.TclError:
                pass
            return
        self._set_detecting(False)
        if st["error"]:
            self._set_cb_text(f"Detection failed: {st['error']}")
            messagebox.showerror("Detect", st["error"], parent=self._dlg_parent())
            return
        results = [r for r in (st["results"] or []) if r is not None]
        found = [r for r in results if r.get("found")]
        dropped = 0
        if found:
            ref = tuple(found[0]["size"])
            kept = [r for r in found if tuple(r["size"]) == ref]
            dropped, found = len(found) - len(kept), kept
        text = (f"{'Stopped' if st['cancel'].is_set() else 'Done'}: {len(results)} of {st['total']} images "
                f"scanned, board found in {len(found)}.")
        if dropped:
            text += f" {dropped} image(s) with a different size than the first were dropped."
        self._set_cb_text(text)
        if not found:
            messagebox.showinfo("Detect", "No chessboard was detected. Check the INNER corner counts (one less "
                                          "than the squares; columns = per row, rows = per column) and the "
                                          "options.", parent=self._dlg_parent())
            return
        cands = [{"path": r["path"], "source_index": int(r["source_index"]),
                  "corners": np.asarray(r["corners"], dtype=np.float32).reshape(-1, 2),
                  "rc": cbs.grid_rc(*np.asarray(r["corners"]).shape[:2]), "sharpness": float(r["sharpness"]),
                  "coverage": float(r["coverage"]), "center_dist": float(r["center_dist"]),
                  "size": tuple(int(v) for v in r["size"])} for r in found]
        cfg = {k: st["cfg"][k] for k in ("ncol", "nrow", "square", "square_y", "max_width")}
        self._cb = {"cfg": cfg, "cands": cands}
        self._cb_use = set()
        if self._world != "points":
            self._world = "points"
        if self._res is not None:
            self._res["cb_poses"] = {}
            self._mark_stale(refresh=False)
        self._show = None
        self._on_auto_select()

    def _on_auto_select(self) -> None:
        if not self._cb:
            return
        try:
            target = int(self._sel_vars["target"].get())
            min_sharp = float(self._sel_vars["min_sharp"].get())
            diversity = float(self._sel_vars["diversity"].get())
        except ValueError:
            messagebox.showerror("Auto-select", "target count, min sharpness and min difference must be numbers.",
                                 parent=self._dlg_parent())
            return
        for c in self._cb["cands"]:
            c["found"] = True                    # rank_calibration_images() expects it
        ranked = rank_board_cands(list(self._cb["cands"]), min_sharp)
        ranked_paths = {c["path"] for c in ranked}
        rest = [c for c in self._cb["cands"] if c["path"] not in ranked_paths]
        for c in rest:
            c["quality"], c["pick_dist"] = None, None
        self._cb["cands"] = ranked + rest
        known = [c for c in ranked if self._cand_axes(c) is not None]
        chosen = take_ranked_images(known, target, diversity)
        self._cb_use = {c["path"] for c in chosen}
        text = (self._cb_text.split(" | ")[0] + f" | {len(chosen)} checked "
                f"(ranked {len(ranked)}, {len(rest)} below the min sharpness)")
        if len(known) < len(ranked):
            text += (f"; {len(ranked) - len(known)} image(s) 'X/Y?' not checked: square X differs from Y and "
                     f"their X direction cannot be told (see the square Y tooltip)")
        self._set_cb_text(text)
        self._mark_stale(refresh=False)
        self._refresh_all()

    def _set_all_cb(self, value: bool) -> None:
        if not self._cb:
            return
        self._cb_use = {c["path"] for c in self._cb["cands"]} if value else set()
        self._mark_stale(refresh=False)
        self._refresh_all()

    def _toggle_cb(self, path: str) -> None:
        if path in self._cb_use:
            self._cb_use.discard(path)
        else:
            self._cb_use.add(path)
        self._mark_stale(refresh=False)
        self._refresh_all()

    def _toggle_selected_cb(self) -> str:
        if self._cb_tree is not None:
            sel = self._cb_tree.selection()
            if sel:
                self._toggle_cb(self._tree_paths[int(sel[0])])
        return "break"

    def _on_tree_click(self, e) -> None:
        tree = self._cb_tree
        if tree.identify_region(e.x, e.y) != "cell":
            return
        item, col = tree.identify_row(e.y), tree.identify_column(e.x)
        if item and col == "#1":
            self._toggle_cb(self._tree_paths[int(item)])

    def _on_tree_select(self, _e=None) -> None:
        sel = self._cb_tree.selection() if self._cb_tree is not None else ()
        if sel:
            path = self._tree_paths[int(sel[0])]
            if path != self._show:
                self._show_board(path)

    def _select_tree_path(self, path: str) -> None:
        if self._cb_tree is None or path not in getattr(self, "_tree_paths", []):
            return
        iid = str(self._tree_paths.index(path))
        self._cb_tree.selection_set(iid)
        self._cb_tree.see(iid)

    def _refresh_cb_tree(self) -> None:
        if self._world_var is not None:
            wp = self._world_pose()
            self._world_var.set("World frame (rvec / tvec outputs): " + (
                wp[2] if wp is not None else ("board image " + os.path.basename(self._world)
                                              if self._world != "points" else "points image")))
        tree = self._cb_tree
        if tree is None:
            return
        keep = tree.selection()
        tree.delete(*tree.get_children())
        self._tree_paths = []
        if not self._cb:
            return
        rank = 0
        for k, c in enumerate(self._cb["cands"]):
            used = c["path"] in self._cb_use
            if c.get("pick_dist") is not None:
                rank += 1
            err = self._cb_err.get(c["path"])
            tags = ("world",) if self._world == c["path"] else (() if used else ("unused",))
            tree.insert("", "end", iid=str(k), tags=tags, values=(
                "✔" if used else "", rank if c.get("pick_dist") is not None else "-",
                c["source_index"] + 1, os.path.basename(c["path"]), f"{c['sharpness']:.0f}",
                f"{100 * c['coverage']:.1f}", "X/Y?" if self._cand_axes(c) is None else (
                    f"{err:.3f}" if err is not None and math.isfinite(err) else ""),
                "★" if self._world == c["path"] else ""))
            self._tree_paths.append(c["path"])
        # Re-select only the board image that is SHOWN: re-selecting another row
        # would fire <<TreeviewSelect>> and switch the image view to it (e.g.
        # right after a click on the points image).
        del keep
        if self._show is not None and self._show in self._tree_paths:
            tree.selection_set(str(self._tree_paths.index(self._show)))

    def _on_set_world(self) -> None:
        sel = self._cb_tree.selection() if self._cb_tree is not None else ()
        if not sel:
            messagebox.showinfo("World frame", "Select a board image in the table first.", parent=self._dlg_parent())
            return
        path = self._tree_paths[int(sel[0])]
        if path not in self._cb_use:
            messagebox.showinfo("World frame", "Check ('use') that image first: only checked images get a pose.",
                                parent=self._dlg_parent())
            return
        self._world = path
        if self._res is None or path not in self._res["cb_poses"]:
            messagebox.showinfo("World frame", "It becomes the world frame after the next Solve.",
                                parent=self._dlg_parent())
        self._refresh_all()
        self._push()

    def _on_world_points(self) -> None:
        self._world = "points"
        self._refresh_all()
        self._push()

    def _on_auto_initial(self) -> None:
        for v in self._init_vars.values():
            v.set("")
        self._mark_stale()

    def _set_initial(self, K, D) -> None:
        K = np.asarray(K, float)
        for n, val in zip(ce.INTR_NAMES, (K[0, 0], K[1, 1], K[0, 2], K[1, 2])):
            self._init_vars[n].set(f"{val:.6g}")
        D = np.asarray(D, float).ravel()
        for j, n in enumerate(ce.DIST_NAMES):
            self._init_vars[n].set(f"{D[j]:.8g}" if j < len(D) else "")
        self._mark_stale()

    def _on_result_to_initial(self) -> None:
        if self._res is None:
            messagebox.showinfo("Initial values", "There is no result yet.", parent=self._dlg_parent())
            return
        self._set_initial(self._res["K"], self._res["D"])

    def _on_pins_to_initial(self) -> None:
        K = None if self._pin_K is None else np.asarray(self._pin_K, float)
        if K is None or K.size != 9:
            messagebox.showinfo("Initial values", "No 3 x 3 camera matrix at the 'cmat in' input pin. Connect it "
                                                  "(and 'dvec in') and run the graph.", parent=self._dlg_parent())
            return
        D = np.zeros(0) if self._pin_D is None else np.asarray(self._pin_D, float).ravel()
        if len(D) > 5 and self._model_var.get() == "simple":
            self._model_var.set({8: "rational", 12: "thin_prism", 14: "tilted"}.get(len(D), "tilted"))
            self._build_param_grid()
        self._set_initial(K.reshape(3, 3), D)

    # ── dialogs / help ─────────────────────────────────────────────────────
    def _dlg_parent(self):
        win = self._inspector_win
        try:
            if win is not None and win.winfo_exists():
                return win
        except tk.TclError:
            pass
        return self.canvas.winfo_toplevel()

    def _on_export(self) -> None:
        path = self.ask_save_as_filename(parent=self._dlg_parent(), title="Export the point table",
                                         defaultextension=".xlsx",
                                         filetypes=[("Excel", "*.xlsx"), ("CSV", "*.csv")])
        if not path:
            return
        try:
            write_table(path, self.rows)
        except Exception as e:
            messagebox.showerror("Export", f"Cannot write the file:\n{e}", parent=self._dlg_parent())

    def _on_import(self) -> None:
        path = self.ask_open_filename(parent=self._dlg_parent(), title="Import a point table", filetypes=[
            ("Tables", "*.xlsx *.csv"), ("Excel", "*.xlsx"), ("CSV", "*.csv"), ("All files", "*.*")])
        if not path:
            return
        try:
            raw = read_table(path)
        except Exception as e:
            messagebox.showerror("Import", f"Cannot read the file:\n{e}", parent=self._dlg_parent())
            return
        replace = True
        if self.rows:
            ans = messagebox.askyesnocancel("Import", f"Replace the {len(self.rows)} rows of the table?\n\n"
                                                      "Yes = replace, No = append, Cancel = do nothing.",
                                            parent=self._dlg_parent())
            if ans is None:
                return
            replace = bool(ans)
        self._add_records(raw, "Import", replace=replace)

    def _open_help(self) -> None:
        if self._help_popup is not None:
            try:
                if self._help_popup.winfo_exists():
                    self._help_popup.lift()
                    return
            except tk.TclError:
                pass
        top = tk.Toplevel(self._dlg_parent())
        top.title("Camera Calibration - Help")
        top.geometry("780x680")
        txt = tk.Text(top, font=("Consolas", 9), wrap="word", padx=8, pady=6)
        vsb = ttk.Scrollbar(top, orient="vertical", command=txt.yview)
        txt.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        txt.pack(fill="both", expand=True)
        txt.insert("1.0", HELP_TEXT)
        txt.configure(state="disabled")
        top.bind("<Escape>", lambda _e: top.destroy())
        self._help_popup = top

    def get_help_text(self) -> str:
        return HELP_TEXT

    # ── project ────────────────────────────────────────────────────────────
    def _cb_pack(self) -> dict:
        """Detections as arrays (rank order); pick_dist -1 = not ranked."""
        cands = self._cb["cands"]
        return {"paths": np.array([self._rel(c["path"]) for c in cands], dtype=str),
                "source_index": np.array([c["source_index"] for c in cands], dtype=np.int64),
                "counts": np.array([len(c["rc"]) for c in cands], dtype=np.int64),
                "corners": np.concatenate([c["corners"].reshape(-1, 2) for c in cands]).astype(np.float32),
                "rc": np.concatenate([c["rc"] for c in cands]).astype(np.int32),
                "support": np.concatenate([np.zeros(len(c["rc"]), np.int32) if c.get("support") is None
                                           else c["support"] for c in cands]).astype(np.int32),
                "n_settings": np.array([c.get("n_settings") or 0 for c in cands], dtype=np.int64),
                "sharpness": np.array([c["sharpness"] for c in cands], dtype=np.float64),
                "coverage": np.array([c["coverage"] for c in cands], dtype=np.float64),
                "center_dist": np.array([c["center_dist"] for c in cands], dtype=np.float64),
                "size": np.array([c["size"] for c in cands], dtype=np.int64),
                "pick_dist": np.array([-1.0 if c.get("pick_dist") is None else c["pick_dist"] for c in cands]),
                "quality": np.array([-1.0 if c.get("quality") is None else c["quality"] for c in cands]),
                "use": np.array([c["path"] in self._cb_use for c in cands], dtype=bool),
                "cfg": np.array(json.dumps(self._cb["cfg"]))}

    _SEARCH_SCALARS = ("found", "partial", "error", "n_det", "best_single", "rejected", "n_lattices",
                       "settings_merged", "sharpness", "coverage", "center_dist", "size", "source_index")

    def _search_pack(self) -> tuple:
        """(JSON-able meta, arrays) of the settings search: settings, image
        list, per-image results (corners etc. concatenated in the arrays)."""
        se = self._search
        st = se["state"]
        results = list(st["results"])
        per, counts = [], []
        corners, rc, support, spread = [], [], [], []
        for r in results:
            if r is None:
                per.append(None)
                counts.append(0)
                continue
            per.append({k: (list(r[k]) if isinstance(r.get(k), (tuple, list)) else
                            (r[k].item() if hasattr(r.get(k), "item") else r[k]))
                        for k in self._SEARCH_SCALARS if k in r})
            n = len(r["rc"]) if r.get("found") else 0
            counts.append(n)
            if n:
                corners.append(np.asarray(r["corners"], np.float32).reshape(-1, 2))
                rc.append(np.asarray(r["rc"], np.int32).reshape(-1, 2))
                support.append(np.asarray(r["support"], np.int32))
                spread.append(np.asarray(r["spread"], np.float32))
        end = se["t_end"] if se["t_end"] is not None else time.monotonic()
        meta = {"settings": [dataclasses.asdict(x) for x in se["settings"]],
                "paths": [self._rel(p) for p in se["paths"]], "results": per,
                "setting_hits": [list(h) for h in st["setting_hits"]], "done": st["done"], "total": st["total"],
                "cancelled": bool(st["cancel"].is_set() or not st["finished"]), "elapsed": end - se["t0"]}
        arrays = {"counts": np.asarray(counts, np.int64),
                  "corners": np.concatenate(corners) if corners else np.zeros((0, 2), np.float32),
                  "rc": np.concatenate(rc) if rc else np.zeros((0, 2), np.int32),
                  "support": np.concatenate(support) if support else np.zeros(0, np.int32),
                  "spread": np.concatenate(spread) if spread else np.zeros(0, np.float32)}
        return meta, arrays

    def _search_unpack(self, meta: dict, arrays: dict) -> None:
        settings = [cbs.Setting(**x) for x in meta["settings"]]
        paths = [self._abs(p) for p in meta["paths"]]
        state = cbs.new_search_state(len(paths), len(settings))
        ends = np.cumsum(np.asarray(arrays["counts"], np.int64))
        starts = np.concatenate([[0], ends[:-1]])
        for i, r in enumerate(meta["results"]):
            if r is None:
                continue
            res = dict(r, path=paths[i])
            if "size" in res:
                res["size"] = tuple(res["size"])
            a, b = int(starts[i]), int(ends[i])
            if res.get("found"):
                res.update(corners=np.asarray(arrays["corners"][a:b], np.float32),
                           rc=np.asarray(arrays["rc"][a:b], np.int32),
                           support=np.asarray(arrays["support"][a:b], np.int32),
                           spread=np.asarray(arrays["spread"][a:b], np.float32))
            state["results"][i] = res
            state["live"][i] = b - a
            state["images_done"] += 1
        state["setting_hits"] = [list(h) for h in meta.get("setting_hits", state["setting_hits"])]
        state["done"], state["total"] = int(meta.get("done", 0)), int(meta.get("total", state["total"]))
        if meta.get("cancelled"):
            state["cancel"].set()
        state["finished"] = True
        t0 = time.monotonic()
        self._search = {"settings": settings, "paths": paths, "state": state, "t0": t0 - float(meta.get("elapsed", 0)),
                        "t_end": t0, "rows": set(), "sel": None}

    def _cb_unpack(self, d) -> None:
        cfg = json.loads(str(d["cfg"]))
        cands = []
        if "counts" in d:                                  # (total, 2) corners + lattice indices
            ends = np.cumsum(np.asarray(d["counts"], dtype=np.int64))
            spans = list(zip(np.concatenate([[0], ends[:-1]]).tolist(), ends.tolist()))
        for k in range(len(d["paths"])):
            pick, qual = float(d["pick_dist"][k]), float(d["quality"][k])
            if "counts" in d:
                a, b = spans[k]
                corners = np.asarray(d["corners"][a:b], dtype=np.float32).reshape(-1, 2)
                rc = np.asarray(d["rc"][a:b], dtype=np.int32).reshape(-1, 2)
                n_set = int(d["n_settings"][k])
                support = np.asarray(d["support"][a:b], dtype=np.int32) if n_set > 0 else None
            else:                                          # older files: (n, nrow, ncol, 2) full grids
                g = np.asarray(d["corners"][k], dtype=np.float32)
                corners, rc, n_set, support = g.reshape(-1, 2), cbs.grid_rc(*g.shape[:2]), 0, None
            cands.append({"path": self._abs(str(d["paths"][k])), "source_index": int(d["source_index"][k]),
                          "corners": corners, "rc": rc, "support": support, "n_settings": n_set or None,
                          "sharpness": float(d["sharpness"][k]), "coverage": float(d["coverage"][k]),
                          "center_dist": float(d["center_dist"][k]),
                          "size": tuple(int(v) for v in d["size"][k]),
                          "pick_dist": None if pick < 0 else pick, "quality": None if qual < 0 else qual})
        self._cb = {"cfg": cfg, "cands": cands} if cands else None
        self._cb_use = {c["path"] for c, u in zip(cands, d["use"]) if bool(u)}

    def _res_to_params(self, res: dict, cb_index: dict) -> dict:
        """A result for the project file: unknowns keyed by ROW POSITION in the
        saved table (row ids are not saved), board poses by the image's
        position in the detection list."""
        out = dict(res)
        ids = {r["id"]: k for k, r in enumerate(self.rows)}
        out["unknowns"] = {str(ids[rid]): list(v) for rid, v in (res.get("unknowns") or {}).items() if rid in ids}
        out["cb_poses"] = {str(cb_index[p]): v for p, v in (res.get("cb_poses") or {}).items() if p in cb_index}
        if out.get("solved"):
            out["solved"] = dict(out["solved"])
            out["solved"]["cb_poses"] = {str(cb_index[p]): v for p, v in out["solved"]["cb_poses"].items()
                                         if p in cb_index}
        return out

    def _res_from_params(self, res: dict, cands: list) -> dict:
        """The inverse of _res_to_params (rows and detections must be restored first)."""
        res = dict(res)
        if "points_pose" not in res:                       # saved by the first version
            res["points_pose"] = [res.pop("rvec"), res.pop("tvec")] if "rvec" in res else None

        def poses_in(d):
            out = {}
            for k, v in (d or {}).items():
                try:
                    out[cands[int(k)]["path"]] = v
                except (IndexError, ValueError):
                    pass
            return out
        res["cb_poses"] = poses_in(res.get("cb_poses"))
        if isinstance(res.get("solved"), dict):
            res["solved"] = dict(res["solved"])
            res["solved"]["cb_poses"] = poses_in(res["solved"].get("cb_poses"))
        unk = {}
        for pos, v in (res.get("unknowns") or {}).items():
            try:
                unk[self.rows[int(pos)]["id"]] = (int(v[0]), float(v[1]))
            except (IndexError, ValueError, TypeError):
                pass
        res["unknowns"] = unk
        return res

    def get_params(self) -> dict:
        params = {
            "model": self._model_var.get(),
            "fixed": sorted(n for n, v in self._fix_vars.items() if v.get()),
            "initial": {n: v.get() for n, v in self._init_vars.items() if v.get().strip()},
            "fix_aspect": bool(self._aspect_var.get()),
            "weights": {t: v.get() for t, v in self._w_vars.items()},
            "loss": self._loss_var.get(), "huber_px": self._huber_var.get(),
            "image_size": [v.get() for v in self._size_vars],
            "follow_pins": bool(self._follow_var.get()),
            "image_path": self._image_path,
            "view": {"names": bool(self._show_names_var.get()), "proj": bool(self._show_proj_var.get()),
                     "lines": bool(self._show_lines_var.get()), "axes": bool(self._show_axes_var.get()),
                     "gain": self._gain_var.get(), "snap": bool(self._snap_var.get()),
                     "line": self._line_var.get()},
        }
        params["cb_settings"] = {
            "det": {k: v.get() for k, v in self._det_vars.items()},
            "sel": {k: v.get() for k, v in self._sel_vars.items()},
            "use_points": bool(self._use_points_var.get()), "use_cb": bool(self._use_cb_var.get()),
            "combine": self._combine_var.get(), "files": [self._rel(f) for f in self._cb_manual_files],
            "force_lsq": bool(self._force_lsq_var.get()),
            "search": dict(self._search_cfg)}
        params["adjust"] = [{k: s[k] for k in ("param", "lo", "hi")} for s in self._adj_rows]
        params["grids"] = [{k: g[k] for k in ("on", "name", "xs", "ys", "zs", "frame", "color")} for g in self._grids]
        params["grid_view"] = {"on": bool(self._grids_on_var.get()), "width": self._grid_width_var.get()}
        params["views"] = {"v3_mode": self._v3_mode_var.get(), "v3_persp": float(self._v3_persp_var.get()),
                           "v3_cam": float(self._v3_cam_var.get()), "und_alpha": float(self._und_alpha_var.get()),
                           "und_pts": bool(self._und_pts_var.get()), "loo_boards": bool(self._loo_boards_var.get()),
                           "loo_points": bool(self._loo_points_var.get())}
        runs_out = []
        cb_index = {c["path"]: k for k, c in enumerate(self._cb["cands"])} if self._cb else {}
        params["world"] = cb_index.get(self._world, "points")
        # Everything goes into the project file: the detections (exact) and a
        # settings search's results in its Data sheet (project_data).
        if self._cb:
            params["cb_data"] = project_data.put_arrays(self, "chessboard", self._cb_pack())
        if self._search is not None:
            meta, arrays = self._search_pack()
            params["search_result"] = meta
            params["search_data"] = project_data.put_arrays(self, "search", arrays)
        if self._res is not None:
            params["result"] = self._res_to_params(self._res, cb_index)
        # Every run with its full result (poses ...), so 'Restore' works after loading too.
        for r in self._runs:
            out = {k: v for k, v in r.items() if k != "full"}
            if r.get("full"):
                out["full"] = self._res_to_params(r["full"], cb_index)
            runs_out.append(out)
        params["runs"] = runs_out
        # Sensitivity (leave-one-out) results: keys by image position / table row position.
        if self._loo_rows:
            ids = {r["id"]: k for k, r in enumerate(self.rows)}
            rows = []
            for row in self._loo_rows:
                row = dict(row)
                if row.get("kind") == "board":
                    row["key"] = cb_index.get(row.get("key"))
                elif row.get("kind") == "row":
                    row["key"] = ids.get(row.get("key"))
                rows.append(row)
            params["loo"] = {"rows": rows, "base": self._loo_base, "text": self._loo_text}
        compact = [[r["name"], r["X"], r["Y"], r["Z"], None if not math.isfinite(r["u"]) else r["u"],
                    None if not math.isfinite(r["v"]) else r["v"], r["line"], r["weight"], r["use"]]
                   for r in self.rows]
        # However long: long params go to the project file's Data sheet
        # (project_io.py); 'rows_file' CSV side files of older projects are still read.
        params["rows"] = compact
        return params

    def set_params(self, params: dict) -> None:
        self._model_var.set(params.get("model", "simple") if params.get("model") in ce.DIST_MODELS else "simple")
        fixed = set(params.get("fixed", DEFAULT_FIXED))
        for n, v in self._fix_vars.items():
            v.set(n in fixed)
        init = params.get("initial", {})
        for n, v in self._init_vars.items():
            v.set(str(init.get(n, "")))
        self._aspect_var.set(bool(params.get("fix_aspect", False)))
        for t, v in self._w_vars.items():
            v.set(str(params.get("weights", {}).get(t, "1")))
        self._loss_var.set(params.get("loss", "linear") if params.get("loss") in ("linear", "huber") else "linear")
        self._huber_var.set(str(params.get("huber_px", "2")))
        size = params.get("image_size", ["", ""])
        for k in range(2):
            self._size_vars[k].set(str(size[k]) if k < len(size) else "")
        self._follow_var.set(bool(params.get("follow_pins", False)))
        view = params.get("view", {})
        self._show_names_var.set(bool(view.get("names", True)))
        self._show_proj_var.set(bool(view.get("proj", True)))
        self._show_lines_var.set(bool(view.get("lines", True)))
        self._show_axes_var.set(bool(view.get("axes", True)))
        self._gain_var.set(str(view.get("gain", "10")))
        self._snap_var.set(bool(view.get("snap", False)))
        self._line_var.set(str(view.get("line", "L1")))
        self._image_path = str(params.get("image_path", "") or "")
        self._file_image = None
        if self._image_path and not self._load_image_file(self._image_path):
            self._file_image = None

        self._next_id = 1
        rows = []
        if params.get("rows_file"):
            try:
                rows, errors = parse_records(read_table(str(Path(get_project_directory()) / params["rows_file"])),
                                             self._new_id, lambda: "P")
                if errors:
                    raise ValueError(errors[0])
            except Exception:
                rows = []
        else:
            for rec in params.get("rows", []):
                name, X, Y, Z, u, v, line, w, use = (list(rec) + [None] * 9)[:9]
                rows.append(new_row(self._new_id(), name or "P", X or "", Y or "", Z or "",
                                    math.nan if u is None else u, math.nan if v is None else v,
                                    line or "", 1.0 if w is None else w, True if use is None else use))
        self.rows = rows

        self._adj_rows = [{"param": str(a.get("param", "fx")), "lo": str(a.get("lo", "")), "hi": str(a.get("hi", ""))}
                          for a in params.get("adjust", []) if isinstance(a, dict)]
        self._grids = []
        for g in params.get("grids", []):
            if isinstance(g, dict):
                self._grids.append({"on": bool(g.get("on", True)), "name": str(g.get("name", "G")),
                                    "xs": str(g.get("xs", "0")), "ys": str(g.get("ys", "0")),
                                    "zs": str(g.get("zs", "0")),
                                    "frame": g.get("frame") if g.get("frame") in ("points", "board") else "points",
                                    "color": g.get("color") if g.get("color") in GRID_COLORS else "cyan"})
        gv = params.get("grid_view", {})
        vw = params.get("views", {})
        try:
            if vw.get("v3_mode") in ("points image (object frame)", "board frame (all boards)", "camera frame"):
                self._v3_mode_var.set(vw["v3_mode"])
            self._v3_persp_var.set(float(vw.get("v3_persp", 0.3)))
            self._v3_cam_var.set(float(vw.get("v3_cam", 0.3)))
            self._und_alpha_var.set(float(vw.get("und_alpha", 0.0)))
        except (TypeError, ValueError):
            pass
        self._und_pts_var.set(bool(vw.get("und_pts", True)))
        self._loo_boards_var.set(bool(vw.get("loo_boards", True)))
        self._loo_points_var.set(bool(vw.get("loo_points", True)))
        self._runs = [dict(r) for r in params.get("runs", []) if isinstance(r, dict) and "K" in r and "D" in r]
        self._grids_on_var.set(bool(gv.get("on", True)))
        self._grid_width_var.set(str(gv.get("width", "1")))
        cbs = params.get("cb_settings", {})
        for k, v in cbs.get("det", {}).items():
            if k in self._det_vars:
                self._det_vars[k].set(v)
        if "square_y" not in cbs.get("det", {}):                # saved before square X / Y
            self._det_vars["square_y"].set("")
        for k, v in cbs.get("sel", {}).items():
            if k in self._sel_vars:
                self._sel_vars[k].set(v)
        self._use_points_var.set(bool(cbs.get("use_points", True)))
        self._use_cb_var.set(bool(cbs.get("use_cb", True)))
        self._combine_var.set(cbs.get("combine", "joint") if cbs.get("combine") in ("joint", "two_step") else "joint")
        self._force_lsq_var.set(bool(cbs.get("force_lsq", False)))
        self._cb_manual_files = [self._abs(f) for f in cbs.get("files", [])]
        for k, v in (cbs.get("search") or {}).items():
            if k in self._search_cfg and (k not in self._SEARCH_FLAG_CHOICES_KEYS or v in self._SEARCH_FLAG_CHOICES):
                self._search_cfg[k] = str(v)
        self._cb, self._cb_use = None, set()
        try:
            if params.get("cb_data") is not None:
                d = project_data.get_arrays(params["cb_data"])
                if d is None:
                    raise ValueError("they are missing from the project file")
                self._cb_unpack(d)
            elif params.get("cb_file"):                       # projects saved before the Data sheet
                with np.load(str(Path(get_project_directory()) / params["cb_file"]), allow_pickle=False) as d:
                    self._cb_unpack({k: d[k] for k in d.files})
            elif params.get("cb_inline"):
                self._cb_unpack({k: np.asarray(v) for k, v in params["cb_inline"].items()})
        except Exception as e:
            self._cb, self._cb_use = None, set()
            self._cb_text = f"The saved chessboard detections could not be read: {e}"
        self._search = None
        if isinstance(params.get("search_result"), dict):
            try:
                arrays = project_data.get_arrays(params.get("search_data"))
                if arrays is None:
                    raise ValueError("its data is missing from the project file")
                self._search_unpack(params["search_result"], arrays)
            except Exception as e:
                self._search = None
                self._cb_text = (self._cb_text + " | " if self._cb_text else "") + \
                    f"The saved settings-search result could not be read: {e}"
        cands = self._cb["cands"] if self._cb else []
        w = params.get("world", "points")
        self._world = cands[w]["path"] if isinstance(w, int) and 0 <= w < len(cands) else "points"
        self._show = None

        res = params.get("result")
        self._res = None
        if isinstance(res, dict) and "K" in res:
            self._res = self._res_from_params(res, cands)
        for r in self._runs:
            if isinstance(r.get("full"), dict) and "K" in r["full"]:
                r["full"] = self._res_from_params(r["full"], cands)
            else:
                r.pop("full", None)
        loo = params.get("loo")
        self._loo_rows, self._loo_base, self._loo_text = [], None, ""
        if isinstance(loo, dict):
            for row in loo.get("rows", []):
                row = dict(row)
                try:
                    if row.get("kind") == "board":
                        row["key"] = cands[int(row["key"])]["path"]
                    elif row.get("kind") == "row":
                        row["key"] = self.rows[int(row["key"])]["id"]
                except (IndexError, TypeError, ValueError, KeyError):
                    row["key"] = None
                self._loo_rows.append(row)
            self._loo_base = loo.get("base")
            self._loo_text = str(loo.get("text", ""))
        self._on_image_changed()
        self._refresh_all()
