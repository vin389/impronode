# node_editor/nodes/multi_camera_points_node.py
"""
Multi-Camera Points: one table of NAMED points seen by several cameras.

Purpose
    For multi-camera measurement each physical point needs its image
    coordinates in every camera, in the SAME row order for all cameras
    (Point Tracking, Stereo Sync and triangulation work by row index). This
    node keeps that table in one place:
      * ArUco markers are detected in all camera images and named
        automatically (ARUCO_<dictionary>_ID<id>), so the same marker gets
        the same row in every camera -- no manual mapping. A marker is
        represented by its CENTRE = intersection of its two diagonals (the
        true centre under perspective, independent of the corner order).
      * Other points (corner stickers, speckles, strokes ...) are added by
        clicking in the camera views and get names P0001, P0002, ...
      * The table can be edited in the inspector, and exported to / imported
        from CSV or Excel to be edited in a spreadsheet.

Table
    One row per point: name, type, and per camera x, y, status. A cell
    without coordinates is NaN. Status values:
        missing      not done yet                      (red)
        auto         found by the ArUco detector       (green)
        suggested    proposed automatically, unchecked (yellow; later versions)
        confirmed    placed / checked by the user      (blue)
        not_visible  the user says this camera cannot see it (grey, NaN)
    Names: letters, digits and underscore only, starting with a letter (safe
    in Python, MATLAB, Excel and file names). Row order is kept: new rows are
    appended, never re-sorted (downstream results are stored by row).

Coordinates follow OpenCV: (0, 0) is the centre of the top-left pixel.
"""

from __future__ import annotations

import csv
import math
import os
import re
import tkinter as tk
from pathlib import Path
from tkinter import messagebox, ttk

import cv2
import numpy as np
from PIL import Image, ImageTk

from node_editor.base_node import BaseNode
from node_editor.execution import ExecutionMode
from node_editor.pin_types import PinDef, PinSchema, PinType
from node_editor.project_context import get_project_directory
from node_editor.tooltip import add_tooltip
from node_editor.nodes.image_display_nodes import ARUCO_DICT_NAMES, _aruco_center

MAX_CAMS = 6
STATUSES = ("missing", "auto", "suggested", "confirmed", "not_visible")   # output code = index
VALID_STATUSES = ("auto", "suggested", "confirmed")                        # have coordinates
STATUS_COLORS = {"missing": "#ffd9d9", "auto": "#d6f2d6", "suggested": "#fff1b0",
                 "confirmed": "#dbe7ff", "not_visible": "#e4e4e4"}
POINT_COLORS = {"auto": "#16a34a", "suggested": "#e69500", "confirmed": "#2563eb"}
SELECT_COLOR = "#ff2db4"
NAME_RE = re.compile(r"[A-Za-z][A-Za-z0-9_]*")


# ── table rows ────────────────────────────────────────────────────────────

def new_row(name: str, kind: str = "manual", meta: dict | None = None) -> dict:
    return {"name": name, "kind": kind, "meta": dict(meta or {}),
            "x": [math.nan] * MAX_CAMS, "y": [math.nan] * MAX_CAMS, "status": ["missing"] * MAX_CAMS}


def cell_valid(row: dict, k: int) -> bool:
    return (row["status"][k] in VALID_STATUSES
            and math.isfinite(row["x"][k]) and math.isfinite(row["y"][k]))


def aruco_name(dict_name: str, marker_id: int) -> str:
    """ARUCO_6X6_1000_ID0123 for DICT_6X6_1000 / id 123 (letters, digits, _)."""
    short = dict_name[5:] if dict_name.startswith("DICT_") else dict_name
    return f"ARUCO_{short.upper()}_ID{int(marker_id):04d}"


def detect_aruco(image: np.ndarray, dict_name: str) -> list[tuple[int, np.ndarray, np.ndarray]]:
    """[(id, 4x2 corners, centre)] of the markers of `dict_name` in an RGB /
    grey image, sorted by id. Centre = diagonal intersection."""
    gray = image if image.ndim == 2 else cv2.cvtColor(image, cv2.COLOR_RGB2GRAY)
    if gray.dtype != np.uint8:
        gray = cv2.normalize(gray, None, 0, 255, cv2.NORM_MINMAX).astype(np.uint8)
    dictionary = cv2.aruco.getPredefinedDictionary(getattr(cv2.aruco, dict_name))
    params = cv2.aruco.DetectorParameters()
    params.cornerRefinementMethod = cv2.aruco.CORNER_REFINE_SUBPIX
    corners, ids, _rej = cv2.aruco.ArucoDetector(dictionary, params).detectMarkers(gray)
    out = []
    if ids is not None:
        for quad, mid in zip(corners, ids.ravel()):
            q = np.asarray(quad, dtype=np.float64).reshape(4, 2)
            out.append((int(mid), q, _aruco_center(q)))
    return sorted(out, key=lambda m: m[0])


# ── file export / import (CSV and Excel share one row format) ─────────────

def table_header(n_cams: int) -> list[str]:
    head = ["name", "type"]
    for k in range(1, n_cams + 1):
        head += [f"cam{k}_x", f"cam{k}_y", f"cam{k}_status"]
    return head


def _fmt(v: float) -> str:
    return "" if not math.isfinite(v) else repr(round(float(v), 6))


def table_records(rows: list[dict], n_cams: int) -> list[list]:
    """Rows as plain lists for CSV / Excel: NaN -> empty cell."""
    out = []
    for r in rows:
        rec = [r["name"], r["kind"]]
        for k in range(n_cams):
            rec += [_fmt(r["x"][k]), _fmt(r["y"][k]), r["status"][k]]
        out.append(rec)
    return out


def write_csv(path: str, rows: list[dict], n_cams: int) -> None:
    tmp = Path(path).with_name(Path(path).name + ".tmp")
    with open(tmp, "w", newline="", encoding="utf-8-sig") as f:      # BOM: Excel opens it as UTF-8
        w = csv.writer(f)
        w.writerow(table_header(n_cams))
        w.writerows(table_records(rows, n_cams))
    tmp.replace(path)


def write_xlsx(path: str, rows: list[dict], n_cams: int) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "points"
    ws.append(table_header(n_cams))
    for c in ws[1]:
        c.font = Font(bold=True)
    for r, rec in zip(rows, table_records(rows, n_cams)):
        # x / y as numbers (empty cell = no coordinate); name, type, status as text.
        ws.append([(float(v) if v else None) if (i >= 2 and (i - 2) % 3 != 2) else v
                   for i, v in enumerate(rec)])
        row_cells = ws[ws.max_row]
        for k in range(n_cams):
            fill = PatternFill("solid", fgColor=STATUS_COLORS[r["status"][k]].lstrip("#"))
            for c in row_cells[2 + 3 * k: 5 + 3 * k]:
                c.fill = fill
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 28
    readme = wb.create_sheet("README")
    for line in (
        ["Multi-Camera Points table"],
        [""],
        ["name", "point name: letters, digits and _ only, starting with a letter; unique"],
        ["type", "aruco / manual (information only)"],
        ["camK_x, camK_y", "image coordinates in camera K (OpenCV: (0, 0) = centre of the top-left pixel); empty = none"],
        ["camK_status", "missing / auto / suggested / confirmed / not_visible"],
        [""],
        ["Rows may be edited, added, deleted or re-ordered. Import the file back into the node;"],
        ["it shows a summary of the changes before applying them. The row order is the order"],
        ["downstream nodes use: re-ordering means results computed before must be recomputed."],
    ):
        readme.append(line)
    readme.column_dimensions["A"].width = 18
    tmp = Path(path).with_name(Path(path).name + ".tmp")
    wb.save(tmp)
    tmp.replace(path)


def read_table_file(path: str) -> list[list]:
    """All rows (header first) of a CSV or .xlsx file, as lists of strings."""
    if Path(path).suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["points"] if "points" in wb.sheetnames else wb.worksheets[0]
        rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    with open(path, newline="", encoding="utf-8-sig") as f:
        return [list(r) for r in csv.reader(f)]


def parse_table(raw: list[list]) -> tuple[list[dict], int, list[str]]:
    """Header + rows -> (rows, number of cameras, errors). Errors name the
    line and column; when there are any, the rows must not be used."""
    errors: list[str] = []
    if not raw:
        return [], 0, ["the file is empty"]
    head = [h.strip().lower() for h in raw[0]]
    if "name" not in head:
        return [], 0, ["no 'name' column in the first line"]
    i_name, i_type = head.index("name"), (head.index("type") if "type" in head else None)
    cams: dict[int, dict[str, int]] = {}
    for i, h in enumerate(head):
        m = re.fullmatch(r"cam(\d+)_(x|y|status)", h)
        if m:
            cams.setdefault(int(m.group(1)), {})[m.group(2)] = i
    if not cams:
        return [], 0, ["no camK_x / camK_y columns"]
    if max(cams) > MAX_CAMS:
        return [], 0, [f"at most {MAX_CAMS} cameras are supported (found cam{max(cams)})"]
    n_cams = max(cams)
    rows, seen = [], set()
    for line_no, rec in enumerate(raw[1:], start=2):
        rec = list(rec) + [""] * (len(head) - len(rec))
        if not any(str(v).strip() for v in rec):
            continue                                            # blank line
        name = str(rec[i_name]).strip()
        if not NAME_RE.fullmatch(name):
            errors.append(f"line {line_no}: name '{name}' must use letters, digits and _ only, "
                          f"starting with a letter")
            continue
        if name in seen:
            errors.append(f"line {line_no}: name '{name}' appears twice")
            continue
        seen.add(name)
        row = new_row(name, str(rec[i_type]).strip() or "manual" if i_type is not None else "manual")
        for k, cols in cams.items():
            j = k - 1
            try:
                x = float(rec[cols["x"]]) if "x" in cols and str(rec[cols["x"]]).strip() else math.nan
                y = float(rec[cols["y"]]) if "y" in cols and str(rec[cols["y"]]).strip() else math.nan
            except ValueError:
                errors.append(f"line {line_no}: cam{k} x / y is not a number")
                continue
            st = str(rec[cols["status"]]).strip().lower() if "status" in cols else ""
            has_xy = math.isfinite(x) and math.isfinite(y)
            if not st:
                st = "confirmed" if has_xy else "missing"
            if st not in STATUSES:
                errors.append(f"line {line_no}: cam{k}_status '{st}' is not one of {', '.join(STATUSES)}")
                continue
            if st in VALID_STATUSES and not has_xy:
                errors.append(f"line {line_no}: cam{k} is '{st}' but has no x / y")
                continue
            if st in ("missing", "not_visible"):
                x = y = math.nan
            row["x"][j], row["y"][j], row["status"][j] = x, y, st
        rows.append(row)
    return rows, n_cams, errors


def change_summary(old: list[dict], new: list[dict]) -> list[str]:
    """Human-readable differences between two tables (for the Import dialog)."""
    old_by = {r["name"]: r for r in old}
    new_names = [r["name"] for r in new]
    added = [n for n in new_names if n not in old_by]
    removed = [r["name"] for r in old if r["name"] not in set(new_names)]
    changed = 0
    for r in new:
        o = old_by.get(r["name"])
        if o is None:
            continue
        for k in range(MAX_CAMS):
            # Exported coordinates are rounded to 6 decimals: compare at that precision.
            same_xy = all((math.isnan(a) and math.isnan(b)) or abs(a - b) < 1e-5
                          for a, b in ((o["x"][k], r["x"][k]), (o["y"][k], r["y"][k])))
            if not same_xy or o["status"][k] != r["status"][k]:
                changed += 1
    lines = [f"{len(new)} rows in the file (now {len(old)})",
             f"{changed} camera cell(s) changed"]
    if added:
        lines.append(f"{len(added)} new row(s): " + ", ".join(added[:8]) + (" ..." if len(added) > 8 else ""))
    if removed:
        lines.append(f"{len(removed)} row(s) removed: " + ", ".join(removed[:8]) + (" ..." if len(removed) > 8 else ""))
    kept = [n for n in new_names if n in old_by]
    if kept != [r["name"] for r in old if r["name"] in set(kept)]:
        lines.append("WARNING: the row order changed -- results computed before (tracking, ...) "
                     "must be recomputed.")
    return lines


# ── widgets ───────────────────────────────────────────────────────────────

class CamView:
    """One camera's zoomable image with the table's points drawn on it.
    Mouse: wheel = zoom at the cursor, right-drag = pan, left-click =
    on_pick(cam, u, v). Screen x of image column u = ox + (u + 0.5) * s, so
    integer (u, v) is a pixel centre (OpenCV convention)."""

    PREVIEW_MAX = 1600            # longest side of the cached downscaled copy

    def __init__(self, parent, cam: int, on_pick):
        self.cam, self.on_pick = cam, on_pick
        self.frame = tk.LabelFrame(parent, text=f"Camera {cam + 1}", font=("Arial", 9, "bold"))
        self.canvas = tk.Canvas(self.frame, bg="#1b1b24", highlightthickness=0, width=300, height=220,
                                cursor="crosshair")
        self.canvas.pack(fill="both", expand=True)
        self.image = None
        self.preview, self.k = None, 1.0              # downscaled copy and its scale
        self.view = None                              # (ox, oy, s)
        self.user_view = False
        self.serial = 0
        self._photo, self._photo_key = None, None
        self._press = None
        self.markers: list = []                       # detected ArUco: (id, quad, centre)
        self.points: list = []                        # (u, v, status, name, selected)
        self.show_names = True
        c = self.canvas
        c.bind("<Configure>", lambda _e: self.render())
        c.bind("<MouseWheel>", self._on_wheel)
        c.bind("<Button-4>", self._on_wheel)
        c.bind("<Button-5>", self._on_wheel)
        c.bind("<ButtonPress-3>", self._on_pan_start)
        c.bind("<B3-Motion>", self._on_pan)
        c.bind("<ButtonPress-1>", lambda e: setattr(self, "_press", (e.x, e.y)))
        c.bind("<ButtonRelease-1>", self._on_click)

    # image / view
    def set_image(self, img) -> None:
        if img is self.image:
            return
        if img is not None:
            a = np.asarray(img)
            if a.dtype != np.uint8:
                a = cv2.normalize(a.astype(np.float64), None, 0, 255, cv2.NORM_MINMAX).astype(np.uint8)
            if a.ndim == 2:
                a = np.repeat(a[:, :, None], 3, axis=2)
            elif a.ndim == 3 and a.shape[2] == 4:
                a = a[:, :, :3]
            old = None if self.image is None else self.image.shape[:2]
            self.image = np.ascontiguousarray(a)
            h, w = self.image.shape[:2]
            f = min(1.0, self.PREVIEW_MAX / max(h, w))
            self.preview = self.image if f >= 1.0 else cv2.resize(
                self.image, (max(1, round(w * f)), max(1, round(h * f))), interpolation=cv2.INTER_AREA)
            self.k = self.preview.shape[1] / w
            if old != (h, w):
                self.user_view = False
        else:
            self.image = self.preview = None
        self.serial += 1
        self.render()

    def _fit(self, cw, ch):
        h, w = self.image.shape[:2]
        s = min((cw - 8) / w, (ch - 8) / h)
        return ((cw - w * s) / 2.0, (ch - h * s) / 2.0, s)

    def fit(self) -> None:
        self.user_view = False
        self.render()

    def to_screen(self, u, v):
        ox, oy, s = self.view
        return ox + (u + 0.5) * s, oy + (v + 0.5) * s

    def to_image(self, x, y):
        ox, oy, s = self.view
        return (x - ox) / s - 0.5, (y - oy) / s - 0.5

    def center_on(self, u: float, v: float, zoom_if_fit: float = 4.0) -> None:
        """Centre the view on (u, v); zoom in once if the view is still 'fit'."""
        c = self.canvas
        if self.image is None or self.view is None:
            return
        cw, ch = c.winfo_width(), c.winfo_height()
        s = self.view[2] if self.user_view else self._fit(cw, ch)[2] * zoom_if_fit
        self.view = (cw / 2.0 - (u + 0.5) * s, ch / 2.0 - (v + 0.5) * s, s)
        self.user_view = True
        self.render()

    # drawing
    def render(self) -> None:
        c = self.canvas
        try:
            cw, ch = c.winfo_width(), c.winfo_height()
        except tk.TclError:
            return
        if cw < 20 or ch < 20:
            return
        if self.image is None:
            c.delete("all")
            self._photo_key = None
            c.create_text(cw / 2, ch / 2, fill="#9aa0b5", justify="center",
                          text=f"No image.\nConnect 'img {self.cam + 1}' (the reference frame).")
            return
        if not self.user_view or self.view is None:
            self.view = self._fit(cw, ch)
        ox, oy, s = self.view
        h, w = self.image.shape[:2]
        # Render only the visible crop, from the downscaled copy when that
        # has enough resolution for the current zoom.
        src, pk = (self.preview, self.k) if s <= self.k else (self.image, 1.0)
        sh, sw = src.shape[:2]
        c0 = max(0, int(math.floor(-ox / s * pk)))
        c1 = min(sw, int(math.ceil((cw - ox) / s * pk)))
        r0 = max(0, int(math.floor(-oy / s * pk)))
        r1 = min(sh, int(math.ceil((ch - oy) / s * pk)))
        c.delete("ov")
        if c1 > c0 and r1 > r0:
            px, py = ox + c0 / pk * s, oy + r0 / pk * s
            dw = max(1, int(round(ox + c1 / pk * s - px)))
            dh = max(1, int(round(oy + r1 / pk * s - py)))
            key = (self.serial, c0, r0, c1, r1, dw, dh)
            if key != self._photo_key:
                crop = np.ascontiguousarray(src[r0:r1, c0:c1])
                interp = cv2.INTER_AREA if dw < crop.shape[1] else (
                    cv2.INTER_NEAREST if s >= 3.0 else cv2.INTER_LINEAR)
                self._photo = ImageTk.PhotoImage(Image.fromarray(cv2.resize(crop, (dw, dh), interpolation=interp)))
                self._photo_key = key
                c.delete("img")
                c.create_image(px, py, anchor="nw", image=self._photo, tags="img")
            else:
                c.coords("img", px, py)
        else:
            c.delete("img")
            self._photo_key = None
        self._draw_overlay()

    def _draw_overlay(self) -> None:
        c = self.canvas
        for _mid, quad, _cen in self.markers:                       # detected markers: thin yellow
            pts = [v for q in quad for v in self.to_screen(*q)]
            c.create_polygon(*pts, outline="#ffd23f", fill="", width=1, tags="ov")
        selected = []
        for u, v, status, name, sel in self.points:
            if sel:
                selected.append((u, v, name))
                continue
            x, y = self.to_screen(u, v)
            col = POINT_COLORS.get(status, "#ffffff")
            c.create_oval(x - 4, y - 4, x + 4, y + 4, outline=col, width=2, tags="ov")
            if self.show_names:
                c.create_text(x + 6, y - 6, text=name, anchor="sw", fill=col, font=("Arial", 7), tags="ov")
        for u, v, name in selected:                                 # selected on top
            x, y = self.to_screen(u, v)
            c.create_oval(x - 9, y - 9, x + 9, y + 9, outline=SELECT_COLOR, width=2, tags="ov")
            c.create_line(x - 14, y, x - 4, y, fill=SELECT_COLOR, width=1, tags="ov")
            c.create_line(x + 4, y, x + 14, y, fill=SELECT_COLOR, width=1, tags="ov")
            c.create_line(x, y - 14, x, y - 4, fill=SELECT_COLOR, width=1, tags="ov")
            c.create_line(x, y + 4, x, y + 14, fill=SELECT_COLOR, width=1, tags="ov")
            c.create_text(x + 11, y - 11, text=name, anchor="sw", fill=SELECT_COLOR,
                          font=("Arial", 8, "bold"), tags="ov")

    # mouse
    def _on_wheel(self, e) -> str:
        if self.image is None or self.view is None:
            return "break"
        steps = 1 if getattr(e, "num", None) == 4 else -1 if getattr(e, "num", None) == 5 else \
            (int(e.delta / 120) if e.delta else 0)
        if steps:
            ox, oy, s = self.view
            ns = min(max(s * (1.25 ** steps), 0.01), 64.0)
            f = ns / s
            self.view = (e.x - (e.x - ox) * f, e.y - (e.y - oy) * f, ns)
            self.user_view = True
            self.render()
        return "break"

    def _on_pan_start(self, e) -> None:
        if self.view is not None:
            self._pan0 = (e.x, e.y, self.view[0], self.view[1])

    def _on_pan(self, e) -> None:
        p = getattr(self, "_pan0", None)
        if p is None or self.view is None:
            return
        self.view = (p[2] + e.x - p[0], p[3] + e.y - p[1], self.view[2])
        self.user_view = True
        self.render()

    def _on_click(self, e) -> None:
        p, self._press = self._press, None
        if p is None or self.image is None or self.view is None or abs(e.x - p[0]) + abs(e.y - p[1]) > 4:
            return
        u, v = self.to_image(e.x, e.y)
        h, w = self.image.shape[:2]
        if -0.5 <= u <= w - 0.5 and -0.5 <= v <= h - 0.5:
            self.on_pick(self.cam, u, v, e)


class Sheet:
    """Spreadsheet-like table drawn on a canvas: one colour per cell (status),
    only the visible rows are drawn (fast for long tables), click selects
    (Ctrl / Shift for several rows), double-click edits a cell."""

    ROW_H = 20
    HEAD_H = 22

    def __init__(self, parent, get_rows, on_select, on_edit):
        self.get_rows, self.on_select, self.on_edit = get_rows, on_select, on_edit
        self.frame = tk.Frame(parent)
        self.canvas = tk.Canvas(self.frame, bg="#ffffff", highlightthickness=0,
                                yscrollincrement=self.ROW_H)
        vsb = ttk.Scrollbar(self.frame, orient="vertical", command=self._yview)
        hsb = ttk.Scrollbar(self.frame, orient="horizontal", command=self._xview)
        self.canvas.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        vsb.pack(side="right", fill="y")
        hsb.pack(side="bottom", fill="x")
        self.canvas.pack(fill="both", expand=True)
        self.columns: list[tuple[str, str, int]] = []
        self.selected: list[int] = []
        self.sel_cam: int | None = None
        self._anchor = None
        self._editor = None
        c = self.canvas
        c.bind("<Configure>", lambda _e: self.redraw())
        c.bind("<ButtonPress-1>", self._on_click)
        c.bind("<Double-Button-1>", self._on_double)
        for seq in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            c.bind(seq, self._on_wheel)

    def set_cameras(self, n_cams: int) -> None:
        cols = [("name", "name", 190), ("kind", "type", 62)]
        for k in range(n_cams):
            cols += [(f"x{k}", f"cam{k + 1}_x", 76), (f"y{k}", f"cam{k + 1}_y", 76),
                     (f"s{k}", f"cam{k + 1}_status", 88)]
        self.columns = cols
        self.redraw()

    def _yview(self, *args):
        self.canvas.yview(*args)
        self.redraw()

    def _xview(self, *args):
        self.canvas.xview(*args)
        self.redraw()

    def _on_wheel(self, e) -> str:
        steps = 1 if getattr(e, "num", None) == 4 else -1 if getattr(e, "num", None) == 5 else \
            (int(e.delta / 120) if e.delta else 0)
        self.canvas.yview_scroll(-3 * steps, "units")
        self.redraw()
        return "break"

    def _col_at(self, x) -> int | None:
        left = 0
        for i, (_k, _t, w) in enumerate(self.columns):
            if left <= x < left + w:
                return i
            left += w
        return None

    def _row_at(self, y) -> int | None:
        i = int((y - self.HEAD_H) // self.ROW_H)
        return i if 0 <= i < len(self.get_rows()) and y >= self.HEAD_H else None

    @staticmethod
    def cell_text(row: dict, key: str) -> str:
        if key in ("name", "kind"):
            return row["name"] if key == "name" else row["kind"]
        k = int(key[1:])
        if key[0] == "s":
            return row["status"][k]
        v = row["x"][k] if key[0] == "x" else row["y"][k]
        return "" if not math.isfinite(v) else f"{v:.2f}"

    def redraw(self) -> None:
        c = self.canvas
        try:
            ch = c.winfo_height()
        except tk.TclError:
            return
        rows = self.get_rows()
        total_w = sum(w for _k, _t, w in self.columns)
        c.configure(scrollregion=(0, 0, total_w, self.HEAD_H + len(rows) * self.ROW_H))
        c.delete("cell")                                         # not the cell editor
        top = c.canvasy(0)
        first = max(0, int((top - self.HEAD_H) // self.ROW_H))
        last = min(len(rows), int((top + ch) // self.ROW_H) + 1)
        sel = set(self.selected)
        for i in range(first, last):
            r, y0 = rows[i], self.HEAD_H + i * self.ROW_H
            left = 0
            for key, _title, w in self.columns:
                if key[0] in "xys" and key[1:].isdigit():
                    fill = STATUS_COLORS[r["status"][int(key[1:])]]
                else:
                    fill = "#ffffff"
                if i in sel:
                    fill = "#ffd1ef" if key in ("name", "kind") else fill
                outline = "#c8c8c8"
                c.create_rectangle(left, y0, left + w, y0 + self.ROW_H, fill=fill, outline=outline, tags="cell")
                c.create_text(left + 4, y0 + self.ROW_H / 2, anchor="w", text=self.cell_text(r, key),
                              font=("Arial", 8, "bold" if (i in sel and key == "name") else "normal"), tags="cell")
                left += w
            if i in sel:
                c.create_rectangle(1, y0 + 1, total_w - 1, y0 + self.ROW_H - 1, outline=SELECT_COLOR, width=2, tags="cell")
        # Sticky header at the top of the visible area.
        left = 0
        for key, title, w in self.columns:
            hl = self.sel_cam is not None and key[1:] == str(self.sel_cam) and key[0] in "xys"
            c.create_rectangle(left, top, left + w, top + self.HEAD_H,
                               fill="#c9d6ea" if hl else "#e3e8f0", outline="#a8b2c4", tags="cell")
            c.create_text(left + 4, top + self.HEAD_H / 2, anchor="w", text=title, font=("Arial", 8, "bold"), tags="cell")
            left += w
        if not rows:
            c.create_text(10, top + self.HEAD_H + 14, anchor="w", fill="#777777", font=("Arial", 9),
                          text="No points yet: press 'Detect ArUco', or click in a camera image to add one.", tags="cell")

    def ensure_visible(self, i: int) -> None:
        c = self.canvas
        top, h = c.canvasy(0), c.winfo_height()
        y0 = self.HEAD_H + i * self.ROW_H
        n = max(1, len(self.get_rows()))
        total = self.HEAD_H + n * self.ROW_H
        if y0 < top + self.HEAD_H or y0 + self.ROW_H > top + h:
            c.yview_moveto(max(0.0, (y0 - self.HEAD_H - h / 3) / total))
        self.redraw()

    def _on_click(self, e) -> None:
        self._close_editor(commit=True)
        x, y = self.canvas.canvasx(e.x), self.canvas.canvasy(e.y)
        if y < self.canvas.canvasy(0) + self.HEAD_H:
            return                                                # header
        i, col = self._row_at(y), self._col_at(x)
        if i is None:
            self.selected = []
        elif e.state & 0x4:                                       # Ctrl: toggle
            self.selected = [j for j in self.selected if j != i] if i in self.selected else self.selected + [i]
            self._anchor = i
        elif e.state & 0x1 and self._anchor is not None:          # Shift: range
            a, b = sorted((self._anchor, i))
            self.selected = list(range(a, b + 1))
        else:
            self.selected, self._anchor = [i], i
        key = self.columns[col][0] if col is not None else ""
        self.sel_cam = int(key[1:]) if key[:1] in ("x", "y", "s") and key[1:].isdigit() else None
        self.redraw()
        self.on_select()

    def _on_double(self, e) -> None:
        x, y = self.canvas.canvasx(e.x), self.canvas.canvasy(e.y)
        i, col = self._row_at(y), self._col_at(x)
        if i is None or col is None:
            return
        key, _t, w = self.columns[col]
        if key == "kind":
            return                                                # information only
        left = sum(cw for _k, _tt, cw in self.columns[:col])
        y0 = self.HEAD_H + i * self.ROW_H
        row = self.get_rows()[i]
        if key[0] == "s":
            widget = ttk.Combobox(self.canvas, values=STATUSES, state="readonly", font=("Arial", 8))
            widget.set(row["status"][int(key[1:])])
            widget.bind("<<ComboboxSelected>>", lambda _e: self._close_editor(commit=True))
        else:
            widget = tk.Entry(self.canvas, font=("Arial", 8))
            widget.insert(0, self.cell_text(row, key) if key == "name" else
                          ("" if not math.isfinite(row[key[0]][int(key[1:])]) else repr(row[key[0]][int(key[1:])])))
            widget.select_range(0, tk.END)
            widget.bind("<Return>", lambda _e: self._close_editor(commit=True))
            widget.bind("<FocusOut>", lambda _e: self._close_editor(commit=True))
        widget.bind("<Escape>", lambda _e: self._close_editor(commit=False))
        win = self.canvas.create_window(left, y0, anchor="nw", window=widget, width=w,
                                        height=self.ROW_H, tags="editor")
        self._editor = (widget, win, i, key)
        widget.focus_set()

    def _close_editor(self, commit: bool) -> None:
        ed, self._editor = self._editor, None
        if ed is None:
            return
        widget, win, i, key = ed
        value = widget.get()
        try:
            self.canvas.delete(win)
            widget.destroy()
        except tk.TclError:
            pass
        if commit:
            self.on_edit(i, key, value)
        self.redraw()


# ── node ──────────────────────────────────────────────────────────────────

class MultiCameraPointsNode(BaseNode):
    """One table of named points across up to 6 cameras: ArUco markers are
    detected and named automatically, other points are clicked; the table
    can be edited here or in a spreadsheet (CSV / Excel)."""

    EXECUTION_MODE = ExecutionMode.SYNC
    NODE_TYPE = "multi_camera_points"
    DISPLAY_NAME = "Multi-Camera Points"
    CATEGORY = "process"
    SEARCH_KEYWORDS = ("multi camera", "correspondence", "points", "aruco", "table", "mapping",
                       "stereo", "triangulation", "name")
    NODE_WIDTH = 210
    NODE_HEIGHT = 200

    HELP_TEXT = """\
Multi-Camera Points
===================
One table of NAMED points seen by several cameras (up to 6), so every point
has its image coordinates in every camera in the SAME row -- what tracking,
synchronisation and triangulation need.

INPUTS
  img 1 ... img 6   the reference image (frame 1) of each camera.

OUTPUTS (all in the same row order)
  pts 1 ... pts 6   (N, 2) x, y of each point in camera k; NaN = none
  names             the N point names
  table             (N, 2K) x_c1, y_c1, x_c2, y_c2, ... (K cameras in use)
  visible           (N, K) True where the point has coordinates
  status            (N, K) 0 missing, 1 auto, 2 suggested, 3 confirmed,
                    4 not visible

ARUCO
  'Detect ArUco' finds the markers of the chosen dictionary in all camera
  images. Each marker is ONE point -- its centre, the intersection of its
  two diagonals (the true centre under perspective, independent of the
  corner order) -- named ARUCO_<dictionary>_ID<id>, e.g.
  ARUCO_6X6_1000_ID0123. The same marker gets the same row in every
  camera automatically. Detection fills cells that are missing / auto;
  cells you placed or marked yourself are kept. New rows go to the end.

OTHER POINTS (corner stickers, speckles, strokes ...)
  Click a row, then click the point in each camera image. Clicking with
  no row selected adds a new point P0001, P0002, ... A click within a few
  pixels of a detected marker centre snaps to it.
  Mouse in the images: wheel = zoom, right-drag = pan, 'Fit' = whole image.

TABLE
  Click a row to select it (Ctrl / Shift: several); the camera views centre
  on it. Double-click a cell to edit name, x, y or status.
  Colours: red = missing, green = auto, yellow = suggested,
  blue = confirmed, grey = not visible.
  Buttons: Add point, Delete, Move up / down, Merge 2 rows (one physical
  point that got two names), Mark not visible / Clear for the camera of
  the clicked column.
  Names: letters, digits and _ only, starting with a letter.

SPREADSHEET
  'Export ...' writes CSV or Excel (.xlsx, with a README sheet). Edit it in
  a spreadsheet and 'Import ...' it back: a summary of the changes (and
  any errors, with line numbers) is shown before anything is applied.
  The ROW ORDER is the order downstream nodes use; re-ordering means
  results computed before must be recomputed (the summary warns).

Coordinates follow OpenCV: (0, 0) is the centre of the top-left pixel.
"""

    def __init__(self, node_id: str, canvas: tk.Canvas):
        super().__init__(node_id, canvas)
        self._rows: list[dict] = []
        self._images: list = [None] * MAX_CAMS
        self._markers: list[list] = [[] for _ in range(MAX_CAMS)]
        self._dict_var = tk.StringVar(value="DICT_6X6_1000")
        self._show_names_var = tk.BooleanVar(value=True)
        self._info_var = tk.StringVar(value="")
        self._face_var = tk.StringVar(value="no points")
        self._views: list[CamView] = []
        self._views_frame = None
        self._sheet: Sheet | None = None
        self._shown_cams = 0
        self._update_face()

    # ── pins ───────────────────────────────────────────────────────────
    def get_pin_schema(self) -> PinSchema:
        return PinSchema(
            inputs=[PinDef(f"img_{k}", PinType.IMAGE, f"img {k}", optional=True)
                    for k in range(1, MAX_CAMS + 1)],
            outputs=[PinDef(f"pts_{k}", PinType.ARRAY, f"pts {k}") for k in range(1, MAX_CAMS + 1)]
            + [PinDef("names", PinType.STRING, "names"),
               PinDef("table", PinType.ARRAY, "table"),
               PinDef("visible", PinType.ARRAY, "visible"),
               PinDef("status", PinType.ARRAY, "status")],
        )

    def n_cams(self) -> int:
        """Cameras in use: up to the last one with an image or any data."""
        used = [k for k in range(MAX_CAMS) if self._images[k] is not None]
        used += [k for r in self._rows for k in range(MAX_CAMS) if r["status"][k] != "missing"]
        return max(used) + 1 if used else 1

    def _outputs(self) -> dict:
        n, K = len(self._rows), self.n_cams()
        xy = np.full((MAX_CAMS, n, 2), np.nan)
        status = np.zeros((n, MAX_CAMS), dtype=np.uint8)
        for i, r in enumerate(self._rows):
            for k in range(MAX_CAMS):
                status[i, k] = STATUSES.index(r["status"][k])
                if cell_valid(r, k):
                    xy[k, i] = (r["x"][k], r["y"][k])
        out = {f"pts_{k + 1}": xy[k] for k in range(MAX_CAMS)}
        out["names"] = [r["name"] for r in self._rows]
        out["table"] = np.concatenate([xy[k] for k in range(K)], axis=1) if n else np.zeros((0, 2 * K))
        out["visible"] = np.isfinite(xy[:K, :, 0]).T if n else np.zeros((0, K), dtype=bool)
        out["status"] = status[:, :K]
        return out

    def compute(self, inputs: dict) -> dict:
        changed = False
        for k in range(MAX_CAMS):
            img = inputs.get(f"img_{k + 1}")
            img = img if isinstance(img, np.ndarray) and img.ndim in (2, 3) and img.size else None
            if img is not self._images[k]:
                self._images[k] = img
                changed = True
        if changed and self.is_inspector_open():
            self._sync_views()
        self._update_face()
        return self._outputs()

    def _push(self) -> None:
        """The table changed: refresh the inspector and send the outputs on."""
        self._update_face()
        if self._sheet is not None:
            if self._shown_cams != self.n_cams():
                self._sync_views()
            self._sheet.redraw()
            self._update_points()
        if self._request_downstream is not None:
            self._request_downstream(self.node_id)

    def _update_face(self) -> None:
        K = self.n_cams()
        missing = sum(1 for r in self._rows for k in range(K) if r["status"][k] == "missing")
        self._face_var.set(f"{len(self._rows)} points, {K} camera(s)"
                           + (f"\n{missing} cell(s) missing" if missing else ""))

    # ── node face ──────────────────────────────────────────────────────
    def build_body(self) -> None:
        x, y, w, h = self.x, self.y, self.width, self.height
        self._body_rect = self.canvas.create_rectangle(
            x, y, x + w, y + h, fill="#eef6ee", outline="#3f7a4a", width=2, tags=(self.node_id, "node_body"))
        self._title_item = self.canvas.create_text(
            x + w / 2, y + 12, text=self.DISPLAY_NAME, font=("Arial", 9, "bold"), fill="#24502c",
            tags=(self.node_id,))
        lbl = tk.Label(self.canvas, textvariable=self._face_var, font=("Arial", 8), bg="#eef6ee",
                       fg="#24502c", justify="center")
        self.canvas.create_window(x + w / 2, y + h / 2, window=lbl, tags=(self.node_id,))
        self._canvas_items += [self._body_rect, self._title_item]

    # ── inspector ──────────────────────────────────────────────────────
    def build_inspector(self, parent: tk.Frame) -> None:
        win = self._inspector_win
        if win is not None:
            sw, sh = win.winfo_screenwidth(), win.winfo_screenheight()
            win.geometry(f"{min(1500, sw - 80)}x{min(920, sh - 100)}")
            for seq in ("<Control-h>", "<Control-H>"):
                win.bind(seq, lambda _e: self._open_help())

        bar = tk.Frame(parent)
        bar.pack(fill="x")
        tk.Label(bar, text="ArUco dictionary:", font=("Arial", 9)).pack(side="left")
        dict_cb = ttk.Combobox(bar, textvariable=self._dict_var, values=list(ARUCO_DICT_NAMES),
                               state="readonly", width=20)
        dict_cb.pack(side="left", padx=(4, 6))
        det_btn = tk.Button(bar, text="Detect ArUco", font=("Arial", 9, "bold"), command=self._on_detect)
        det_btn.pack(side="left")
        names_cb = tk.Checkbutton(bar, text="Show names", variable=self._show_names_var,
                                  command=self._update_points)
        names_cb.pack(side="left", padx=(12, 0))
        fit_btn = tk.Button(bar, text="Fit views", font=("Arial", 8), command=lambda: [v.fit() for v in self._views])
        fit_btn.pack(side="left", padx=(6, 0))
        tk.Button(bar, text="Help (Ctrl+H)", font=("Arial", 8), command=self._open_help).pack(side="right")
        add_tooltip(dict_cb, "The dictionary the printed markers belong to (it MUST match them, "
                             "e.g. DICT_6X6_1000).")
        add_tooltip(det_btn, "Find the markers in every camera image. Each marker is one point, "
                             "its centre (diagonal intersection), named ARUCO_<dictionary>_ID<id>; "
                             "the same marker gets the same row in every camera. Cells you placed or "
                             "marked yourself are not overwritten; new rows go to the end.")
        add_tooltip(names_cb, "Write the point names next to the points in the camera images.")
        add_tooltip(fit_btn, "Show the whole image in every camera view.")
        tk.Label(parent, textvariable=self._info_var, font=("Arial", 8), fg="#334455", anchor="w",
                 justify="left").pack(fill="x")

        panes = tk.PanedWindow(parent, orient=tk.VERTICAL, sashwidth=6, sashrelief=tk.RAISED)
        panes.pack(fill="both", expand=True)
        self._views_frame = tk.Frame(panes)
        bottom = tk.Frame(panes)
        panes.add(self._views_frame, minsize=200, stretch="always")
        panes.add(bottom, minsize=180, height=330, stretch="always")

        tk.Label(bottom, font=("Arial", 8), fg="#555555", anchor="w", justify="left", text=(
            "Select a row, then click the point in each camera image (click with no row selected = new "
            "point).  Images: wheel = zoom, right-drag = pan.  Table: double-click a cell to edit; "
            "Ctrl / Shift-click for several rows.")).pack(fill="x")
        btns = tk.Frame(bottom)
        btns.pack(fill="x", pady=(2, 2))
        buttons = (
            ("Add point", self._on_add, "Add an empty row P#### at the end; then click it in the images."),
            ("Delete", self._on_delete, "Delete the selected row(s)."),
            ("Move up", lambda: self._on_move(-1), "Move the selected row(s) up (the row order is the "
                                                   "order downstream nodes use)."),
            ("Move down", lambda: self._on_move(1), "Move the selected row(s) down."),
            ("Merge 2 rows", self._on_merge, "The two selected rows are the same physical point (e.g. it "
                                             "got different names in different cameras): combine them "
                                             "into the first one."),
            ("Mark not visible", self._on_not_visible, "The camera of the clicked column cannot see the "
                                                       "selected point(s): no coordinates (NaN), shown grey."),
            ("Clear cell", self._on_clear, "Remove the selected point(s)' coordinates in the camera of the "
                                           "clicked column (back to 'missing')."),
        )
        for text, cmd, tip in buttons:
            b = tk.Button(btns, text=text, font=("Arial", 8), command=cmd)
            b.pack(side="left", padx=(0, 4))
            add_tooltip(b, tip)
        imp = tk.Button(btns, text="Import ...", font=("Arial", 8), command=self._on_import)
        imp.pack(side="right")
        exp = tk.Button(btns, text="Export ...", font=("Arial", 8), command=self._on_export)
        exp.pack(side="right", padx=(0, 4))
        add_tooltip(exp, "Save the table as Excel (.xlsx, with a README sheet) or CSV, to edit it in a "
                         "spreadsheet.")
        add_tooltip(imp, "Read an edited CSV / Excel table back. A summary of the changes (and any "
                         "errors, with line numbers) is shown first; nothing changes until you confirm.")
        self._sheet = Sheet(bottom, lambda: self._rows, self._on_select, self._on_edit)
        self._sheet.frame.pack(fill="both", expand=True)
        self._sync_views()

    def close_inspector(self) -> None:
        self._views, self._views_frame, self._sheet = [], None, None
        self._shown_cams = 0
        super().close_inspector()

    def _sync_views(self) -> None:
        """(Re)build the camera views when the number of cameras changed,
        and give each its current image."""
        if self._views_frame is None:
            return
        K = self.n_cams()
        if K != self._shown_cams:
            for v in self._views:
                v.frame.destroy()
            self._views = []
            per_row = min(K, 3)
            for k in range(K):
                v = CamView(self._views_frame, k, self._on_pick)
                v.frame.grid(row=k // per_row, column=k % per_row, sticky="nsew", padx=2, pady=2)
                self._views.append(v)
            for col in range(per_row):
                self._views_frame.grid_columnconfigure(col, weight=1, uniform="cam")
            for row in range((K + per_row - 1) // per_row):
                self._views_frame.grid_rowconfigure(row, weight=1)
            self._shown_cams = K
            if self._sheet is not None:
                self._sheet.set_cameras(K)
        for k, v in enumerate(self._views):
            v.markers = self._markers[k]
            v.set_image(self._images[k])
        self._update_points()

    def _update_points(self) -> None:
        sel = set(self._sheet.selected) if self._sheet is not None else set()
        for k, v in enumerate(self._views):
            v.show_names = bool(self._show_names_var.get())
            v.markers = self._markers[k]
            v.points = [(r["x"][k], r["y"][k], r["status"][k], r["name"], i in sel)
                        for i, r in enumerate(self._rows) if cell_valid(r, k)]
            v.render()

    # ── selection / picking / editing ─────────────────────────────────
    def _selected(self) -> list[int]:
        return sorted(self._sheet.selected) if self._sheet is not None else []

    def _select(self, indices: list[int], center: bool = True) -> None:
        if self._sheet is None:
            return
        self._sheet.selected = list(indices)
        if indices:
            self._sheet.ensure_visible(indices[0])
        else:
            self._sheet.redraw()
        self._on_select(center=center)

    def _on_select(self, center: bool = True) -> None:
        sel = self._selected()
        if center and sel:
            r = self._rows[sel[0]]
            for k, v in enumerate(self._views):
                if cell_valid(r, k):
                    v.center_on(r["x"][k], r["y"][k])
        self._update_points()

    def _next_manual_name(self) -> str:
        used = [int(m.group(1)) for r in self._rows for m in [re.fullmatch(r"P(\d+)", r["name"])] if m]
        return f"P{(max(used) + 1 if used else 1):04d}"

    def _on_pick(self, cam: int, u: float, v: float, event) -> None:
        """A click in a camera image: place the selected point there (or a
        new one). Within a few screen pixels of a detected marker centre the
        click snaps to it."""
        view = self._views[cam]
        snapped = None
        for mid, _quad, cen in self._markers[cam]:
            sx, sy = view.to_screen(*cen)
            if (sx - event.x) ** 2 + (sy - event.y) ** 2 <= 12 ** 2:
                snapped = (mid, cen)
                break
        sel = self._selected()
        if len(sel) > 1:
            self._info_var.set("Several rows are selected: select one row to place its point.")
            return
        if snapped is not None:
            u, v = float(snapped[1][0]), float(snapped[1][1])
        if sel:
            i = sel[0]
        else:
            if snapped is not None:
                name = aruco_name(self._dict_var.get(), snapped[0])
                i = next((j for j, r in enumerate(self._rows) if r["name"] == name), None)
                if i is None:
                    self._rows.append(new_row(name, "aruco", {"dictionary": self._dict_var.get(),
                                                              "id": snapped[0]}))
                    i = len(self._rows) - 1
            else:
                self._rows.append(new_row(self._next_manual_name()))
                i = len(self._rows) - 1
        r = self._rows[i]
        r["x"][cam], r["y"][cam], r["status"][cam] = u, v, "confirmed"
        self._info_var.set(f"{r['name']}: camera {cam + 1} = ({u:.2f}, {v:.2f})"
                           + ("  (snapped to the marker centre)" if snapped is not None else ""))
        self._select([i], center=False)
        self._push()

    def _on_edit(self, i: int, key: str, value: str) -> None:
        r = self._rows[i]
        value = value.strip()
        if key == "name":
            if value == r["name"]:
                return
            if not NAME_RE.fullmatch(value):
                messagebox.showwarning("Name", "Use letters, digits and _ only, starting with a letter "
                                               "(e.g. P0012, MARK_A).", parent=self._dlg_parent())
                return
            if any(o["name"] == value for o in self._rows):
                messagebox.showwarning("Name", f"'{value}' is already used. To join two rows that are the "
                                               "same point use 'Merge 2 rows'.", parent=self._dlg_parent())
                return
            r["name"] = value
        elif key[0] == "s":
            k = int(key[1:])
            if value in VALID_STATUSES and not (math.isfinite(r["x"][k]) and math.isfinite(r["y"][k])):
                messagebox.showwarning("Status", f"'{value}' needs x and y. Enter them, or click the point "
                                                 "in the camera image.", parent=self._dlg_parent())
                return
            if value in ("missing", "not_visible"):
                r["x"][k] = r["y"][k] = math.nan
            r["status"][k] = value
        else:
            k, axis = int(key[1:]), key[0]
            if value == "":
                r[axis][k] = math.nan
                if r["status"][k] in VALID_STATUSES:
                    r["x"][k] = r["y"][k] = math.nan
                    r["status"][k] = "missing"
            else:
                try:
                    r[axis][k] = float(value)
                except ValueError:
                    messagebox.showwarning("Coordinate", f"'{value}' is not a number.", parent=self._dlg_parent())
                    return
                if math.isfinite(r["x"][k]) and math.isfinite(r["y"][k]):
                    r["status"][k] = "confirmed"
        self._push()

    def _need_selection(self, n: int | None = None) -> list[int] | None:
        sel = self._selected()
        if not sel or (n is not None and len(sel) != n):
            messagebox.showinfo("Table", "Select " + (f"exactly {n} rows" if n else "one or more rows")
                                + " first (Ctrl / Shift-click for several).", parent=self._dlg_parent())
            return None
        return sel

    def _on_add(self) -> None:
        self._rows.append(new_row(self._next_manual_name()))
        self._select([len(self._rows) - 1], center=False)
        self._info_var.set(f"Added {self._rows[-1]['name']}: now click it in each camera image.")
        self._push()

    def _on_delete(self) -> None:
        sel = self._need_selection()
        if sel is None:
            return
        names = [self._rows[i]["name"] for i in sel]
        if not messagebox.askyesno("Delete", f"Delete {len(sel)} row(s)?\n" + ", ".join(names[:10])
                                   + (" ..." if len(names) > 10 else ""), parent=self._dlg_parent()):
            return
        self._rows = [r for i, r in enumerate(self._rows) if i not in set(sel)]
        self._select([], center=False)
        self._push()

    def _on_move(self, step: int) -> None:
        sel = self._need_selection()
        if sel is None:
            return
        if (step < 0 and sel[0] == 0) or (step > 0 and sel[-1] == len(self._rows) - 1):
            return
        order = list(range(len(self._rows)))
        for i in (sel if step < 0 else reversed(sel)):
            order[i], order[i + step] = order[i + step], order[i]
        self._rows = [self._rows[j] for j in order]
        self._select([i + step for i in sel], center=False)
        self._push()

    def _on_merge(self) -> None:
        sel = self._need_selection(2)
        if sel is None:
            return
        a, b = self._rows[sel[0]], self._rows[sel[1]]
        conflicts = []
        for k in range(MAX_CAMS):
            if cell_valid(b, k) and not cell_valid(a, k):
                a["x"][k], a["y"][k], a["status"][k] = b["x"][k], b["y"][k], b["status"][k]
            elif cell_valid(a, k) and cell_valid(b, k):
                if abs(a["x"][k] - b["x"][k]) > 1e-6 or abs(a["y"][k] - b["y"][k]) > 1e-6:
                    conflicts.append(k + 1)
            elif a["status"][k] == "missing" and b["status"][k] == "not_visible":
                a["status"][k] = "not_visible"
        if not messagebox.askyesno(
                "Merge 2 rows", f"Merge '{b['name']}' into '{a['name']}' and delete '{b['name']}'?"
                + (f"\n\nBoth have coordinates in camera(s) {conflicts}: the ones of '{a['name']}' are kept."
                   if conflicts else ""), parent=self._dlg_parent()):
            return
        del self._rows[sel[1]]
        self._select([sel[0]], center=False)
        self._info_var.set(f"Merged into {a['name']}.")
        self._push()

    def _cam_of_click(self) -> int | None:
        cam = self._sheet.sel_cam if self._sheet is not None else None
        if cam is None:
            messagebox.showinfo("Table", "Click a cell in the camera's columns (camK_x / y / status) of the "
                                         "row first, so it is clear which camera is meant.",
                                parent=self._dlg_parent())
        return cam

    def _on_not_visible(self) -> None:
        self._set_cells("not_visible")

    def _on_clear(self) -> None:
        self._set_cells("missing")

    def _set_cells(self, status: str) -> None:
        sel = self._need_selection()
        cam = self._cam_of_click() if sel else None
        if sel is None or cam is None:
            return
        for i in sel:
            r = self._rows[i]
            r["x"][cam] = r["y"][cam] = math.nan
            r["status"][cam] = status
        self._push()

    # ── ArUco ──────────────────────────────────────────────────────────
    def _on_detect(self) -> None:
        dict_name = self._dict_var.get()
        counts, new_rows, kept, failed = [], 0, 0, []
        for k in range(MAX_CAMS):
            img = self._images[k]
            if img is None:
                self._markers[k] = []
                continue
            try:
                self._markers[k] = detect_aruco(np.asarray(img), dict_name)
            except Exception as e:
                failed.append(f"camera {k + 1}: {e}")
                continue
            counts.append(f"cam{k + 1}: {len(self._markers[k])}")
            for mid, _quad, cen in self._markers[k]:
                name = aruco_name(dict_name, mid)
                row = next((r for r in self._rows if r["name"] == name), None)
                if row is None:
                    row = new_row(name, "aruco", {"dictionary": dict_name, "id": mid})
                    self._rows.append(row)
                    new_rows += 1
                if row["status"][k] in ("missing", "auto", "suggested"):
                    row["x"][k], row["y"][k], row["status"][k] = float(cen[0]), float(cen[1]), "auto"
                else:
                    kept += 1
        if not counts and not failed:
            self._info_var.set("No camera image yet: connect img 1 ... (the reference frames).")
            return
        self._info_var.set(f"{dict_name}: markers found {', '.join(counts)};  {new_rows} new row(s)"
                           + (f";  {kept} cell(s) you edited were kept" if kept else "")
                           + (";  " + "; ".join(failed) if failed else ""))
        self._push()

    # ── export / import ────────────────────────────────────────────────
    def _dlg_parent(self):
        win = getattr(self, "_inspector_win", None)
        try:
            if win is not None and win.winfo_exists():
                return win
        except tk.TclError:
            pass
        return self.canvas.winfo_toplevel()

    def _on_export(self) -> None:
        path = self.ask_save_as_filename(
            parent=self._dlg_parent(), title="Export the point table", defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx"), ("CSV", "*.csv")])
        if not path:
            return
        try:
            (write_csv if path.lower().endswith(".csv") else write_xlsx)(path, self._rows, self.n_cams())
        except Exception as e:
            messagebox.showerror("Export", f"Cannot write the file:\n{e}", parent=self._dlg_parent())
            return
        self._info_var.set(f"Exported {len(self._rows)} rows to {Path(path).name}")

    def _on_import(self) -> None:
        path = self.ask_open_filename(
            parent=self._dlg_parent(), title="Import a point table",
            filetypes=[("Tables", "*.xlsx *.csv"), ("Excel", "*.xlsx"), ("CSV", "*.csv"), ("All files", "*.*")])
        if not path:
            return
        try:
            rows, _n, errors = parse_table(read_table_file(path))
        except Exception as e:
            messagebox.showerror("Import", f"Cannot read the file:\n{e}", parent=self._dlg_parent())
            return
        if errors:
            messagebox.showerror("Import", f"{len(errors)} problem(s) -- nothing was changed:\n\n"
                                 + "\n".join(errors[:12]) + ("\n..." if len(errors) > 12 else ""),
                                 parent=self._dlg_parent())
            return
        summary = change_summary(self._rows, rows)
        if not messagebox.askyesno("Import", f"{Path(path).name}:\n\n" + "\n".join(summary)
                                   + "\n\nApply these changes?", parent=self._dlg_parent()):
            return
        old_meta = {r["name"]: r["meta"] for r in self._rows}
        for r in rows:
            r["meta"] = old_meta.get(r["name"], r["meta"])
        self._rows = rows
        self._select([], center=False)
        self._info_var.set(f"Imported {len(rows)} rows from {Path(path).name}")
        self._push()

    def _open_help(self) -> None:
        top = tk.Toplevel(self._dlg_parent())
        top.title("Multi-Camera Points - Help")
        top.geometry("720x640")
        txt = tk.Text(top, font=("Courier", 9), wrap="word")
        txt.pack(fill="both", expand=True)
        txt.insert("1.0", self.HELP_TEXT)
        txt.configure(state="disabled")
        top.bind("<Escape>", lambda _e: top.destroy())

    def get_help_text(self) -> str:
        return self.HELP_TEXT

    # ── project ────────────────────────────────────────────────────────
    def get_params(self) -> dict:
        params = {"dictionary": self._dict_var.get(), "show_names": bool(self._show_names_var.get())}
        compact = [[r["name"], r["kind"], r["meta"],
                    [None if not math.isfinite(v) else v for v in r["x"]],
                    [None if not math.isfinite(v) else v for v in r["y"]], r["status"]] for r in self._rows]
        # However long: long params go to the project file's Data sheet
        # (project_io.py); 'rows_file' CSV side files of older projects are still read.
        params["rows"] = compact
        return params

    def set_params(self, params: dict) -> None:
        self._dict_var.set(str(params.get("dictionary", "DICT_6X6_1000")))
        self._show_names_var.set(bool(params.get("show_names", True)))
        rows = []
        if params.get("rows_file"):
            base = get_project_directory()
            try:
                rows, _n, errors = parse_table(read_table_file(str(Path(base) / params["rows_file"])))
                if errors:
                    raise ValueError(errors[0])
                for r in rows:
                    r["kind"] = params.get("kinds", {}).get(r["name"], r["kind"])
                    r["meta"] = params.get("meta", {}).get(r["name"], {})
            except Exception as e:
                rows = []
                self._info_var.set(f"The saved table {params['rows_file']} could not be read: {e}")
        else:
            for name, kind, meta, xs, ys, st in params.get("rows", []):
                r = new_row(str(name), str(kind), meta)
                for k in range(MAX_CAMS):
                    r["x"][k] = math.nan if xs[k] is None else float(xs[k])
                    r["y"][k] = math.nan if ys[k] is None else float(ys[k])
                    r["status"][k] = st[k] if st[k] in STATUSES else "missing"
                rows.append(r)
        self._rows = rows
        self._update_face()
        if self._sheet is not None:
            self._push()
