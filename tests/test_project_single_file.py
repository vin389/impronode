"""Everything a project holds is saved in the ONE .xlsx file (no side files)
and comes back exactly."""

import gc
import json
import time
import tkinter as tk

import cv2
import numpy as np
import pytest

from node_editor import project_data
from node_editor.nodes.camera_calibration_node import CameraCalibrationNode, new_row
from node_editor.project_context import set_project_file_path
from node_editor.project_io import CHUNK, DATA_SHEET, load_project, save_project

from project_roundtrip import restore, save_and_load
from test_camera_calibration_search import POSES, render_view


# ── project_io ────────────────────────────────────────────────────────────

def test_long_params_and_data_go_to_the_hidden_data_sheet(tmp_path):
    long_list = [f"images/frame_{k:06d}.png" for k in range(20000)]          # ~500,000 characters
    text = "x" * (3 * CHUNK + 5)
    path = tmp_path / "p.xlsx"
    save_project(path, [{"node_id": "n0", "node_type": "t", "x": 1, "y": 2, "width": 3, "height": 4,
                         "params": {"files": long_list, "short": 1}}], [], {"canvas_width": 10},
                 data={"n0/blob": text})
    assert [p.name for p in tmp_path.iterdir()] == ["p.xlsx"]
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    assert wb[DATA_SHEET].sheet_state == "hidden"
    assert all(len(str(c)) <= 32767 for row in wb[DATA_SHEET].iter_rows(values_only=True) for c in row if c)
    wb.close()
    nodes, _links, meta = load_project(path)
    assert nodes[0]["params"] == {"files": long_list, "short": 1}
    assert meta["data"] == {"n0/blob": text}
    assert "extra_tables" not in meta                 # the Data sheet is not a user table
    assert meta["canvas_width"] == 10


@pytest.mark.parametrize("arr", [np.arange(24, dtype=np.int16).reshape(2, 3, 4), np.array([]),
                                 np.random.default_rng(0).normal(size=(500, 300)).astype(np.float32),
                                 np.array([True, False]), np.array(3.5)])
def test_arrays_are_exact(arr):
    text = project_data.encode_arrays({"a": arr})
    back = project_data.decode_arrays(text)["a"]
    assert back.dtype == arr.dtype and back.shape == arr.shape and np.array_equal(back, arr)


def test_unchanged_arrays_are_not_encoded_again(monkeypatch):
    arr = np.ones((200, 200))
    project_data.begin_save()
    project_data.put_array("n", "a", arr)
    calls = []
    real = project_data.encode_arrays
    monkeypatch.setattr(project_data, "encode_arrays", lambda a: calls.append(1) or real(a))
    project_data.put_array("n", "a", arr.copy())
    assert calls == []
    project_data.put_array("n", "a", arr + 1)
    assert calls == [1]
    project_data.end_save()


# ── the calibration node: settings, corners, search, runs, sensitivity ──────

@pytest.fixture(scope="module")
def root():
    gc.collect()
    r = tk.Tk()
    r.withdraw()
    yield r
    r.destroy()


def test_calibration_node_round_trip_through_one_xlsx(tmp_path, root):
    img_dir = tmp_path / "images"
    img_dir.mkdir()
    paths = []
    for k, (rv, tv) in enumerate(POSES[:5]):
        p = img_dir / f"view{k}.png"
        cv2.imwrite(str(p), render_view(rv, tv))
        paths.append(str(p))
    proj_dir = tmp_path / "project"
    proj_dir.mkdir()
    try:
        canvas = tk.Canvas(root, width=300, height=200)
        node = CameraCalibrationNode("node_3", canvas)
        node.build_body()
        set_project_file_path(str(proj_dir / "calib.xlsx"))
        node._cb_manual_files = paths
        for key, value in (("square", "1"), ("square_y", "1.02"), ("ncol", "11"), ("nrow", "8"), ("threads", "4")):
            node._det_vars[key].set(value)
        node._force_lsq_var.set(True)
        node._on_open_search()
        sv = node._search_vars
        for key, value in (("cols", "9 11"), ("rows", "8"), ("widths", "0"), ("adaptive", "on"),
                           ("normalize", "off"), ("fast", "on")):
            sv[key].set(value)
        node._on_search_start()
        t0 = time.time()
        while not node._search["state"]["finished"]:
            root.update()
            time.sleep(0.02)
            assert time.time() - t0 < 300
        node._on_search_apply()
        node._use_points_var.set(False)
        node._solve_sync()
        node._fix_vars["k3"].set(False)
        node._solve_sync()
        assert len(node._runs) == 2 and node._runs[-1]["full"]["cb_poses"]
        board = node._cb["cands"][1]["path"]
        node.rows = [new_row(node._new_id(), f"P{k}", X=str(k), Y="2", Z="?", u=10.5 + k, v=20.25)
                     for k in range(1500)]                                    # far beyond one Excel cell
        node._loo_rows = [{"label": "board 2", "kind": "board", "key": board, "K": node._res["K"],
                           "D": node._res["D"], "rms": 0.1, "success": True, "error": None, "seconds": 1.0},
                          {"label": "row 3", "kind": "row", "key": node.rows[2]["id"], "K": node._res["K"],
                           "D": node._res["D"], "rms": 0.2, "success": True, "error": None, "seconds": 1.0}]
        node._loo_base = {"K": node._res["K"], "D": node._res["D"]}
        node._loo_text = "Done: 2 solves"

        params, stored = save_and_load(proj_dir / "calib.xlsx", node, "camera_calibration")
        assert "cb_file" not in params and "rows_file" not in params

        node2 = CameraCalibrationNode("node_3", canvas)
        node2.build_body()
        restore(node2, params, stored)

        assert [(r["name"], r["X"], r["Z"], r["u"], r["v"]) for r in node2.rows] == \
               [(r["name"], r["X"], r["Z"], r["u"], r["v"]) for r in node.rows]
        assert node2._square_xy() == (1.0, 1.02)
        assert node2._force_lsq_var.get() is True
        assert node2._search_cfg == node._search_cfg
        a = {c["path"]: c for c in node._cb["cands"]}
        assert {c["path"] for c in node2._cb["cands"]} == set(a)
        for c in node2._cb["cands"]:
            assert np.array_equal(c["corners"], a[c["path"]]["corners"])        # exact, not rounded
            assert np.array_equal(c["rc"], a[c["path"]]["rc"])
            assert np.array_equal(c["support"], a[c["path"]]["support"])
        assert node2._cb_use == node._cb_use
        # the settings search (even one not applied) comes back
        s1, s2 = node._search, node2._search
        assert s2 is not None and s2["settings"] == s1["settings"] and s2["paths"] == s1["paths"]
        for r1, r2 in zip(s1["state"]["results"], s2["state"]["results"]):
            assert r1["found"] == r2["found"]
            if r1["found"]:
                for k in ("corners", "rc", "support", "spread"):
                    assert np.array_equal(r1[k], r2[k])
                assert r1["settings_merged"] == r2["settings_merged"]
        # result, runs (with their poses) and sensitivity rows
        assert node2._res["K"] == node._res["K"] and node2._res["cb_poses"].keys() == node._res["cb_poses"].keys()
        assert [r["K"] for r in node2._runs] == [r["K"] for r in node._runs]
        assert node2._runs[0]["full"]["cb_poses"].keys() == node._runs[0]["full"]["cb_poses"].keys()
        assert [r["key"] for r in node2._loo_rows] == [board, node2.rows[2]["id"]]
        assert node2._loo_base == node._loo_base and node2._loo_text == "Done: 2 solves"
    finally:
        set_project_file_path(None)


def test_old_projects_with_side_files_still_load(tmp_path, root):
    """A project saved before the Data sheet (chessboard .npz side file)."""
    try:
        canvas = tk.Canvas(root)
        node = CameraCalibrationNode("n", canvas)
        node.build_body()
        corners = np.random.default_rng(1).uniform(0, 500, (3, 6, 9, 2)).astype(np.float32)
        side = "old.n.chessboard.npz"
        np.savez_compressed(tmp_path / side, paths=np.array(["a.png", "b.png", "c.png"]),
                            source_index=np.arange(3), corners=corners, sharpness=np.ones(3), coverage=np.ones(3),
                            center_dist=np.zeros(3), size=np.array([[640, 480]] * 3), pick_dist=-np.ones(3),
                            quality=-np.ones(3), use=np.array([True, False, True]),
                            cfg=np.array(json.dumps({"ncol": 9, "nrow": 6, "square": 25.0})))
        set_project_file_path(str(tmp_path / "old.xlsx"))
        node.set_params({"cb_file": side, "rows": []})
        assert len(node._cb["cands"]) == 3
        assert np.array_equal(node._cb["cands"][1]["corners"], corners[1].reshape(-1, 2))
        assert node._cb["cands"][1]["rc"].tolist()[:2] == [[0, 0], [0, 1]]
    finally:
        set_project_file_path(None)
